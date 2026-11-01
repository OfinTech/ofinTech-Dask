"""Browser smoke test of the real UI against a running engine (desktop mode) and a fake model.

    cd ui && node build.mjs
    python e2e/ui_smoke.py            # starts the fake model and the engine, drives Chromium, saves screenshots

Needs: Python Playwright with Chromium (`pip install playwright && playwright install chromium`).
"""

import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
TOKEN = "e2e" + "k" * 40
PORT = 47822


def wait_line(proc: subprocess.Popen, needle: str, timeout: float = 60) -> None:
    end = time.time() + timeout
    while time.time() < end:
        line = proc.stdout.readline().decode()
        if needle in line:
            return
    raise RuntimeError(f"timed out waiting for {needle}")


def sample_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Management P&L"
    ws.append(["Management accounts"])
    ws.append([])
    ws.append(["Section", "Account", "Jan 2026", "Feb 2026", "Mar 2026"])
    ws.append(["Income", "Freight revenue", 180000, 176500, 191200])
    ws.append(["Cost of sales", "Fuel", 61000, 59800, 64400])
    ws.append(["Expenses", "Rent", 14000, 14000, 14000])
    wb.save(path)


def main() -> int:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="desk-e2e-"))
    out.mkdir(parents=True, exist_ok=True)
    data = Path(tempfile.mkdtemp(prefix="desk-data-"))
    fake = subprocess.Popen([sys.executable, str(ROOT / "e2e" / "fake_llm.py")])
    env = {**os.environ, "DESK_DEV": "1", "DESK_DATA_DIR": str(data), "DESK_UI_DIR": str(ROOT / "ui" / "dist"),
           "DESK_LOCAL_TOKEN": TOKEN, "DESK_PORT": str(PORT)}
    engine = subprocess.Popen([sys.executable, "-m", "desk_engine"], cwd=ROOT / "engine", env=env,
                              stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    errors: list[str] = []
    try:
        wait_line(engine, "DESK_READY")
        base = f"http://localhost:{PORT}"
        xlsx = out / "pl.xlsx"
        sample_workbook(xlsx)
        with sync_playwright() as p:
            browser = p.chromium.launch()
            pg = browser.new_page(viewport={"width": 1360, "height": 860})
            pg.on("pageerror", lambda e: errors.append(str(e)))
            try:
                run_steps(pg, base, out, xlsx)
            except Exception:
                pg.screenshot(path=str(out / "failure.png"), full_page=True)
                raise
            browser.close()
    finally:
        engine.terminate()
        try:
            log = engine.stdout.read().decode(errors="replace") if engine.stdout else ""
            (out / "engine.log").write_text(log)
        except Exception:
            pass
        fake.terminate()
    if errors:
        print("Page errors:", errors)
        return 1
    print(f"UI smoke test passed; screenshots in {out}")
    return 0


def hide_toasts(pg, hide: bool = True) -> None:
    """Keeps toasts out of a screenshot without hiding them from later waits."""
    if hide:
        pg.evaluate("document.head.insertAdjacentHTML('beforeend', '<style id=\"hide-toasts\">.toasts{visibility:hidden}</style>')")
    else:
        pg.evaluate("document.getElementById('hide-toasts')?.remove()")


def run_steps(pg, base, out, xlsx):
    pg.goto(f"{base}/#token={TOKEN}")
    pg.locator("label:has(> span:text-is('Company or workspace name')) input").fill("E2E Ltd")
    pg.get_by_role("button", name="Create workspace").click()
    pg.wait_for_selector("text=No general ledger yet")   # the overview is the landing page; it asks for a ledger first
    pg.goto(f"{base}/#/settings/models")
    pg.get_by_role("button", name="Add provider").click()
    m = pg.locator(".modal")
    m.locator("label:has(> span:text-is('Provider')) select").select_option("openai_compatible")
    m.locator("label:has(> span:text-is('Display name')) input").fill("Fake model")
    m.locator("label:has(> span:text-is('Base URL')) input").fill("http://127.0.0.1:47999/v1")
    pg.get_by_role("button", name="Save").click()
    pg.wait_for_selector("strong:text-is('Fake model')")
    pg.goto(f"{base}/#/chat")
    pg.get_by_role("button", name="New conversation").first.click()
    pg.wait_for_selector("text=What do you need to know?")
    pg.wait_for_timeout(300)
    pg.get_by_label("Message").fill("What is the NPV and IRR of these cash flows at 8%?")
    pg.keyboard.press("Enter")
    pg.wait_for_selector("text=Check before relying on", timeout=20000)
    pg.screenshot(path=str(out / "chat.png"))
    pg.get_by_label("Message").fill("Remind me to post the accruals")
    pg.keyboard.press("Enter")
    pg.get_by_role("button", name="Approve and run").click(timeout=20000)
    pg.wait_for_selector("text=Done.", timeout=20000)
    pg.goto(f"{base}/#/data")
    pg.get_by_role("button", name="Import Excel or CSV").click()
    pg.set_input_files("input[type=file]", str(xlsx))
    pg.wait_for_selector("text=Periods are in columns")
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector("text=Export CSV")
    # documents: upload a policy, search it, ask the assistant and see citations
    import docx

    d = docx.Document()
    d.add_heading("Travel and expenses", level=1)
    d.add_paragraph("Hotels may cost at most 180 per night. Economy class for flights under six hours.")
    d.add_heading("Approvals", level=1)
    d.add_paragraph("Expenses above 1,000 need approval from the finance manager.")
    d.save(out / "policy.docx")
    pg.goto(f"{base}/#/documents")
    pg.set_input_files("input[type=file][multiple]", str(out / "policy.docx"))
    pg.wait_for_selector("td:has-text('keyword')", timeout=20000)
    pg.get_by_label("Search documents").fill("hotel limit per night")
    pg.get_by_role("button", name="Search", exact=True).click()
    pg.wait_for_selector(".hit")
    pg.screenshot(path=str(out / "documents.png"))
    pg.goto(f"{base}/#/chat")
    pg.get_by_role("button", name="New conversation").first.click()
    pg.wait_for_selector("text=What do you need to know?")
    pg.wait_for_timeout(300)
    pg.get_by_label("Message").fill("What does our expense policy say about hotels?")
    pg.keyboard.press("Enter")
    pg.wait_for_selector("ol.sources", timeout=20000)
    pg.wait_for_selector("text=finance manager's approval", timeout=20000)
    pg.screenshot(path=str(out / "chat-citations.png"))
    pg.locator("ol.sources .linklike").first.click()
    pg.wait_for_selector(".doc-text")
    pg.screenshot(path=str(out / "document-view.png"))
    # attach a file in the assistant and ask about it
    pg.goto(f"{base}/#/chat")
    pg.get_by_role("button", name="New conversation").first.click()
    pg.wait_for_selector("text=What do you need to know?")
    pg.wait_for_timeout(300)
    import docx as _docx
    att = _docx.Document()
    att.add_heading("Delegation of authority", level=1)
    att.add_paragraph("Payments above 25,000 need two approvers: the CFO and the finance manager.")
    att.save(out / "authority.docx")
    pg.set_input_files("input[aria-label='Attach files']", str(out / "authority.docx"))
    pg.wait_for_selector(".attach-chip:has-text('authority.docx')")
    pg.get_by_label("Message").fill("What does the attached file say about approvals?")
    pg.keyboard.press("Enter")
    pg.wait_for_selector(".msg.user a.attach-chip:has-text('authority.docx')", timeout=20000)
    pg.wait_for_selector("ol.sources", timeout=20000)
    pg.screenshot(path=str(out / "chat-attachment.png"))
    # evidence-linked workbook: create it in chat, open it, select figures, wait for the reviewer
    pg.goto(f"{base}/#/chat")
    pg.get_by_role("button", name="New conversation").first.click()
    pg.wait_for_selector("text=What do you need to know?")
    pg.wait_for_timeout(300)
    pg.get_by_label("Message").fill("Build the Q1 management workbook")
    pg.keyboard.press("Enter")
    pg.get_by_role("button", name="Open Q1 2026 management accounts").click(timeout=20000)
    pg.wait_for_selector(".workpaper")
    pg.locator("button.fig.sourced").nth(2).click()
    pg.wait_for_selector("text=Underlying rows", timeout=10000)
    pg.wait_for_selector("text=February fuel looks low", timeout=20000)
    pg.screenshot(path=str(out / "file-evidence.png"))
    pg.get_by_role("tab", name="Adjustments").click()
    pg.locator("button.fig.unverified").first.click()
    pg.wait_for_selector("text=No source found")
    pg.screenshot(path=str(out / "file-unverified.png"))
    pg.goto(f"{base}/#/files")
    pg.wait_for_selector("text=without source")
    pg.goto(f"{base}/#/connections")
    pg.wait_for_selector("text=Other accounting system")
    pg.screenshot(path=str(out / "connections.png"), full_page=True)
    # webhooks under Connections: an outgoing one delivers a signed test event to the fake receiver, an incoming one
    # shows its endpoint and secret once
    pg.get_by_role("button", name="Add outgoing webhook").click()
    m = pg.locator(".modal")
    m.locator("label:has(> span:text-is('Name')) input").fill("Data warehouse")
    m.locator("label:has(> span:text-is('Address')) input").fill("http://127.0.0.1:47999/hook")
    m.locator("label.check:has-text('approval') input").check()
    m.locator("label.check:has-text('test events') input").check()
    m.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector(".modal >> text=Webhook secret")
    secret = pg.locator(".modal code.excel").first.inner_text()
    assert secret.startswith("who_") and len(secret) > 30, secret
    pg.screenshot(path=str(out / "webhook-secret.png"))
    pg.get_by_role("button", name="I have stored it").click()
    pg.get_by_role("button", name="Send a test").click()
    pg.wait_for_selector("text=Test event queued")
    for _ in range(60):
        pg.wait_for_timeout(500)
        pg.get_by_role("button", name="Deliveries (1)").click()
        pg.wait_for_selector(".modal >> text=test.ping")
        if pg.locator(".modal .tag", has_text="delivered").count():
            break
        pg.keyboard.press("Escape")
    else:
        raise AssertionError("the test event was not delivered")
    pg.locator(".modal tr.clickable").first.click()
    pg.wait_for_selector(".modal pre >> text=Test event from OfinTech Desk")
    pg.screenshot(path=str(out / "webhook-deliveries.png"))
    pg.keyboard.press("Escape")
    pg.get_by_role("button", name="Add incoming webhook").click()
    m = pg.locator(".modal")
    m.locator("label:has(> span:text-is('Name')) input").fill("Bank portal")
    m.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector(".modal code:has-text('/api/hooks/in/wh_')")
    pg.get_by_role("button", name="I have stored it").click()
    pg.wait_for_selector("strong:text-is('Bank portal')")
    hide_toasts(pg)
    pg.evaluate("[...document.querySelectorAll('h2')].find((h) => h.textContent === 'Webhooks').scrollIntoView({ block: 'start' })")
    pg.wait_for_timeout(300)
    pg.screenshot(path=str(out / "webhooks.png"))
    hide_toasts(pg, False)
    # a new workflow starts as a draft and can be saved before it is complete
    pg.goto(f"{base}/#/workflows")
    pg.get_by_role("button", name="New workflow").click()
    pg.wait_for_selector("text=This workflow is a draft")
    pg.get_by_role("button", name="Save", exact=True).click()
    pg.wait_for_selector("text=Workflow saved")
    pg.locator("label:has(> span:text-is('Instruction')) textarea").fill("List overdue invoices and propose next steps.")
    pg.wait_for_selector("text=This workflow is a draft", state="detached")
    pg.screenshot(path=str(out / "workflow-draft.png"))
    pg.goto(f"{base}/#/dashboards")
    pg.get_by_role("button", name="New dashboard").click()
    pg.locator(".modal select").nth(1).select_option(index=1)
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector(".widget svg")
    pg.screenshot(path=str(out / "dashboard.png"), full_page=True)
    pg.get_by_role("button", name="Refresh").click()
    pg.wait_for_selector("text=Figures recalculated from the latest data")
    pg.wait_for_selector("text=data as of")
    for path in ("workflows", "skills", "connections", "inbox", "settings/members", "settings/files"):
        pg.goto(f"{base}/#/{path}")
        pg.wait_for_timeout(600)
    # search across everything
    pg.goto(f"{base}/#/search")
    pg.get_by_label("Search the workspace").fill("hotel per night")
    pg.get_by_role("button", name="Search", exact=True).click()
    pg.wait_for_selector(".result-group")
    pg.screenshot(path=str(out / "search.png"))
    # memory added by hand
    pg.goto(f"{base}/#/memory")
    pg.get_by_role("button", name="Add memory").click()
    pg.get_by_label("What should the assistant remember?").fill("Accruals are posted by working day 3.")
    pg.locator(".panel .btn.primary:text-is('Save')").click()
    pg.wait_for_selector("text=Accruals are posted by working day 3.")
    pg.get_by_label("Search memory").fill("accru")
    pg.wait_for_timeout(600)
    assert pg.locator(".memory-item").count() == 1
    pg.get_by_label("Search memory").fill("nothing like this")
    pg.wait_for_selector("text=Nothing matches")
    pg.get_by_label("Search memory").fill("")
    pg.screenshot(path=str(out / "memory.png"))
    # the user's own Excel file
    pg.goto(f"{base}/#/files")
    pg.set_input_files("input[type=file][accept='.xlsx,.docx']", str(xlsx))
    pg.wait_for_selector("text=Original, as uploaded")
    pg.get_by_role("button", name="Show contents").click()
    pg.wait_for_selector("pre.peek")
    pg.screenshot(path=str(out / "user-file.png"))
    # workflow templates, webhook trigger and schedule preview
    pg.goto(f"{base}/#/workflows")
    pg.locator(".template-card:has-text('Month-end pack') button").click()
    pg.wait_for_selector("text=Next runs")
    pg.screenshot(path=str(out / "workflow-template.png"), full_page=True)
    # delete a conversation
    pg.goto(f"{base}/#/chat")
    pg.wait_for_selector(".thread-item")
    before = pg.locator(".thread-item").count()
    pg.locator(".thread-item").first.hover()
    pg.locator(".thread-item .thread-del").first.click()
    pg.locator(".modal-foot .btn.danger, .modal-foot .btn.primary").last.click()
    for _ in range(50):
        if pg.locator(".thread-item").count() == before - 1:
            break
        pg.wait_for_timeout(100)
    assert pg.locator(".thread-item").count() == before - 1
    # cost centres: a general ledger imported from Excel, rules, and the budget read from the customer's file
    sys.path.insert(0, str(ROOT / "engine"))
    from tests.test_cost_centres import budget_xlsx, ledger_xlsx

    (out / "gl.xlsx").write_bytes(ledger_xlsx())
    (out / "Riyadh budget.xlsx").write_bytes(budget_xlsx())
    pg.goto(f"{base}/#/data")
    pg.get_by_role("button", name="Import Excel or CSV").click()
    pg.set_input_files("input[type=file]", str(out / "gl.xlsx"))
    pg.wait_for_selector(".modal-foot .btn.primary")
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector("text=Export CSV")
    pg.goto(f"{base}/#/costcentres")
    pg.get_by_role("button", name="New cost centre").click()
    m = pg.locator(".modal")
    m.locator("label:has(> span:text-is('Name')) input").fill("Riyadh tower")
    m.locator("label:has(> span:text-is('Starts')) input").fill("2026-05-01")
    m.locator("label:has(> span:text-is('Planned end')) input").fill("2026-12-31")
    m.locator("label:has(> span:text-is('Tag in the ledger')) input").first.fill("Project: Riyadh tower")
    pg.get_by_role("button", name="Add a rule").click()
    m.locator("label:has(> span:text-is('Accounts')) input").nth(1).fill("5")
    m.locator("label:has(> span:text-is('Accounts')) input").nth(1).press("Enter")
    m.locator("label:has(> span:text-is('Words in the description')) input").nth(1).fill("riyadh tower")
    m.locator("label:has(> span:text-is('Words in the description')) input").nth(1).press("Enter")
    pg.wait_for_selector("text=These rules pick up 5 ledger lines")
    pg.screenshot(path=str(out / "cost-centre-editor.png"))
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector("text=When the budget runs out")
    pg.set_input_files("label:has-text('Read budget from a file') input[type=file]", str(out / "Riyadh budget.xlsx"))
    pg.wait_for_selector("a:text-is('Budget!C4')")
    pg.wait_for_selector("text=At this pace the budget lasts until")
    pg.add_style_tag(content=".toasts { visibility: hidden; }")
    pg.set_viewport_size({"width": 1360, "height": 2400})  # the page scrolls inside the app frame
    pg.wait_for_timeout(400)
    pg.screenshot(path=str(out / "cost-centre.png"))
    pg.set_viewport_size({"width": 1360, "height": 860})
    pg.get_by_role("button", name="Review").first.click()
    pg.wait_for_selector(".modal >> text=BANK-7")
    pg.keyboard.press("Escape")
    pg.goto(f"{base}/#/costcentres")
    pg.wait_for_selector("td >> text=Riyadh tower")
    pg.screenshot(path=str(out / "cost-centres.png"))
    # the general ledger template, filled in and uploaded from the Cost centres page; both templates download
    from openpyxl import load_workbook

    pg.get_by_role("button", name="Import ledger").click()
    with pg.expect_download() as dl:
        pg.get_by_role("button", name="General ledger template").click()
    tpl = out / dl.value.suggested_filename
    dl.value.save_as(str(tpl))
    assert tpl.name == "OfinTech-Desk-general-ledger-template.xlsx", tpl.name
    with pg.expect_download() as dl2:
        pg.get_by_role("button", name="Budget template").click()
    assert dl2.value.suggested_filename == "OfinTech-Desk-budget-template.xlsx"
    wb = load_workbook(tpl)
    sheet = wb["General ledger"]
    sheet.delete_rows(2, sheet.max_row)
    src = load_workbook(out / "gl.xlsx").active
    for row in src.iter_rows(min_row=2, values_only=True):
        sheet.append(list(row))
    wb.save(out / "GL October.xlsx")
    pg.set_input_files("label:has-text('Upload ledger file') input[type=file]", str(out / "GL October.xlsx"))
    pg.wait_for_selector("text=now in the ledger “GL October · General ledger”")
    hide_toasts(pg)
    pg.wait_for_timeout(300)
    pg.screenshot(path=str(out / "cost-centres-import.png"))
    hide_toasts(pg, False)
    assert pg.locator("select[aria-label='Ledger to replace'] option", has_text="GL October").count() == 1
    # financial analysis: a 15-month ledger, the KPI scorecard, a KPI explained, account roles, a forecast approved
    from tests.test_financial_analysis import ledger_xlsx as fa_ledger

    (out / "trading-gl.xlsx").write_bytes(fa_ledger())
    pg.goto(f"{base}/#/data")
    pg.get_by_role("button", name="Import Excel or CSV").click()
    pg.set_input_files("input[type=file]", str(out / "trading-gl.xlsx"))
    pg.wait_for_selector(".modal-foot .btn.primary")
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector("text=Export CSV")
    pg.goto(f"{base}/#/analysis")
    pg.wait_for_selector("text=Growth and profitability")
    trading = pg.locator("select[aria-label='General ledger'] option", has_text="trading").get_attribute("value")
    pg.get_by_label("General ledger").select_option(trading)
    pg.wait_for_selector(".kpi-card:has-text('Gross margin') >> text=60.0%")
    pg.add_style_tag(content=".toasts { visibility: hidden; }")
    pg.set_viewport_size({"width": 1360, "height": 2200})
    pg.wait_for_timeout(300)
    pg.screenshot(path=str(out / "analysis.png"))
    pg.set_viewport_size({"width": 1360, "height": 860})
    pg.locator(".kpi-card:has-text('Days sales outstanding')").click()
    pg.wait_for_selector(".modal >> text=Largest accounts")
    pg.locator(".modal label:has(> span:text-is('Target (at most)')) input").fill("30")
    pg.locator(".modal label:has(> span:text-is('Warning level')) input").fill("45")
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector(".kpi-card.watch:has-text('Days sales outstanding')")
    pg.locator(".kpi-card:has-text('Days sales outstanding')").click()
    pg.wait_for_selector(".modal >> text=1200 Trade receivables")
    pg.screenshot(path=str(out / "analysis-kpi.png"))
    pg.keyboard.press("Escape")
    pg.goto(f"{base}/#/analysis/roles")
    pg.wait_for_selector("text=Confirm all suggestions")
    pg.screenshot(path=str(out / "analysis-roles.png"))
    # executive overview and the drill-down to one journal entry, the month kept at every step
    pg.goto(f"{base}/#/analysis")
    pg.get_by_label("General ledger").select_option(trading)
    pg.wait_for_selector(".kpi-card:has-text('Gross margin') >> text=60.0%")
    pg.goto(f"{base}/#/overview")
    pg.wait_for_selector(".health >> text=measures on track")
    pg.wait_for_selector(".exec-tile:has-text('Net profit')")
    assert pg.locator(".exec-risk:has-text('Days sales outstanding')").count() == 0  # the executive name is used
    pg.wait_for_selector(".exec-risk:has-text('Days to collect a sale')")
    pg.screenshot(path=str(out / "overview.png"))
    pg.locator(".exec-tile:has-text('Net profit')").click()
    pg.wait_for_selector("h1:text-is('Profitability')")
    assert "month=" in pg.evaluate("location.hash")
    month = pg.evaluate("new URLSearchParams(location.hash.split('?')[1]).get('month')")
    pg.locator("table.pl a", has_text="Other operating expenses").click()
    pg.wait_for_selector("h1:text-is('Other operating expenses')")
    assert pg.evaluate("location.hash").endswith(f"month={month}")
    pg.locator("table.ledger tbody tr.clickable").first.click()
    pg.wait_for_selector("text=to open its journal entry")
    pg.screenshot(path=str(out / "overview-transactions.png"))
    pg.locator("table.ledger tbody tr.clickable").first.click()
    pg.wait_for_selector("h1:has-text('Journal entry')")
    pg.wait_for_selector("text=balanced")
    pg.wait_for_selector("text=dataset.import")
    assert pg.locator(".crumbs a").count() == 4, pg.locator(".crumbs").inner_text()
    pg.screenshot(path=str(out / "overview-entry.png"))
    pg.locator(".crumbs a", has_text="Overview").click()
    pg.wait_for_selector(".health")
    assert pg.evaluate("location.hash").endswith(f"month={month}")
    # the view switch only changes the rail; the executive view shows five sections
    pg.get_by_role("tab", name="Executive").click()
    assert pg.locator(".nav a").count() == 5 + 2, pg.locator(".nav").inner_text()   # plus Connections and Settings
    pg.get_by_role("tab", name="Accounting").click()
    assert pg.locator(".nav a").count() == 14 + 2
    pg.get_by_role("tab", name="Finance").click()
    pg.goto(f"{base}/#/analysis/forecasts")
    pg.get_by_role("button", name="New forecast").click()
    pg.locator(".modal label:has(> span:text-is('Name')) input").fill("Rolling forecast")
    pg.locator(".modal-foot .btn.primary").click()
    pg.wait_for_selector("text=Submit for approval")
    pg.get_by_label("Method for Revenue").select_option("growth")
    pg.locator(".assumptions tr:has-text('Revenue') input").first.fill("2")
    pg.wait_for_timeout(800)
    pg.get_by_role("button", name="Submit for approval").click()
    pg.get_by_role("button", name="Approve").click()
    pg.get_by_role("button", name="Make it the plan").click()
    pg.wait_for_selector(".tag:text-is('plan')")
    pg.set_viewport_size({"width": 1360, "height": 2200})
    pg.wait_for_timeout(300)
    pg.screenshot(path=str(out / "forecast.png"))
    pg.set_viewport_size({"width": 1360, "height": 860})
    # month-end close: a ledger uploaded, September prepared, the statement matched, entries decided
    from tests.test_close import ledger_xlsx as close_ledger, statement_csv, BOOK

    (out / "Close GL.xlsx").write_bytes(close_ledger())
    (out / "september-statement.csv").write_bytes(statement_csv(None))
    pg.goto(f"{base}/#/costcentres")
    pg.get_by_role("button", name="Import ledger").click()
    pg.set_input_files("label:has-text('Upload ledger file') input[type=file]", str(out / "Close GL.xlsx"))
    pg.wait_for_selector("text=now in the ledger “Close GL · General ledger”")
    pg.goto(f"{base}/#/close")
    pg.wait_for_selector("text=Prepare a month")
    gl = pg.locator("select option", has_text="Close GL").get_attribute("value")
    pg.locator(".close-start select").select_option(gl)
    pg.locator(".close-start input[type=month]").fill("2026-09")
    pg.get_by_role("button", name="Prepare the close").click()
    pg.wait_for_selector("text=Ready to close")
    pg.wait_for_selector(".close-item:has-text('Accrue Rent')")
    pg.set_input_files("label:has-text('Upload statement') input[type=file]", str(out / "september-statement.csv"))
    pg.wait_for_selector("text=Adjusted ledger balance")
    pg.wait_for_selector(".close-item:has-text('Bank charge')")
    rent = pg.locator(".close-item:has-text('Accrue Rent')")
    rent.get_by_role("button", name="Approve").click()
    pg.wait_for_selector(".close-item.approved:has-text('Accrue Rent')")
    dep = pg.locator(".close-item:has-text('Charge Depreciation')")
    dep.get_by_role("button", name="Reject").click()
    dep.get_by_label("Why is this entry not needed?").fill("Vehicle sold in August")
    dep.locator("form button:text-is('Reject')").click()
    pg.wait_for_selector(".close-item.rejected:has-text('Charge Depreciation')")
    fees = pg.locator(".close-item:has-text('Professional fees')")
    fees.get_by_label("Explanation").fill("Due diligence for the Jeddah branch, one-off")
    fees.get_by_role("button", name="Explain").click()
    pg.wait_for_selector(".close-item.explained:has-text('Professional fees')")
    hide_toasts(pg)
    pg.set_viewport_size({"width": 1360, "height": 3000})
    pg.wait_for_timeout(400)
    pg.screenshot(path=str(out / "close.png"))
    pg.set_viewport_size({"width": 1360, "height": 860})
    hide_toasts(pg, False)
    with pg.expect_download() as dl:
        pg.get_by_role("button", name="Journal (Excel)").click()
    assert dl.value.suggested_filename == "OfinTech-Desk-close-journal-2026-09.xlsx"
    with pg.expect_download() as dl2:
        pg.get_by_role("button", name="Close pack").click()
    assert dl2.value.suggested_filename == "OfinTech-Desk-close-pack-2026-09.xlsx"
    pg.goto(f"{base}/#/close")
    pg.wait_for_selector("td:has-text('September 2026')")
    pg.screenshot(path=str(out / "close-list.png"))
    # Arabic (right to left), Spanish and French
    pg.goto(f"{base}/#/settings/account")
    pg.get_by_role("button", name="العربية").click()
    pg.wait_for_selector("html[dir=rtl][lang=ar]", state="attached")
    pg.wait_for_selector("text=اللغة")
    pg.evaluate("document.fonts.ready")
    font = pg.evaluate("getComputedStyle(document.body).fontFamily")
    assert font.startswith('"IBM Plex Sans Arabic"'), font
    assert pg.evaluate("document.fonts.check('600 16px \"IBM Plex Sans Arabic\"', 'مرحبا')")
    pg.screenshot(path=str(out / "ar-settings.png"))
    for path, needle in (("chat", "المساعد"), ("files", "الملفات"), ("workflows", "البدء من قالب"),
                         ("costcentres", "مراكز التكلفة"), ("analysis", "النمو والربحية"), ("close", "إقفال الشهر"),
                         ("overview", "على المسار"), ("overview/profit", "الربحية")):
        pg.goto(f"{base}/#/{path}")
        pg.wait_for_selector(f"text={needle}")
        pg.wait_for_timeout(300)
        pg.screenshot(path=str(out / f"ar-{path.replace('/', '-')}.png"))
    pg.goto(f"{base}/#/settings/account")
    pg.get_by_role("button", name="Español").click()
    pg.wait_for_selector("html[dir=ltr][lang=es]", state="attached")
    pg.goto(f"{base}/#/memory")
    pg.wait_for_selector("text=Memoria")
    pg.screenshot(path=str(out / "es-memory.png"))
    pg.goto(f"{base}/#/settings/account")
    pg.get_by_role("button", name="Français").click()
    pg.goto(f"{base}/#/search")
    pg.wait_for_selector("text=Recherche")
    pg.screenshot(path=str(out / "fr-search.png"))
    pg.goto(f"{base}/#/settings/account")
    pg.get_by_role("button", name="English").click()
    pg.wait_for_selector("html[lang=en]", state="attached")


if __name__ == "__main__":
    sys.exit(main())
