"""Month-end close on a small ledger worked out by hand.

June to September 2026. Rent (5,000) and depreciation (1,000) were posted every month until August and are missing in
September; professional fees jump to 9,000 in September; customer B, invoiced 10,000 each month, is not invoiced in
September. The September bank statement has a bank charge and a receipt nobody recorded, and a cheque in the ledger
has not cleared yet.
"""

from __future__ import annotations

import csv
import io
import unittest
from datetime import date
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from .helpers import Harness

BANK = ("1000", "Bank", "Bank")
AR = ("1200", "Receivables", "Receivable")
VEH = ("1500", "Vehicles", "Asset")
ACC_DEP = ("1590", "Accumulated depreciation", "Asset")
AP = ("2000", "Payables", "Payable")
ACCRUED = ("2100", "Accrued expenses", "Liability")
CAPITAL = ("3000", "Capital", "Equity")
SALES = ("4000", "Sales", "Income")
SALARIES = ("6000", "Salaries", "Expense")
RENT = ("6100", "Rent", "Expense")
FEES = ("6400", "Professional fees", "Expense")
DEP = ("6500", "Depreciation", "Expense")

ROWS: list[list] = []


def entry(day: str, ref: str, party: str | None, *lines: tuple) -> None:
    for acc, dr, cr, desc in lines:
        ROWS.append([date.fromisoformat(day), ref, "Journal", acc[0], acc[1], acc[2], dr or None, cr or None, desc, party, None])


entry("2026-06-01", "OPEN", None, (BANK, 100000, 0, "Opening"), (VEH, 48000, 0, "Opening"), (ACCRUED, 0, 3000, "Opening"),
      (CAPITAL, 0, 145000, "Opening"))
for m in ("06", "07", "08"):
    entry(f"2026-{m}-05", f"INV-A{m}", "Customer A", (AR, 30000, 0, "Freight"), (SALES, 0, 30000, "Freight"))
    entry(f"2026-{m}-06", f"INV-B{m}", "Customer B", (AR, 10000, 0, "Storage"), (SALES, 0, 10000, "Storage"))
    entry(f"2026-{m}-20", f"RCPT-A{m}", "Customer A", (BANK, 30000, 0, "Receipt"), (AR, 0, 30000, "Receipt"))
    entry(f"2026-{m}-21", f"RCPT-B{m}", "Customer B", (BANK, 10000, 0, "Receipt"), (AR, 0, 10000, "Receipt"))
    entry(f"2026-{m}-01", f"RENT{m}", "Landlord", (RENT, 5000, 0, "Warehouse rent"), (BANK, 0, 5000, "Rent"))
    entry(f"2026-{m}-28", f"SAL{m}", None, (SALARIES, 20000, 0, "Salaries"), (BANK, 0, 20000, "Salaries"))
    entry(f"2026-{m}-30", f"DEP{m}", None, (DEP, 1000, 0, "Depreciation"), (ACC_DEP, 0, 1000, "Depreciation"))
entry("2026-07-15", "FEE07", "Audit LLP", (FEES, 1000, 0, "Tax advice"), (BANK, 0, 1000, "Tax advice"))
entry("2026-08-25", "BILL-S1", "Steel Co", (FEES, 2000, 0, "Survey"), (AP, 0, 2000, "Survey"))
# September: no rent, no depreciation, no invoice for B; a large fee
entry("2026-09-05", "INV-A09", "Customer A", (AR, 30000, 0, "Freight"), (SALES, 0, 30000, "Freight"))
entry("2026-09-10", "RCPT-A09", "Customer A", (BANK, 30000, 0, "Receipt"), (AR, 0, 30000, "Receipt"))
entry("2026-09-15", "FEE09", "Advisory LLP", (FEES, 9000, 0, "Due diligence"), (BANK, 0, 9000, "Due diligence"))
entry("2026-09-28", "SAL09", None, (SALARIES, 20000, 0, "Salaries"), (BANK, 0, 20000, "Salaries"))
entry("2026-09-30", "PAY-S1", "Steel Co", (AP, 2000, 0, "Cheque 1042"), (BANK, 0, 2000, "Cheque 1042"))
entry("2026-10-02", "INV-A10", "Customer A", (AR, 30000, 0, "Freight"), (SALES, 0, 30000, "Freight"))

# bank balance at 30 September from the ledger
BOOK = 100000 + 3 * (30000 + 10000 - 5000 - 20000) - 1000 + 30000 - 9000 - 20000 - 2000


def ledger_xlsx() -> bytes:
    from desk_engine.connectors.accounting import GL_COLUMNS

    wb = Workbook()
    ws = wb.active
    ws.append([c["name"] for c in GL_COLUMNS])
    for r in ROWS:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def statement_csv(closing: float | None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Value Date", "Details", "Money out", "Money in", "Balance"])
    bal = BOOK - 30000 + 9000 + 20000 + 2000   # balance at 31 August (the cheque is in the ledger only)
    rows = [("31/08/2026", "Opening balance", "", "", bal)]
    for d, text, out, inc in (("11/09/2026", "TRANSFER CUSTOMER A", "", "30,000.00"),
                              ("16/09/2026", "ADVISORY LLP", "9,000.00", ""),
                              ("28/09/2026", "PAYROLL SEPTEMBER", "20,000.00", ""),
                              ("29/09/2026", "TRANSFER FROM ACME TRADING", "", "1,500.00"),
                              ("30/09/2026", "MONTHLY SERVICE CHARGE", "25.00", "")):
        bal += (float(inc.replace(",", "")) if inc else 0) - (float(out.replace(",", "")) if out else 0)
        rows.append((d, text, out, inc, round(bal, 2) if closing is None else ""))
    if closing is not None:
        rows[-1] = (*rows[-1][:4], closing)
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode()


class Close(unittest.TestCase):
    def setUp(self) -> None:
        self.h = Harness()
        self.ws = self.h.setup_workspace()
        up = self.h.api("POST", f"/api/w/{self.ws}/cost-centres/ledger/upload", files={"file": ("GL.xlsx", ledger_xlsx())})
        self.did = up["dataset"]["id"]

    def tearDown(self) -> None:
        self.h.close()

    def item(self, view: dict, key: str) -> dict:
        return next(i for i in view["items"] if i["data"]["key"] == key)

    def test_close_end_to_end(self) -> None:
        h, ws = self.h, self.ws
        with patch("desk_engine.close.service._today", return_value="2026-11-15"):
            ov = h.api("GET", f"/api/w/{ws}/close")
        self.assertEqual(ov["suggested"], {"dataset_id": self.did, "month": "2026-10"})
        with patch("desk_engine.close.service._today", return_value="2026-10-29"):
            self.assertEqual(h.api("GET", f"/api/w/{ws}/close")["suggested"]["month"], "2026-09")
        h.api("POST", f"/api/w/{ws}/close", json={"month": "2026-12"}, expect=422)
        v = h.api("POST", f"/api/w/{ws}/close", json={"month": "2026-09"})
        self.assertEqual((v["month"], v["status"], v["month_end"]), ("2026-09", "open", "2026-09-30"))
        self.assertEqual(v["health"]["unbalanced_count"], 0)
        self.assertEqual(v["materiality"], 500.0)        # 1% of 40,000 a month, at least 500
        # missing recurring costs: rent as an accrual to the accrued-expenses account, depreciation to accumulated
        # depreciation; salaries were posted, fees are irregular
        entries = [i for i in v["items"] if i["kind"] in ("accrual", "depreciation")]
        self.assertEqual([(i["kind"], i["data"]["amount"]) for i in entries], [("accrual", 5000.0), ("depreciation", 1000.0)])
        rent = entries[0]["data"]
        self.assertEqual([(ln["account_code"], ln["debit"], ln["credit"]) for ln in rent["lines"]],
                         [("6100", 5000.0, 0.0), ("2100", 0.0, 5000.0)])
        self.assertIsNone(rent["warning"])
        self.assertEqual([x["amount"] for x in rent["history"]], [5000.0, 5000.0, 5000.0])
        self.assertEqual(len(rent["evidence"]), 3)
        dep = entries[1]["data"]
        self.assertEqual([ln["account_code"] for ln in dep["lines"]], ["6500", "1590"])
        # variances: fees up (9,000 against an average of 1,000), sales down 10,000; rent is not repeated here
        var = {i["data"]["account"]: i["data"] for i in v["items"] if i["kind"] == "variance"}
        self.assertEqual(set(var), {"Professional fees", "Sales"})
        self.assertEqual((var["Professional fees"]["average"], var["Professional fees"]["difference"]), (1000.0, 8000.0))
        self.assertEqual(var["Professional fees"]["evidence"][0]["entry"], "FEE09")
        self.assertEqual((var["Sales"]["difference"], var["Sales"]["good"]), (-10000.0, False))
        unb = [i for i in v["items"] if i["kind"] == "unbilled"]
        self.assertEqual([(i["data"]["party"], i["data"]["amount"]) for i in unb], [("Customer B", 10000.0)])
        checks = {c["key"]: c["status"] for c in v["readiness"]["checks"]}
        self.assertEqual(checks, {"balanced": "ok", "roles": "warn", "entries": "todo", "bank": "todo",
                                  "variances": "todo", "unbilled": "todo"})
        self.assertEqual(v["readiness"]["score"], 25)

        # the bank statement: matched within days; the cheque is outstanding; the charge becomes an entry; the
        # unknown receipt is to review; with the closing balance, nothing is unexplained
        r = v["id"]
        h.api("POST", f"/api/w/{ws}/close/{r}/statement", files={"file": ("sep.csv", b"just,some\ntext,1\n")}, expect=422)
        v = h.api("POST", f"/api/w/{ws}/close/{r}/statement", files={"file": ("sep.csv", statement_csv(None))})
        b = v["bank"]
        self.assertEqual((b["statement_lines"], b["matched"]), (5, 3))
        self.assertEqual([x["entry"] for x in b["ledger_only"]], ["PAY-S1"])
        self.assertEqual(b["book_balance"], BOOK)
        self.assertEqual(b["closing_balance"], BOOK + 2000 + 1500 - 25)
        self.assertEqual(b["difference"], 0.0)
        fee = next(i for i in v["items"] if i["kind"] == "bank_fee")
        self.assertEqual([(ln["account"], ln["debit"], ln["credit"]) for ln in fee["data"]["lines"]],
                         [("Bank charges", 25.0, 0.0), ("Bank", 0.0, 25.0)])
        self.assertIn("not in this ledger", fee["data"]["warning"])
        unknown = next(i for i in v["items"] if i["kind"] == "bank_unknown")
        self.assertEqual(unknown["data"]["amount"], 1500.0)
        self.assertEqual({c["key"]: c["status"] for c in v["readiness"]["checks"]}["bank"], "todo")

        # decisions: approve rent with a different amount, reject depreciation (needs a reason), fix the fee account
        rent_id = entries[0]["id"]
        h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{entries[1]['id']}", json={"status": "rejected"}, expect=422)
        h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{rent_id}", json={"amount": 5250})
        v = h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{rent_id}", json={"status": "approved", "note": "Rent +5% from September"})
        self.assertEqual([ln["debit"] + ln["credit"] for ln in self.item(v, rent["key"])["data"]["lines"]], [5250.0, 5250.0])
        h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{entries[1]['id']}",
              json={"status": "rejected", "note": "Vehicle sold in August"})
        h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{fee['id']}",
              json={"lines": [{"index": 0, "account": "Bank fees", "account_code": "6900"}], "status": "approved"})
        h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{unknown['id']}", json={"status": "reviewed", "note": "Deposit for October job"})
        for i in v["items"]:
            if i["kind"] == "variance":
                h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{i['id']}", json={"status": "explained"}, expect=422)
                h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{i['id']}", json={"status": "explained", "note": "One-off"})
        v = h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{unb[0]['id']}", json={"status": "reviewed", "note": "Contract ended"})
        checks = {c["key"]: c["status"] for c in v["readiness"]["checks"]}
        self.assertEqual(checks, {"balanced": "ok", "roles": "warn", "entries": "ok", "bank": "ok", "variances": "ok",
                                  "unbilled": "ok"})
        self.assertEqual(v["readiness"]["score"], 92)
        self.assertEqual(v["journal"], {"entries": 2, "total": 5275.0})

        # refreshing keeps the decisions and the edited amount
        v = h.api("POST", f"/api/w/{ws}/close/{r}/refresh", json={})
        self.assertEqual((self.item(v, rent["key"])["status"], self.item(v, rent["key"])["data"]["amount"]), ("approved", 5250.0))

        # the journal: GL template columns, numbered, balanced; the close pack opens with its sheets
        jr = h.client.get(f"/api/w/{ws}/close/{r}/journal.csv")
        rows = list(csv.reader(io.StringIO(jr.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0][:4], ["Date", "Entry", "Type", "Account code"])
        self.assertEqual([(x[0], x[1], x[3], x[6], x[7]) for x in rows[1:]],
                         [("2026-09-30", "CLOSE-2026-09-001", "6100", "5250.0", ""),
                          ("2026-09-30", "CLOSE-2026-09-001", "2100", "", "5250.0"),
                          ("2026-09-30", "CLOSE-2026-09-002", "6900", "25.0", ""),
                          ("2026-09-30", "CLOSE-2026-09-002", "1000", "", "25.0")])
        wb = load_workbook(io.BytesIO(h.client.get(f"/api/w/{ws}/close/{r}/pack.xlsx").content))
        self.assertEqual(wb.sheetnames, ["Summary", "Proposed entries", "Bank reconciliation", "Variances", "To review", "Evidence"])
        self.assertEqual(wb["Summary"]["B4"].value, 0.92)
        self.assertTrue(wb["Proposed entries"]["I2"].hyperlink.location or wb["Proposed entries"]["I2"].hyperlink.target)
        self.assertEqual(wb["Bank reconciliation"]["B11"].value, 0.0)
        self.assertEqual(len(load_workbook(io.BytesIO(h.client.get(f"/api/w/{ws}/close/{r}/journal.xlsx").content))["Journal"]["A"]), 6)

        # posting the rent in the accounting system and re-importing resolves the item
        ROWS.append([date(2026, 9, 30), "RENT09", "Journal", "6100", "Rent", "Expense", 5250, None, "Rent", "Landlord", None])
        ROWS.append([date(2026, 9, 30), "RENT09", "Journal", "2100", "Accrued expenses", "Liability", None, 5250, "Rent", None, None])
        try:
            h.api("POST", f"/api/w/{ws}/cost-centres/ledger/upload", files={"file": ("GL.xlsx", ledger_xlsx())},
                  data={"dataset_id": self.did})
        finally:
            del ROWS[-2:]
        v = h.api("POST", f"/api/w/{ws}/close/{r}/refresh", json={})
        self.assertEqual(self.item(v, rent["key"])["status"], "resolved")
        self.assertEqual(v["journal"]["entries"], 1)

        # closing is for administrators and freezes the month
        h.api("POST", f"/api/w/{ws}/close/{r}/status", json={"status": "closed"})
        h.api("POST", f"/api/w/{ws}/close/{r}/refresh", json={}, expect=409)
        h.api("PATCH", f"/api/w/{ws}/close/{r}/items/{unknown['id']}", json={"status": "open"}, expect=409)
        v = h.api("POST", f"/api/w/{ws}/close/{r}/status", json={"status": "open"})
        self.assertEqual(v["status"], "open")
        # the assistant can say where the close stands
        h.llm.script = [[("month_end_close", {"month": "2026-09"})], "Done."]
        t = h.api("POST", f"/api/w/{ws}/threads", json={})
        h.api("POST", f"/api/w/{ws}/threads/{t['id']}/messages", json={"text": "Where is the September close?"})
        tool = next(m for m in h.wait_run(t["id"])["messages"] if m["role"] == "tool")
        self.assertIn("Month-end close 2026-09", tool["content"])
        audit = [a["action"] for a in h.engine.db.all("SELECT action FROM audit_log WHERE workspace_id = ?", (ws,))]
        self.assertIn("close.closed", audit)

    def test_statement_closing_balance_and_overnight_template(self) -> None:
        h, ws = self.h, self.ws
        v = h.api("POST", f"/api/w/{ws}/close", json={"month": "2026-09", "dataset_id": self.did})
        v = h.api("POST", f"/api/w/{ws}/close/{v['id']}/statement", files={"file": ("sep.csv", statement_csv(BOOK))},
                  data={"closing_balance": ""})
        # the file's own closing balance was wrong by 3,475: the difference is shown, not hidden
        self.assertEqual(v["bank"]["difference"], -3475.0)
        self.assertEqual({c["key"]: c["status"] for c in v["readiness"]["checks"]}["bank"], "warn")
        v = h.api("PUT", f"/api/w/{ws}/close/{v['id']}/statement", json={"closing_balance": BOOK + 3475})
        self.assertEqual(v["bank"]["difference"], 0.0)
        # the workflow template prepares the close unattended and passes the link to the notification
        from desk_engine.workflows.templates import get_template

        tpl = get_template("close_overnight")
        wf = h.api("POST", f"/api/w/{ws}/workflows", json={"name": tpl["name"], "definition": tpl["definition"],
                                                            "enabled": True})
        with patch("desk_engine.close.service._today", return_value="2026-11-15"):
            run = h.wait_workflow(h.api("POST", f"/api/w/{ws}/workflows/{wf['id']}/run", json={})["run_id"])
        self.assertEqual(run["status"], "done", run)
        row = h.engine.db.one("SELECT * FROM close_runs WHERE workspace_id = ? AND month = ?", (ws, "2026-10"))
        self.assertEqual(row["prepared_via"], "workflow")
        note = h.engine.db.one("SELECT * FROM notifications WHERE workspace_id = ? AND kind = 'workflow' ORDER BY "
                               "created_at DESC", (ws,))
        self.assertIn(f"/close/{row['id']}", note["title"] + (note["link"] or ""))


if __name__ == "__main__":
    unittest.main()
