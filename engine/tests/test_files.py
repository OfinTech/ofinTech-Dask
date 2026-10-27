"""Evidence-linked files: formulas, sources, Excel and Word output, the reviewer, versions and drill-down."""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from desk_engine.errors import Invalid
from desk_engine.files import expr as X

from .helpers import Harness

CSV = "Period,Account,Scenario,Amount\n" + "\n".join(
    f"2026-{m:02d}-01,{acct},{sc},{amt}"
    for m in (1, 2, 3)
    for acct, sc, amt in [("Revenue", "Actual", 50000 + m * 1000), ("Revenue", "Budget", 52000),
                          ("Rent", "Actual", 4500), ("Rent", "Budget", 4200),
                          ("Travel", "Actual", 3100 + m * 250), ("Travel", "Budget", 0)]) + "\n"

QUERY_ARGS = {"dataset": "Actuals", "measures": [{"column": "Amount", "agg": "sum"}],
              "group_by": [{"column": "Account"}, {"column": "Scenario"}],
              "filters": [{"column": "Period", "op": "between", "value": ["2026-01-01", "2026-03-31"]}]}

WORKBOOK = {
    "title": "Q1 2026 actual vs budget",
    "sheets": [
        {"name": "Variance", "blocks": [
            {"type": "heading", "text": "Q1 2026: actual against budget"},
            {"type": "table", "id": "va", "title": "By account", "query": QUERY_ARGS,
             "pivot": {"column": "Scenario", "row_total": False},
             "columns": [{"label": "Account"}, {"label": "Actual"}, {"label": "Budget"},
                         {"label": "Variance", "formula": "actual - budget"},
                         {"label": "Variance %", "formula": "(actual - budget) / budget", "format": "percent"}]},
            {"type": "kpis", "items": [
                {"label": "Revenue variance", "value": "[va:Revenue:variance]", "format": "money"},
                {"label": "Cost share of revenue", "value": "([va:Rent:actual] + [va:Travel:actual]) / va.Revenue.actual",
                 "format": "percent"}]},
        ]},
        {"name": "Notes & adjustments", "blocks": [
            {"type": "table", "id": "adj", "title": "Proposed accruals",
             "rows": [["Rent", 13500], ["Audit fee", 7250.5]],
             "columns": [{"label": "Item"}, {"label": "Amount", "format": "money"},
                         {"label": "With VAT 15%", "formula": "ROUND(amount * 1.15, 2)", "format": "money"}]},
        ]},
    ],
}

REPORT = {
    "title": "Q1 2026 management commentary",
    "sheets": WORKBOOK["sheets"][:1],
    "sections": [{"heading": "Results", "blocks": [
        {"type": "paragraph", "text": "Revenue was {{va.Revenue.actual}} against a budget of {{[va:Revenue:budget]}}, "
                                      "a variance of {{va.Revenue.variance}} ({{va.Revenue.variance / va.Revenue.budget | percent}})."},
        {"type": "bullets", "items": ["Rent stayed at 13,500.00 for the quarter.",
                                      "Marketing spend reached 48,213.77, which needs explaining."]},
        {"type": "table", "table": "va"}]}],
}


def soffice_values(data: bytes) -> dict[str, dict[str, object]]:
    """Recalculates a workbook in LibreOffice and returns {sheet: {coord: value}}."""
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in.xlsx"
        src.write_bytes(data)
        env = {**os.environ, "HOME": tmp}
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                        str(Path(tmp) / "out"), str(src)], check=True, capture_output=True, timeout=120, env=env)
        wb = load_workbook(Path(tmp) / "out" / "in.xlsx", data_only=True)
        return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row} for ws in wb.worksheets}


class Formulas(unittest.TestCase):
    def test_parse_compile_evaluate(self) -> None:
        ast = X.parse("(actual - budget) / budget")
        addr = {"actual": "B5", "budget": "C5"}
        self.assertEqual(X.to_excel(ast, lambda n: addr[n.key]), '=IFERROR((B5-C5)/C5,"")')
        self.assertAlmostEqual(X.evaluate(ast, lambda n: {"actual": 110, "budget": 100}[n.key]), 0.1)
        self.assertIsNone(X.evaluate(ast, lambda n: {"actual": 110, "budget": 0}[n.key]))  # Excel: #DIV/0! → ""
        self.assertEqual(X.evaluate(X.parse("a - b"), lambda n: {"a": 5, "b": None}[n.key]), 5)  # blank is zero
        self.assertEqual(X.to_excel(X.parse("a - (b - c)"), lambda n: n.key.upper() + "1"), "=A1-(B1-C1)")
        self.assertEqual(X.to_excel(X.parse("-(a + b) * 2"), lambda n: n.key.upper() + "1"), "=-(A1+B1)*2")
        self.assertEqual(X.evaluate(X.parse("ROUND(2.5, 0)"), lambda n: 0), 3)  # half away from zero, like Excel
        self.assertEqual(X.evaluate(X.parse("ROUND(-2.345, 2)"), lambda n: 0), -2.35)
        ref = X.parse("[pl:Gross profit:amount] / pl.total.amount")
        self.assertEqual(X.references(ref), [X.Ref("pl", "Gross profit", "amount"), X.Ref("pl", "total", "amount")])
        for bad in ("__import__('os')", "a ** 2", "open(1)", "a; b", "[x]", "1 +"):
            with self.assertRaises(Invalid):
                X.parse(bad)


class FilesApi(unittest.TestCase):
    def setUp(self) -> None:
        self.h = Harness()
        self.ws = self.h.setup_workspace()
        h, ws = self.h, self.ws
        h.api("PUT", f"/api/w/{ws}/files-settings", json={"auto_review": False})
        up = h.api("POST", f"/api/w/{ws}/imports", files=[("file", ("actuals.csv", CSV.encode(), "text/csv"))])
        h.api("POST", f"/api/w/{ws}/imports/{up['upload_id']}/commit", json={"name": "Actuals"})

    def tearDown(self) -> None:
        self.h.close()

    def chat(self, text: str) -> dict:
        h, ws = self.h, self.ws
        t = h.api("POST", f"/api/w/{ws}/threads", json={"title": "Q1"})
        h.api("POST", f"/api/w/{ws}/threads/{t['id']}/messages", json={"text": text})
        return h.wait_run(t["id"])

    def file_id(self, data: dict) -> str:
        for m in data["messages"]:
            for b in (m.get("payload") or {}).get("blocks") or []:
                if b.get("type") == "link" and b.get("kind") == "file":
                    return b["id"]
        raise AssertionError("no file link in the conversation: " + str([m["content"][:300] for m in data["messages"]]))

    def test_workbook_sources_formulas_and_review(self) -> None:
        h, ws = self.h, self.ws
        h.llm.script = [[("query_dataset", QUERY_ARGS)], [("create_workbook", WORKBOOK)], "The workbook is ready."]
        data = self.chat("Build me a Q1 actual vs budget workbook")
        self.assertEqual(data["run"]["status"], "done", data["messages"][-1]["content"])
        fid = self.file_id(data)
        f = h.api("GET", f"/api/w/{ws}/files/{fid}")
        model = f["model"]
        va = model["sheets"][0]["blocks"][1]
        keys = [c["key"] for c in va["columns"]]
        self.assertEqual(keys, ["account", "actual", "budget", "variance", "variance_pct"])
        rows = {r[0]["v"]: r for r in va["rows"]}
        self.assertEqual(rows["Revenue"][1]["v"], 156000)          # 51k + 52k + 53k
        self.assertEqual(rows["Revenue"][3]["v"], 0)               # 156,000 − 156,000
        self.assertEqual(rows["Travel"][1]["v"], 10800)             # 3,350 + 3,600 + 3,850
        self.assertIsNone(rows["Travel"][4]["v"])                   # budget 0 → blank, like Excel
        self.assertEqual(rows["Rent"][1]["st"], "sourced")
        self.assertIn("Account = Rent", rows["Rent"][1]["n"])
        self.assertIn("over 3 rows", rows["Rent"][1]["n"])
        self.assertEqual(rows["Rent"][3]["x"], "=B5-C5")
        self.assertEqual(rows["Rent"][4]["x"], '=IFERROR((B5-C5)/C5,"")')
        total = va["total"]
        self.assertEqual(total[1]["x"], "=SUM(B5:B7)")
        self.assertEqual(total[4]["x"], '=IFERROR((B8-C8)/C8,"")')  # ratio of totals, not a sum of ratios
        self.assertAlmostEqual(total[4]["v"], (total[1]["v"] - total[2]["v"]) / total[2]["v"])
        kpis = model["sheets"][0]["blocks"][2]["items"]
        self.assertEqual(kpis[0]["cell"]["x"], "=D6")
        self.assertEqual(kpis[1]["cell"]["x"], '=IFERROR((B5+B7)/B6,"")')
        adj = model["sheets"][1]["blocks"][0]
        self.assertEqual(adj["rows"][0][1]["st"], "sourced")         # 13,500 = 3 × 4,500 appears in the query result
        self.assertEqual(adj["rows"][1][1]["st"], "unverified")      # 7,250.50 appears nowhere
        self.assertEqual(adj["rows"][1][2]["v"], 8338.08)            # ROUND(7250.5 × 1.15, 2)
        self.assertEqual(model["summary"]["cells_unverified"], 1)
        checks = [x for x in f["findings"] if x["origin"] == "check"]
        self.assertTrue(any("without a source" in x["title"] for x in checks))
        self.assertTrue(any("Division by zero" in x["title"] for x in checks))
        src = next(s for s in model["sources"] if s["kind"] == "dataset")
        self.assertIn("Actuals", src["title"])
        self.assertEqual(src["link"], f"data/{src['link'].split('/')[1]}")

        # drill-down from an aggregated figure to the underlying rows
        cell = h.api("GET", f"/api/w/{ws}/files/{fid}/cell", params={"a": rows["Rent"][1]["a"]})
        self.assertEqual(cell["rows"]["total_rows"], 3)
        self.assertTrue(all(r[1] == "Rent" and r[2] == "Actual" for r in cell["rows"]["rows"]))
        self.assertEqual(cell["source"]["id"], rows["Rent"][1]["s"])

        # the Excel file
        resp = h.client.get(f"/api/w/{ws}/files/{fid}/download", params={"format": "xlsx"})
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["Variance", "Notes & adjustments", "Review", "Sources"])
        sh = wb["Variance"]
        self.assertEqual(sh["D5"].value, "=B5-C5")
        self.assertEqual(sh["B5"].font.color.rgb, "001F4E79")       # inputs in blue
        self.assertIn("Account = Rent", sh["B5"].comment.text)
        self.assertEqual(sh["F5"].value, rows["Rent"][1]["s"])      # Ref column
        self.assertEqual(sh["F5"].hyperlink.location, "'Sources'!A8")
        notes = wb["Notes & adjustments"]
        self.assertEqual(notes["B4"].fill.fgColor.rgb, "00FFF2CC")  # no source → yellow
        self.assertIn("No source", notes["B4"].comment.text)
        self.assertEqual(wb["Sources"]["A8"].value, "S1")

        if shutil.which("soffice"):  # Excel's own arithmetic agrees with the model
            calc = soffice_values(resp.content)
            for c in [*sum(va["rows"], []), *va["total"], *(k["cell"] for k in kpis), *sum(adj["rows"], [])]:
                if c.get("st") == "computed":
                    sheet, coord = c["a"].rsplit("!", 1)
                    got = calc[sheet][coord]
                    if c["v"] is None:
                        self.assertIn(got, ("", None), c["a"])
                    else:
                        self.assertAlmostEqual(got, c["v"], places=6, msg=c["a"])

        # independent review: separate model call, findings land on cells, the file is re-rendered
        h.llm.script = [[("add_finding", {"target": rows["Travel"][2]["a"], "severity": "high",
                                          "title": "Travel has no budget",
                                          "detail": "Budget is zero for travel; confirm whether it was omitted."}),
                         ("add_finding", {"target": "Nowhere!Z99", "severity": "low", "title": "Add the period"})],
                        "One material point: travel budget missing."]
        h.api("POST", f"/api/w/{ws}/files/{fid}/review", json={})
        for _ in range(100):
            f = h.api("GET", f"/api/w/{ws}/files/{fid}")
            if f["file"]["review_status"] != "running":
                break
            time.sleep(0.05)
        self.assertEqual(f["file"]["review_status"], "done", f["file"].get("review_error"))
        reviewer_call = h.llm.calls[-2]
        system = reviewer_call["messages"][0]["content"]
        self.assertIn("independent reviewer", system)
        convo = " ".join(str(m.get("content")) for m in reviewer_call["messages"])
        self.assertNotIn("Build me a Q1", convo)                    # the reviewer never sees the chat
        self.assertIn("[NO SOURCE]", convo)
        self.assertNotIn("create_workbook", [t["function"]["name"] for t in reviewer_call["tools"]])
        found = [x for x in f["findings"] if x["origin"] == "reviewer"]
        self.assertEqual({x["target"] for x in found}, {rows["Travel"][2]["a"], "file"})
        wb = load_workbook(io.BytesIO(h.client.get(f"/api/w/{ws}/files/{fid}/download").content))
        self.assertIn("Travel has no budget", wb["Variance"]["C7"].comment.text)
        self.assertIn("Travel has no budget", [c.value for c in wb["Review"]["C"]])
        # resolving a finding removes it from the file
        fnd = next(x for x in found if x["title"] == "Travel has no budget")
        h.api("PATCH", f"/api/w/{ws}/files/{fid}/findings/{fnd['id']}", json={"status": "resolved"})
        wb = load_workbook(io.BytesIO(h.client.get(f"/api/w/{ws}/files/{fid}/download").content))
        self.assertNotIn("Travel has no budget", (wb["Variance"]["C7"].comment or type("x", (), {"text": ""})).text)

        # refresh with new data: version 2 and a list of changed figures
        up = h.api("POST", f"/api/w/{ws}/imports", files=[("file", ("more.csv", (CSV + "2026-03-15,Rent,Actual,500\n")
                                                                         .encode(), "text/csv"))])
        ds = next(d for d in h.api("GET", f"/api/w/{ws}/datasets") if d["name"] == "Actuals")
        h.api("DELETE", f"/api/w/{ws}/datasets/{ds['id']}", expect=None)
        h.api("POST", f"/api/w/{ws}/imports/{up['upload_id']}/commit", json={"name": "Actuals"})
        v2 = h.api("POST", f"/api/w/{ws}/files/{fid}/refresh", json={})
        self.assertEqual(v2["version"], 2)
        changed = {c["a"]: c for c in v2["changes"]}
        self.assertEqual(changed["Variance!B5"]["after"] - changed["Variance!B5"]["before"], 500)
        self.assertEqual(len(v2["versions"]), 2)
        # other workspaces cannot see the file
        other = h.api("POST", "/api/workspaces", json={"name": "Other"})["id"]
        h.api("GET", f"/api/w/{other}/files/{fid}", expect=404)

    def test_report_with_linked_figures(self) -> None:
        h, ws = self.h, self.ws
        h.llm.script = [[("query_dataset", QUERY_ARGS)], [("create_report", REPORT)], "Report ready."]
        data = self.chat("Write the Q1 commentary")
        fid = self.file_id(data)
        f = h.api("GET", f"/api/w/{ws}/files/{fid}")
        self.assertEqual(f["formats"], ["docx", "xlsx"])
        sec = f["model"]["report"]["sections"][0]
        runs = sec["blocks"][0]["runs"]
        figs = [r for r in runs if r.get("fig")]
        self.assertEqual(figs[0]["t"], "156,000.00")
        self.assertEqual(figs[0]["fig"]["st"], "sourced")
        self.assertEqual(figs[3]["fig"]["st"], "computed")
        bullets = sec["blocks"][1]["items"]
        self.assertEqual([r["fig"]["st"] for r in bullets[0]["runs"] if r.get("fig")], ["sourced"])
        self.assertEqual([r["fig"]["st"] for r in bullets[1]["runs"] if r.get("fig")], ["unverified"])
        self.assertTrue(any(x["target"] == bullets[1]["id"] for x in f["findings"]))
        docx_bytes = h.client.get(f"/api/w/{ws}/files/{fid}/download", params={"format": "docx"}).content
        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
            body = z.read("word/document.xml").decode()
            comments = z.read("word/comments.xml").decode()
        self.assertIn('w:anchor="ref_S1"', body)
        self.assertIn('w:name="ref_S1"', body)
        self.assertIn("w:highlight", body)
        self.assertIn("No source found", comments)
        self.assertIn("Sources and calculations", body)
        self.assertIn("ofintech-desk://open/data/", self._rels(docx_bytes))

    @staticmethod
    def _rels(data: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return z.read("word/_rels/document.xml.rels").decode()

    def test_bad_definitions_come_back_as_guidance(self) -> None:
        h, ws = self.h, self.ws
        bad = {"title": "x", "sheets": [{"name": "S", "blocks": [
            {"type": "table", "id": "t", "rows": [["A", 1]], "columns": [
                {"label": "Item"}, {"label": "Amount"}, {"label": "Bad", "formula": "amount - missing"}]}]}]}
        h.llm.script = [[("create_workbook", bad)], "Sorry."]
        data = self.chat("make it")
        tool = next(m for m in data["messages"] if m["role"] == "tool")
        self.assertIn("uses 'missing'", tool["content"])
        self.assertIn("Column keys: item, amount, bad", tool["content"])
        self.assertEqual(h.api("GET", f"/api/w/{ws}/files"), [])


if __name__ == "__main__":
    unittest.main()
