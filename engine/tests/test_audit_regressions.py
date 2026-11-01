"""Regression tests for the findings of the enterprise audit (docs/audit/02-SECURITY-AUDIT.md).

One test per finding, named after it, so a reader of the audit can find the proof that each fix holds. The
tests repeat the attempt that was possible before the fix and expect it to be refused; where the fix is a
hardening with no earlier attack, they pin the new behaviour.
"""

from __future__ import annotations

import unittest
from urllib.parse import parse_qs, urlparse

from .helpers import Harness

PW = "Correct-Horse-9"


class OB01ZohoCallback(unittest.TestCase):
    """The accounts server and API domain named by the OAuth callback are accepted only when they are exactly one
    of Zoho's data centres: a look-alike host must receive neither the code nor the client secret."""

    def test_lookalike_accounts_server_gets_no_token_request(self) -> None:
        h = Harness()
        try:
            ws = h.setup_workspace(with_model=False)
            h.api("PUT", "/api/oauth-apps/zoho_books", json={"client_id": "zoho-cid", "client_secret": "zs"})
            for host in ("https://accounts.zoho.attacker.test", "https://accounts.zohocloud.attacker.test",
                         "https://accounts.zoho.com.attacker.test", "https://accounts.zoho.com@attacker.test"):
                url = h.api("POST", f"/api/w/{ws}/connections/zoho_books/authorize", json={})["url"]
                state = parse_qs(urlparse(url).query)["state"][0]
                r = h.client.get("/oauth/callback", params={"code": "c", "state": state, "accounts-server": host})
                self.assertEqual(r.status_code, 400, host)
            self.assertEqual([str(r.url) for r in h.router.requests if "attacker" in str(r.url)], [])
        finally:
            h.close()

    def test_api_domain_from_token_response_is_allow_listed(self) -> None:
        from desk_engine.connectors.flows import zoho_accounts_server, zoho_api_domain
        from desk_engine.errors import Invalid

        self.assertEqual(zoho_accounts_server(None, "eu"), "https://accounts.zoho.eu")
        self.assertEqual(zoho_accounts_server("https://accounts.zoho.sa/", None), "https://accounts.zoho.sa")
        self.assertEqual(zoho_api_domain("https://www.zohoapis.eu"), "https://www.zohoapis.eu")
        self.assertEqual(zoho_api_domain(None), "https://www.zohoapis.com")
        for bad in ("https://www.zohoapis.eu.attacker.test", "https://attacker.test", "http://www.zohoapis.eu"):
            with self.assertRaises(Invalid):
                zoho_api_domain(bad)
        with self.assertRaises(Invalid):
            zoho_accounts_server("https://accounts.zoho", None)


class OB02OriginCheck(unittest.TestCase):
    """Cookie-authenticated writes compare the whole origin, never a prefix."""

    def test_lookalike_origins_are_refused(self) -> None:
        h = Harness("server")
        try:
            c = h.client
            r = c.post("/api/auth/signup", json={"email": "owner@acme.test", "name": "Olivia", "password": PW})
            self.assertEqual(r.status_code, 200, r.text)
            body = {"name": "Acme", "base_currency": "USD"}
            refused = ["https://desk.example.com.attacker.test", "https://desk.example.community",
                       "http://desk.example.com", "https://desk.example.com:8443", "https://attacker.test", "null"]
            for origin in refused:
                r = c.post("/api/workspaces", json=body, headers={"origin": origin})
                self.assertEqual((r.status_code, r.json()["error"]), (403, "bad_origin"), origin)
                r = c.post("/api/workspaces", json=body, headers={"referer": origin + "/#/x"})
                self.assertEqual((r.status_code, r.json()["error"]), (403, "bad_origin"), origin)
            for ok in ("https://desk.example.com", "HTTPS://DESK.EXAMPLE.COM"):
                r = c.post("/api/workspaces", json=body, headers={"origin": ok})
                self.assertEqual(r.status_code, 200, (ok, r.text))
            r = c.post("/api/workspaces", json=body, headers={"referer": "https://desk.example.com/#/settings"})
            self.assertEqual(r.status_code, 200, r.text)
        finally:
            h.close()

    def test_website_forms_use_the_same_exact_match(self) -> None:
        h = Harness("server", site_url="https://ofintech.example")
        try:
            data = {"email": "owner@acme.test", "name": "Olivia", "password": PW}
            r = h.client.post("/auth/form/sign-up", data=data, headers={"origin": "https://ofintech.example.attacker.test"},
                              follow_redirects=False)
            self.assertEqual(r.status_code, 403)
            r = h.client.post("/auth/form/sign-up", data=data, headers={"origin": "https://ofintech.example"},
                              follow_redirects=False)
            self.assertEqual(r.status_code, 303, r.text)
        finally:
            h.close()


class OB03FormulaInjection(unittest.TestCase):
    """Text imported from ledgers and typed by people never becomes a formula in a file the engine writes."""

    HOSTILE = ["=HYPERLINK(\"https://attacker.test\",\"x\")", "+1+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)",
               "\t=1+1", "\r=1+1", "==2*3"]

    def _view(self, text: str) -> dict:
        line = {"account_code": "6100", "account": text, "account_class": "Expense", "debit": 100.0, "credit": 0}
        return {"month": "2026-09", "month_end": "2026-09-30",
                "items": [{"id": "it1", "kind": "accrual", "status": "approved", "note": text,
                           "data": {"title": text, "lines": [line, {**line, "debit": 0, "credit": 100.0}]}}]}

    def test_journal_csv_neutralises_every_prefix(self) -> None:
        import csv
        import io

        from desk_engine.close.pack import journal_csv

        for text in self.HOSTILE:
            rows = list(csv.reader(io.StringIO(journal_csv(self._view(text)).decode("utf-8-sig"))))
            self.assertEqual(rows[1][4], "'" + text, text)   # account
            self.assertEqual(rows[1][8], "'" + text, text)   # description
        self.assertEqual(rows[1][6], "100.0")                # numbers untouched

    def test_journal_and_pack_workbooks_store_text_as_text(self) -> None:
        import io

        from openpyxl import load_workbook

        from desk_engine.close.pack import journal_xlsx

        for text in self.HOSTILE:
            ws = load_workbook(io.BytesIO(journal_xlsx(self._view(text))))["Journal"]
            self.assertEqual(ws["E2"].data_type, "s", text)
            self.assertEqual(ws["E2"].value, text)
            self.assertEqual(ws["I2"].data_type, "s", text)
        # the engine's own totals are still formulas
        self.assertEqual(ws["G4"].data_type, "f")

    def test_dataset_export_and_cell_helper(self) -> None:
        from openpyxl import Workbook

        from desk_engine.files.cells import csv_safe, set_cell

        for text in self.HOSTILE:
            self.assertEqual(csv_safe(text), "'" + text)
            self.assertEqual(set_cell(Workbook().active["A1"], text).data_type, "s")
        self.assertEqual(csv_safe("plain"), "plain")
        self.assertEqual(csv_safe(12.5), 12.5)
        self.assertIsNone(csv_safe(None))

    def test_evidence_workbook_double_equals_is_text(self) -> None:
        from openpyxl import Workbook

        from desk_engine.files.xlsx import _set

        c = Workbook().active["A1"]
        _set(c, "==2*3")
        self.assertEqual(c.data_type, "s")


class OB15FormulaEngineRobustness(unittest.TestCase):
    """Hostile formulas get a validation error or a blank cell, never an unhandled exception."""

    def test_deep_nesting_is_refused_cleanly(self) -> None:
        from desk_engine.analysis import formula
        from desk_engine.errors import Invalid
        from desk_engine.files import expr

        deep = "(" * 300 + "1" + ")" * 300
        with self.assertRaises(Invalid):
            expr.parse(deep)
        with self.assertRaises(Invalid):
            formula.parse(deep[:399])
        expr.parse("(" * 10 + "1" + ")" * 10)   # ordinary nesting still parses

    def test_round_with_absurd_digits_gives_blank(self) -> None:
        from desk_engine.files import expr

        node = expr.parse("ROUND(1.5, 1000000)")
        self.assertIsNone(expr.evaluate(node, lambda ref: 0))
        self.assertEqual(expr.evaluate(expr.parse("ROUND(2.345, 2)"), lambda ref: 0), 2.35)


if __name__ == "__main__":
    unittest.main()
