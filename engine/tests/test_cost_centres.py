"""Cost centres: rules over the general ledger, untagged share, accrual against cash, budget run-out, budgets read
from the customer's own file, and the journal lines each accounting connector produces.

The ledger below is small enough to check by hand; the expected figures are worked out in the comments.
"""

from __future__ import annotations

import io
import math
import unittest
from datetime import date, timedelta

from desk_engine.connectors.accounting import gl_rows
from desk_engine.costcentres.analysis import Analysis
from desk_engine.costcentres.budget import match_accounts, parse_docx, parse_xlsx
from desk_engine.costcentres.rules import Matcher, norm, parse_tags, prepare

from .helpers import Harness

RIYADH = "Project: Riyadh tower"


def entry(day: str, ref: str, party: str | None, *lines: tuple[str, str, str, float, float, str | None, str | None]
          ) -> list[dict]:
    out = []
    for code, account, cls, dr, cr, tag, desc in lines:
        tags = []
        if tag:
            value, _, pct = tag.partition("@")
            dim, _, val = value.partition(": ")
            tags.append({"dimension": dim, "value": val} | ({"percent": float(pct)} if pct else {}))
        out.append({"date": day, "entry": ref, "type": "Journal", "account_code": code, "account": account,
                    "account_class": cls, "debit": dr, "credit": cr, "description": desc, "party": party,
                    "tags": tags})
    return out


BANK = ("1010", "Bank", "Bank")
AP = ("2000", "Accounts payable", "Payable")
VAT = ("1300", "VAT receivable", "Asset")
MAT = ("5200", "Materials", "Cost of sales")
SUB = ("5100", "Subcontractors", "Cost of sales")
RENT = ("6100", "Site rent", "Expense")


def L(acc: tuple[str, str, str], dr: float, cr: float, tag: str | None = None, desc: str | None = None) -> tuple:
    return (*acc, dr, cr, tag, desc)


LEDGER = [
    # bill B-1: materials 10,000 tagged to the project, paid in full on 5 June
    *entry("2026-05-10", "B-1", "Gulf Steel", L(MAT, 10000, 0, RIYADH, "Steel"), L(VAT, 1500, 0),
           L(AP, 0, 11500)),
    *entry("2026-06-05", "PAY-1", "Gulf Steel", L(AP, 11500, 0), L(BANK, 0, 11500)),
    # bill B-2: 20,000 not tagged, picked up by the keyword rule; 10,000 of 23,000 paid on 1 July
    *entry("2026-06-15", "B-2", "Gulf Steel", L(MAT, 20000, 0, None, "Rebar for Riyadh tower"), L(VAT, 3000, 0),
           L(AP, 0, 23000)),
    *entry("2026-07-01", "PAY-2", "Gulf Steel", L(AP, 10000, 0), L(BANK, 0, 10000)),
    # paid straight from the bank
    *entry("2026-07-20", "CASH-1", "Al Noor Properties", L(RENT, 6000, 0, RIYADH, "Site rent July"),
           L(BANK, 0, 6000)),
    # accrual split 50/50 between two projects: 4,500 here, not cash
    *entry("2026-08-31", "ACC-1", None, L(SUB, 9000, 0, RIYADH + "@50"), ("2100", "Accruals", "Liability", 0, 9000,
                                                                           None, None)),
    *entry("2026-08-31", "ACC-1", None, ("5100", "Subcontractors", "Cost of sales", 0, 0, "Project: Jeddah mall@50",
                                        None)),
    # another project's bill
    *entry("2026-08-10", "B-3", "Desert Cranes", L(SUB, 4000, 0, "Project: Jeddah mall"), L(AP, 0, 4000)),
    # money out to the project's supplier with no bill and no cost centre: may belong here
    *entry("2026-09-05", "BANK-7", "Gulf Steel", ("9999", "Suspense", "Liability", 2500, 0, None, "Gulf Steel advance"),
           L(BANK, 0, 2500)),
    # revenue
    *entry("2026-09-10", "INV-1", "Client A", ("1200", "Receivables", "Receivable", 50000, 0, None, None),
           ("4000", "Contract revenue", "Income", 0, 50000, RIYADH, None)),
    # the project's supplier, not tagged and not matched: possibly missing
    *entry("2026-09-20", "B-4", "Gulf Steel", L(MAT, 1000, 0, None, "Office shelves"), L(AP, 0, 1000)),
]
RULES = {"include": [{"tag": RIYADH}, {"keywords": ["riyadh tower"], "accounts": ["5"]}]}
CC = {"name": "Riyadh tower", "rules": RULES, "start_date": "2026-05-01", "end_date": "2026-12-31", "burn_months": 3}


def ledger_lines() -> list[dict]:
    rows = gl_rows(LEDGER)
    keys = ["date", "entry", "type", "account_code", "account", "account_class", "debit", "credit", "description",
            "party", "tags"]
    return [prepare(dict(zip(keys, r, strict=True)) | {"row": i + 1}) for i, r in enumerate(rows)]


def budget_xlsx(total: float | str = 110000) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Riyadh tower — budget 2026"
    ws.append([])
    ws.append(["Code", "Item", "Budget (SAR)"])
    ws.append([5200, "Materials", 60000])
    ws.append([5100, "Subcontractors", "30,000"])
    ws.append([None, "Site rent", 12000])
    ws.append([None, "Contingency", 8000])
    ws.append([None, "Total", total])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Rules(unittest.TestCase):
    def test_normalising_and_tags(self) -> None:
        self.assertEqual(norm("  مُؤسّسة  الإنشاءات "), norm("مؤسسة الانشاءات"))
        self.assertEqual(norm("مكتبة"), norm("مكتبه"))
        self.assertEqual(parse_tags("Project: Riyadh | Department: Site (60%)"),
                         [("project", "riyadh", 1.0), ("department", "site", 0.6)])

    def test_conditions(self) -> None:
        m = Matcher({"include": [{"tag": "Riyadh tower"}, {"accounts": ["61"], "parties": ["al noor properties"]}],
                     "exclude": [{"keywords": ["reversal"]}]})
        line = prepare({"date": "2026-01-01", "account_code": "6100", "account": "Site rent", "party": "Al Noor Properties",
                        "description": "July", "tags": ""})
        self.assertEqual(m.match(line), (1.0, 1, False))
        self.assertEqual(m.match(prepare(line | {"description": "Reversal of July"}))[0], 0.0)
        tagged = prepare(line | {"tags": "Project: Riyadh tower (25%)", "party": None})
        self.assertEqual(m.match(tagged), (0.25, 0, True))
        with self.assertRaises(Exception):
            Matcher({"include": [{"date_from": "2026-01-01"}]})  # a date alone is not a rule


class Analyse(unittest.TestCase):
    def analysis(self, **cc: object) -> Analysis:
        return Analysis(ledger_lines(), CC | cc, as_of=date(2026, 9, 20))

    def test_rules_and_untagged_share(self) -> None:
        t = self.analysis().totals()
        self.assertEqual(t["cost"], 40500.0)          # 10,000 + 20,000 + 6,000 + 9,000 × 50%
        self.assertEqual(t["revenue"], 50000.0)
        self.assertEqual(t["margin"], 9500.0)
        self.assertEqual(t["tagged_cost"], 20500.0)
        self.assertEqual(t["rule_cost"], 20000.0)
        self.assertEqual(t["untagged_share"], round(20000 / 40500 * 100, 1))
        self.assertEqual([r["cost"] for r in t["by_rule"]], [20500.0, 20000.0])

    def test_accrual_against_cash(self) -> None:
        c = self.analysis().cash()
        months = c.pop("_months")
        paid_b2 = round(20000 * 10000 / 23000, 2)  # FIFO: B-1 settled first, then 10,000 of B-2's 23,000
        self.assertEqual(c["accrual_cost"], 40500.0)
        self.assertEqual(c["paid_direct"], 6000.0)
        self.assertEqual(c["through_payables"], 30000.0)
        self.assertEqual(c["paid_through_payables"], round(10000 + paid_b2, 2))
        self.assertEqual(c["unpaid"], round(20000 - paid_b2, 2))
        self.assertEqual(c["not_cash"], 4500.0)
        self.assertEqual(c["cash_paid"], round(6000 + 10000 + paid_b2, 2))
        self.assertEqual(c["gap"], round(40500 - 16000 - paid_b2, 2))
        self.assertAlmostEqual(c["paid_direct"] + c["through_payables"] + c["not_cash"], c["accrual_cost"])
        self.assertAlmostEqual(months["2026-06"]["cash"], 10000)
        self.assertAlmostEqual(months["2026-07"]["cash"], 6000 + 20000 * 10000 / 23000)
        self.assertEqual(c["supplier_missing"], 0.0)

    def test_candidates_and_overlaps(self) -> None:
        a = Analysis(ledger_lines(), CC, as_of=date(2026, 9, 20),
                     others=[("Jeddah mall", Matcher({"include": [{"tag": "Project: Jeddah mall"}]}), False)])
        missing = a.possibly_missing()
        self.assertEqual((missing["count"], missing["amount"]), (1, 1000.0))
        self.assertEqual(missing["lines"][0]["entry"], "B-4")
        bank = a.unassigned_bank()
        self.assertEqual((bank["count"], bank["amount"]), (1, 2500.0))
        self.assertEqual(bank["lines"][0]["entry"], "BANK-7")
        self.assertEqual(a.overlaps()["count"], 0)  # a 50/50 split is not counted twice
        greedy = Analysis(ledger_lines(), CC, as_of=date(2026, 9, 20),
                          others=[("All steel", Matcher({"include": [{"parties": ["Gulf Steel"]}]}), False)])
        self.assertEqual(greedy.overlaps()["amount"], 30000.0)  # B-1 and B-2 taken by both

    def test_run_out(self) -> None:
        # window 21 June – 20 September (92 days): 6,000 + 4,500 = 10,500, so 114.13 a day
        per_day = 10500 / 92
        ro = self.analysis(budget_amount=45000).runout(45000, self.analysis().daily_costs())
        self.assertEqual(ro["window_days"], 92)
        self.assertEqual(ro["spent_in_window"], 10500.0)
        expected = date(2026, 9, 20) + timedelta(days=math.ceil(4500 / per_day))
        self.assertEqual(ro["run_out_date"], expected.isoformat())
        self.assertEqual(ro["status"], "before_end")
        self.assertEqual(ro["days_before_end"], (date(2026, 12, 31) - expected).days)
        self.assertEqual(ro["projected_at_end"], round(40500 + per_day * 102, 2))
        over = self.analysis().runout(30000, self.analysis().daily_costs())
        self.assertEqual((over["status"], over["exceeded_on"]), ("exceeded", "2026-06-15"))
        self.assertEqual(self.analysis().runout(None, {})["status"], "no_budget")


class BudgetFiles(unittest.TestCase):
    def test_excel_budget(self) -> None:
        r = parse_xlsx(budget_xlsx())
        self.assertEqual([(x["label"], x["amount"], x["code"], x["cell"]) for x in r["lines"]], [
            ("Materials", 60000.0, "5200", "Budget!C4"), ("Subcontractors", 30000.0, "5100", "Budget!C5"),
            ("Site rent", 12000.0, None, "Budget!C6"), ("Contingency", 8000.0, None, "Budget!C7")])
        self.assertEqual(r["total"], 110000.0)
        self.assertIn("Budget!C8", r["check"])
        self.assertEqual(r["columns"], {"amount": "C", "label": "B", "code": "A"})
        off = parse_xlsx(budget_xlsx(100000))
        self.assertIn("110,000.00", off["warnings"][0])
        mapped = match_accounts(r["lines"], [("5200", "Materials"), ("5100", "Subcontractors"), ("6100", "Site rent")])
        self.assertEqual([x["accounts"] for x in mapped], [["5200"], ["5100"], ["6100"], []])

    def test_title_line_that_mentions_the_budget(self) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Budget"
        ws["A1"] = "Eastport warehouse fit-out - approved budget"   # a title, not the header
        ws.append([])
        ws.append(["Code", "Item", "Budget (USD)", "Notes"])
        ws.append([6400, "Contractors", 58000, None])
        ws.append([1600, "Pallet racking", 86000, None])
        ws.append([None, "Total", 144000, None])
        buf = io.BytesIO()
        wb.save(buf)
        r = parse_xlsx(buf.getvalue())
        self.assertEqual([(x["label"], x["amount"], x["code"], x["cell"]) for x in r["lines"]],
                         [("Contractors", 58000.0, "6400", "Budget!C4"), ("Pallet racking", 86000.0, "1600", "Budget!C5")])
        self.assertIn("Budget!C6", r["check"])

    def test_arabic_word_budget(self) -> None:
        from docx import Document

        d = Document()
        d.add_paragraph("موازنة المشروع")
        t = d.add_table(rows=4, cols=2)
        for i, (a, b) in enumerate([("البند", "الموازنة"), ("المواد", "٦٠٬٠٠٠"), ("الإيجار", "12,000.50"),
                                    ("الإجمالي", "72,000.50")]):
            t.cell(i, 0).text, t.cell(i, 1).text = a, b
        buf = io.BytesIO()
        d.save(buf)
        r = parse_docx(buf.getvalue())
        self.assertEqual([(x["label"], x["amount"]) for x in r["lines"]], [("المواد", 60000.0), ("الإيجار", 12000.5)])
        self.assertIn("72,000.50", r["check"])
        self.assertEqual(match_accounts(r["lines"], [("5200", "المواد"), (None, "إيجار الموقع")])[0]["accounts"], ["5200"])


class Api(unittest.TestCase):
    def setUp(self) -> None:
        self.h = Harness()
        self.ws = self.h.setup_workspace()

    def tearDown(self) -> None:
        self.h.close()

    def test_end_to_end(self) -> None:
        h, ws = self.h, self.ws
        conn = h.api("POST", f"/api/w/{ws}/connections/custom", json={"label": "ERP", "mode": "push"})
        r = h.client.post("/api/ingest/v1/accounting/journal-lines", json={"lines": LEDGER},
                          headers={"authorization": f"Bearer {conn['ingest_key']}"})
        self.assertEqual(r.status_code, 200, r.text)
        listing = h.api("GET", f"/api/w/{ws}/cost-centres")
        self.assertEqual([x["name"] for x in listing["ledgers"]], ["ERP · General ledger"])
        did = listing["ledgers"][0]["id"]
        values = h.api("GET", f"/api/w/{ws}/cost-centres/ledgers/{did}")
        self.assertIn("Project: Riyadh tower", values["tags"])
        self.assertIn("Gulf Steel", values["parties"])
        pv = h.api("POST", f"/api/w/{ws}/cost-centres/preview", json={"dataset_id": did, "rules": RULES,
                                                                      "start_date": "2026-05-01"})
        self.assertEqual(pv["cost"], 40500.0)
        h.api("POST", f"/api/w/{ws}/cost-centres", expect=422, json={"name": "x", "dataset_id": did,
                                                                    "rules": {"include": []}})
        cc = h.api("POST", f"/api/w/{ws}/cost-centres", json=CC | {"dataset_id": did, "code": "PRJ-7"})
        h.api("POST", f"/api/w/{ws}/cost-centres", json={"name": "Jeddah mall", "dataset_id": did,
                                                         "rules": {"include": [{"tag": "Project: Jeddah mall"}]}})
        up = h.client.post(f"/api/w/{ws}/files/upload", files={"file": ("Riyadh budget.xlsx", budget_xlsx())})
        fid = up.json()["file"]["id"]
        imp = h.api("POST", f"/api/w/{ws}/cost-centres/{cc['id']}/budget/import", json={"file_id": fid})
        self.assertEqual((len(imp["lines"]), imp["mapped"]), (4, 3))
        self.assertEqual(imp["lines"][2]["source"]["cell"], "Budget!C6")
        rep = h.api("GET", f"/api/w/{ws}/cost-centres/{cc['id']}/report")
        self.assertEqual(rep["as_of"], "2026-09-20")  # the ledger's last date
        self.assertEqual(rep["totals"]["budget"], 110000.0)
        self.assertEqual(rep["totals"]["cost"], 40500.0)
        self.assertEqual([(b["label"], b["actual"]) for b in rep["budget"]["lines"]],
                         [("Materials", 30000.0), ("Subcontractors", 4500.0), ("Site rent", 6000.0),
                          ("Contingency", None)])
        self.assertEqual(rep["budget"]["unbudgeted_cost"], 0.0)
        self.assertEqual(rep["runout"]["status"], "after_end")
        self.assertEqual(rep["overlaps"]["count"], 0)
        lines = h.api("GET", f"/api/w/{ws}/cost-centres/PRJ-7/lines", params={"kind": "matched"})
        self.assertEqual(lines["total"], 5)
        self.assertEqual({x["basis"] for x in lines["lines"]}, {"tag", "rule"})
        # editing the rules keeps the history
        new = h.api("PATCH", f"/api/w/{ws}/cost-centres/{cc['id']}", json={
            "rules": {"include": [{"tag": RIYADH}]}, "note": "tags only"})
        self.assertEqual([v["version"] for v in new["versions"]], [2, 1])
        self.assertEqual(h.api("GET", f"/api/w/{ws}/cost-centres/{cc['id']}/report")["totals"]["cost"], 20500.0)
        # correcting a budget line keeps where the old figure came from
        bl = new["budget_lines"]
        bl[3] = bl[3] | {"amount": 5000, "accounts": []}
        got = h.api("PUT", f"/api/w/{ws}/cost-centres/{cc['id']}/budget/lines", json={"lines": [
            {k: x[k] for k in ("id", "label", "amount", "accounts")} for x in bl]})
        self.assertEqual(got["budget_lines"][3]["source"]["was"]["cell"], "Budget!C7")
        self.assertEqual(got["budget_lines"][0]["source"]["cell"], "Budget!C4")
        # the assistant reads the same report
        h.llm.script = [[("cost_centre_report", {"cost_centre": "Riyadh tower"})], "Done."]
        t = h.api("POST", f"/api/w/{ws}/threads", json={})
        h.api("POST", f"/api/w/{ws}/threads/{t['id']}/messages", json={"text": "How is the Riyadh tower project?"})
        tool = next(m for m in h.wait_run(t["id"])["messages"] if m["role"] == "tool")
        self.assertIn("Cost 20,500.00", tool["content"])
        self.assertIn("Budget 107,000.00", tool["content"])
        self.assertIn("possibly missing from this cost centre: 2 for 21,000.00", tool["content"])
        sent = [x["function"]["name"] for x in h.llm.calls[0]["tools"]]
        self.assertIn("cost_centre_report", sent)


def ledger_xlsx() -> bytes:
    from openpyxl import Workbook

    from desk_engine.connectors.accounting import GL_COLUMNS

    wb = Workbook()
    ws = wb.active
    ws.title = "GL"
    ws.append([c["name"] for c in GL_COLUMNS])
    for r in gl_rows(LEDGER):
        r[3] = int(r[3])  # codes typed as numbers, as people do
        ws.append([date.fromisoformat(r[0]), *r[1:]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportedLedger(unittest.TestCase):
    def test_excel_general_ledger(self) -> None:
        h = Harness()
        try:
            ws = h.setup_workspace()
            up = h.api("POST", f"/api/w/{ws}/imports", files={"file": ("GL export.xlsx", ledger_xlsx())})
            ds = h.api("POST", f"/api/w/{ws}/imports/{up['upload_id']}/commit", json={"name": "GL 2026"})
            self.assertEqual([x["id"] for x in h.api("GET", f"/api/w/{ws}/cost-centres")["ledgers"]], [ds["id"]])
            pv = h.api("POST", f"/api/w/{ws}/cost-centres/preview", json={"dataset_id": ds["id"], "rules": RULES})
            self.assertEqual((pv["cost"], pv["matched_lines"]), (40500.0, 5))
        finally:
            h.close()


class ConnectorLines(unittest.TestCase):
    def test_xero_journals(self) -> None:
        from desk_engine.connectors.xero import parse_journals

        lines = parse_journals([{"JournalNumber": 41, "JournalDate": "/Date(1781049600000+0000)/", "SourceType": "ACCPAY",
                                 "JournalLines": [
                                     {"AccountID": "a1", "AccountCode": "5200", "AccountType": "DIRECTCOSTS",
                                      "AccountName": "Materials", "NetAmount": 1000, "Description": "Steel",
                                      "TrackingCategories": [{"Name": "Project", "Option": "Riyadh tower"}]},
                                     {"AccountID": "ap", "AccountCode": "800", "AccountType": "CURRLIAB",
                                      "AccountName": "Accounts Payable", "NetAmount": -1000}]}], {"ap": "Payable"})
        self.assertEqual([(x["date"], x["entry"], x["account_class"], x["debit"], x["credit"]) for x in lines],
                         [("2026-06-10", "J41", "Cost of sales", 1000.0, 0.0), ("2026-06-10", "J41", "Payable", 0.0, 1000.0)])
        self.assertEqual(gl_rows(lines)[0][10], RIYADH)

    def test_quickbooks_general_ledger(self) -> None:
        from desk_engine.connectors.quickbooks import parse_general_ledger

        keys = ["tx_date", "txn_type", "doc_num", "name", "memo", "account_name", "klass_name", "dept_name", "debt_amt",
                "credit_amt", "subt_nat_amount"]
        report = {"Columns": {"Column": [{"ColTitle": k, "MetaData": [{"Name": "ColKey", "Value": k}]} for k in keys]},
                  "Rows": {"Row": [{"Header": {"ColData": [{"value": "Materials", "id": "57"}]}, "type": "Section",
                                    "Rows": {"Row": [
                                        {"type": "Data", "ColData": [{"value": "Beginning Balance"}]},
                                        {"type": "Data", "ColData": [
                                            {"value": "2026-06-15"}, {"value": "Bill", "id": "145"}, {"value": "B-2"},
                                            {"value": "Gulf Steel", "id": "9"}, {"value": "Rebar"},
                                            {"value": "Materials", "id": "57"}, {"value": "Riyadh tower"}, {"value": ""},
                                            {"value": "20000.00"}, {"value": ""}, {"value": "20000.00"}]}]}}]}}
        lines = parse_general_ledger(report, {"57": {"code": "5200", "class": "Cost of sales"}})
        self.assertEqual(len(lines), 1)
        ln = lines[0]
        self.assertEqual((ln["entry"], ln["party"], ln["account_code"], ln["debit"], ln["credit"]),
                         ("Bill:145", "Gulf Steel", "5200", 20000.0, 0.0))
        self.assertEqual(ln["tags"], [{"dimension": "Class", "value": "Riyadh tower"}])

    def test_odoo_analytic_split(self) -> None:
        from desk_engine.connectors.odoo import odoo_line

        ln = odoo_line({"date": "2026-08-31", "move_id": [7, "MISC/2026/0007"], "move_name": "MISC/2026/0007",
                        "journal_id": [3, "Miscellaneous"], "account_id": [12, "510000 Subcontractors"],
                        "name": "Crane hire", "debit": 9000, "credit": 0, "partner_id": False,
                        "analytic_distribution": {"4": 60, "5": 40}},
                       {12: {"code": "510000", "name": "Subcontractors", "account_type": "expense_direct_cost"}},
                       {"4": ("Projects", "Riyadh tower"), "5": ("Projects", "Jeddah mall")})
        self.assertEqual(ln["account_class"], "Cost of sales")
        self.assertEqual(gl_rows([ln])[0][10], "Projects: Riyadh tower (60%) | Projects: Jeddah mall (40%)")
        m = Matcher({"include": [{"tag": "Riyadh tower"}]})
        line = prepare({"date": "2026-08-31", "account_code": "510000", "account": "Subcontractors", "party": None,
                        "description": "", "tags": gl_rows([ln])[0][10]})
        self.assertEqual(m.match(line)[0], 0.6)


if __name__ == "__main__":
    unittest.main()
