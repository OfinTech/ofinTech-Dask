"""The financial golden scenario: fixed inputs through calculators, dataset queries, the formula engine and the
close journal, with every output recorded in tests/fixtures/financial_golden.json.

The recorded file was produced by the pre-audit engine (commit e8de370) and is byte-identical to the current
engine's output. Any difference the test reports is a change in financial behaviour and must be explained in
docs/audit/08-REMEDIATION-ROADMAP.md before the file is regenerated with:

    cd engine && python -m tests.golden_financial --update
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import random
import sys
from pathlib import Path
from typing import Any

GOLDEN = Path(__file__).parent / "fixtures" / "financial_golden.json"


def scenario() -> dict[str, Any]:
    from .helpers import Harness
    from desk_engine.agent.tools import ToolContext, registry
    from desk_engine.close import pack
    from desk_engine.files import expr

    out: dict[str, Any] = {}
    h = Harness()
    ws = h.setup_workspace(with_model=False)
    ctx = ToolContext(h.engine, h.engine.db.one("SELECT * FROM workspaces WHERE id = ?", (ws,)), "u")

    def run(name, args):
        o = asyncio.run(registry.get_static(name).run(ctx, args))
        return {"text": o.text, "data": json.loads(json.dumps(o.data, default=str))}

    out["calc"] = {
     "npv_irr": run("calc_npv_irr", {"cashflows": [-100000, 25000, 30000, 35000, 40000], "discount_rate": 0.08}),
     "xnpv": run("calc_xnpv_xirr", {"cashflows": [-5000, 1500, 2000, 2500], "dates": ["2026-01-15", "2026-07-01", "2027-01-01", "2027-06-30"], "discount_rate": 0.1}),
     "loan": run("calc_loan_schedule", {"principal": 800000, "annual_rate": 0.065, "periods": 48}),
     "dep_sl": run("calc_depreciation", {"cost": 593000, "salvage_value": 30000, "useful_life_years": 5, "method": "straight_line"}),
     "dep_ddb": run("calc_depreciation", {"cost": 120000, "salvage_value": 10000, "useful_life_years": 4, "method": "double_declining"}),
     "dep_syd": run("calc_depreciation", {"cost": 120000, "salvage_value": 10000, "useful_life_years": 4, "method": "sum_of_years"}),
     "ratios": run("calc_financial_ratios", {"revenue": 5920200, "cogs": 830620, "operating_income": 1400000, "net_income": 1336001, "total_assets": 7000000, "total_equity": 5500000, "current_assets": 6500000, "current_liabilities": 1200000, "inventory": 0, "receivables": 897250, "payables": 186593, "cash": 5625524, "total_debt": 475352, "interest_expense": 45000}),
     "variance": run("calc_variance", {"lines": [{"name": "Revenue", "actual": 5920200, "budget": 5500000, "higher_is_better": True}, {"name": "Rent", "actual": 120000.5, "budget": 110000, "higher_is_better": False}, {"name": "Zero", "actual": 0, "budget": 0}]}),
     "tax_incl": run("calc_sales_tax", {"amount": 1150, "rate": 0.15, "amount_includes_tax": True}),
     "tax_excl": run("calc_sales_tax", {"amount": 1000.33, "rate": 0.15, "amount_includes_tax": False}),
     "breakeven": run("calc_breakeven", {"fixed_costs": 250000, "unit_price": 45, "unit_variable_cost": 27.5}),
     "growth": run("calc_growth", {"values": [881000, 3735000, 5920200], "years": 2}),
     "fx": run("calc_currency_convert", {"amounts": [1000, 2500.75, -300], "rate": 3.75, "from_currency": "USD", "to_currency": "SAR", "rate_date": "2026-10-31"}),
    }

    # dataset import + queries
    rnd = random.Random(7)
    rows = [["date", "account", "party", "description", "debit", "credit"]]
    for i in range(5000):
        rows.append([f"2026-{1 + i % 12:02d}-{1 + i % 28:02d}", str(4000 + i % 30), f"Party {i % 40}", f"Line {i}", f"{rnd.randint(1, 9000)}.{i % 100:02d}", "0" if i % 3 else f"{rnd.randint(1, 500)}.50"])
    buf = io.StringIO(); csv.writer(buf).writerows(rows)
    up = h.api("POST", f"/api/w/{ws}/imports", files={"file": ("gl.csv", buf.getvalue().encode())})
    ds = h.api("POST", f"/api/w/{ws}/imports/{up['upload_id']}/commit", json={"name": "GL", "sheet": None, "header_row": 0})
    out["dataset"] = {"row_count": ds["row_count"], "columns": [c["name"] for c in ds["columns"]]}
    qs = {
     "total": {"dataset": ds["id"], "measures": [{"column": "debit", "agg": "sum"}, {"column": "credit", "agg": "sum"}, {"agg": "count"}]},
     "by_month": {"dataset": ds["id"], "measures": [{"column": "debit", "agg": "sum"}], "group_by": [{"column": "date", "grain": "month"}], "sort": [{"by": "date", "dir": "asc"}], "limit": 100},
     "by_fy": {"dataset": ds["id"], "measures": [{"column": "debit", "agg": "sum"}], "group_by": [{"column": "date", "grain": "fiscal_year"}]},
     "by_account_filtered": {"dataset": ds["id"], "measures": [{"column": "debit", "agg": "sum"}, {"column": "debit", "agg": "avg"}, {"column": "debit", "agg": "max"}], "group_by": [{"column": "account"}], "filters": [{"column": "party", "op": "eq", "value": "Party 7"}], "sort": [{"by": "account", "dir": "asc"}], "limit": 100},
     "contains": {"dataset": ds["id"], "measures": [{"agg": "count"}], "filters": [{"column": "description", "op": "contains", "value": "Line 12"}]},
    }
    out["queries"] = {k: h.api("POST", f"/api/w/{ws}/query", json=v) for k, v in qs.items()}
    exp = h.client.get(f"/api/w/{ws}/datasets/{ds['id']}/export.csv").text
    out["export_csv_head"] = exp.splitlines()[:5]

    # formula engine
    out["expr"] = {}
    for f in ["ROUND(7250.5 * 1.15, 2)", "1/3", "ROUND(2.675, 2)", "ROUND(-2.5, 0)", "SUM(1, 2.5, -0.5) / 3", "MAX(1, 2) - MIN(3, 4)", "ABS(-7.25) * 100", "1 / 0", "ROUND(1234567.891, -3)", "(1 + 2) * (3 + 4)"]:
        out["expr"][f] = expr.evaluate(expr.parse(f), lambda ref: 0)

    # close journal and pack with benign ledger text
    view = {"month": "2026-09", "month_end": "2026-09-30", "items": [
      {"id": "it1", "kind": "accrual", "status": "approved", "note": "Audit fee accrual", "data": {"title": "Audit fee Q3", "lines": [
         {"account_code": "6300", "account": "Audit fees", "account_class": "Expense", "debit": 12500.0, "credit": 0},
         {"account_code": "2100", "account": "Accrued expenses", "account_class": "Liability", "debit": 0, "credit": 12500.0}]}},
      {"id": "it2", "kind": "bank_fee", "status": "approved", "note": None, "data": {"title": "Bank charges", "date": "2026-09-14", "lines": [
         {"account_code": "6900", "account": "Bank charges", "account_class": "Expense", "debit": 85.5, "credit": 0},
         {"account_code": "1010", "account": "Bank – SNB", "account_class": "Asset", "debit": 0, "credit": 85.5}]}},
      {"id": "it3", "kind": "accrual", "status": "rejected", "note": None, "data": {"title": "Not approved", "lines": [{"account_code": "1", "account": "x", "account_class": "", "debit": 1, "credit": 0}]}}]}
    out["journal_rows"] = pack.journal_rows(view)
    out["journal_csv"] = pack.journal_csv(view).decode("utf-8-sig").splitlines()
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(pack.journal_xlsx(view)))
    out["journal_xlsx"] = [[(c.value, c.data_type) for c in r] for r in wb["Journal"].iter_rows()]
    h.close()
    return json.loads(json.dumps(out, sort_keys=True, default=str))


def main(argv: list[str]) -> None:
    data = scenario()
    if "--update" in argv:
        GOLDEN.parent.mkdir(exist_ok=True)
        GOLDEN.write_text(json.dumps(data, sort_keys=True, indent=1) + "\n")
        print(f"wrote {GOLDEN}")
    else:
        print(json.dumps(data, sort_keys=True, indent=1))


if __name__ == "__main__":
    main(sys.argv[1:])
