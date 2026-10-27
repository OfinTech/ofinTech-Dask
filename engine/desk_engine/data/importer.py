"""Excel and CSV import: header detection, type inference (numbers in any common notation,
dates in day-first or month-first order), and an optional wide-to-long unpivot for reports that
put periods in columns."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from ..errors import Invalid

MAX_ROWS = 200_000
MAX_COLS = 200

_NUM_CLEAN = re.compile(r"[\s '’]")
_CURRENCY = re.compile(r"^[A-Z]{3}\s*|\s*[A-Z]{3}$|[$€£¥₹₩₺₪₦₱฿₫₴₸₽﷼]|R\$|CHF|kr|zł|Rp|RM|S\$|A\$|C\$")
_EU_NUMBER = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$|^-?\d+,\d+$")
_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
_MONTH_HEADER = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*[\s\-/']*(\d{2,4})?$", re.I)


def parse_number(value: Any, decimal_comma: bool = False) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s or s in {"-", "–", "—"}:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    if s.endswith("-") and s[:-1].replace(",", "").replace(".", "").isdigit():
        negative, s = True, s[:-1]  # trailing minus used by some ledgers
    s = _CURRENCY.sub("", s).strip()
    percent = s.endswith("%")
    if percent:
        s = s[:-1]
    s = _NUM_CLEAN.sub("", s)
    if s.startswith("-"):
        negative, s = not negative, s[1:]
    if decimal_comma:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        number = float(s)
    except ValueError:
        return None
    if percent:
        number /= 100
    return -number if negative else number


def parse_date(value: Any, order: str = "auto") -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?$", s)
    if m:
        return _safe_date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})$", s)
    if m:
        a, b, y = int(m[1]), int(m[2]), int(m[3])
        y = y + 2000 if y < 100 else y
        if order == "MDY" or (order == "auto" and b > 12 >= a):
            return _safe_date(y, a, b)
        return _safe_date(y, b, a)
    m = re.match(r"^(\d{1,2})[\s\-]([A-Za-z]{3,9})[\s\-,]+(\d{2,4})$", s)
    if m and m[2][:3].lower() in _MONTHS:
        y = int(m[3])
        return _safe_date(y + 2000 if y < 100 else y, _MONTHS[m[2][:3].lower()], int(m[1]))
    m = re.match(r"^([A-Za-z]{3,9})\s+(\d{1,2}),?\s+(\d{4})$", s)
    if m and m[1][:3].lower() in _MONTHS:
        return _safe_date(int(m[3]), _MONTHS[m[1][:3].lower()], int(m[2]))
    m = _MONTH_HEADER.match(s)
    if m and m[2]:
        y = int(m[2])
        return _safe_date(y + 2000 if y < 100 else y, _MONTHS[m[1][:3].lower()], 1)
    return None


def _safe_date(y: int, m: int, d: int) -> date | None:
    try:
        return date(y, m, d)
    except ValueError:
        return None


@dataclass
class Column:
    key: str
    name: str
    type: str  # number | date | text

    def to_dict(self) -> dict[str, str]:
        return {"key": self.key, "name": self.name, "type": self.type}


@dataclass
class Table:
    name: str
    raw: list[list[Any]]
    header_row: int = 0
    columns: list[Column] = field(default_factory=list)
    rows: list[list[Any]] = field(default_factory=list)


# --- reading -----------------------------------------------------------------------------------


def read_workbook(data: bytes, filename: str) -> list[Table]:
    lower = filename.lower()
    if lower.endswith((".csv", ".txt", ".tsv")):
        return [Table(name=filename.rsplit(".", 1)[0], raw=_read_csv(data))]
    if lower.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    if lower.endswith(".xls"):
        raise Invalid("Old .xls files are not supported. Open the file in Excel and save it as .xlsx.",
                      code="unsupported_file")
    raise Invalid("Upload an .xlsx, .xlsm or .csv file.", code="unsupported_file")


def _read_csv(data: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i > MAX_ROWS:
            break
        rows.append(row[:MAX_COLS])
    return rows


def _read_xlsx(data: bytes) -> list[Table]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise Invalid("This file could not be read as an Excel workbook.", code="unreadable_file") from exc
    tables = []
    for ws in wb.worksheets:
        raw = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS:
                break
            raw.append(list(row[:MAX_COLS]))
        while raw and all(v in (None, "") for v in raw[-1]):
            raw.pop()
        if raw:
            tables.append(Table(name=ws.title, raw=raw))
    wb.close()
    if not tables:
        raise Invalid("The workbook has no data.", code="empty_file")
    return tables


# --- shaping -----------------------------------------------------------------------------------


def detect_header_row(raw: list[list[Any]]) -> int:
    best, best_score = 0, -1.0
    for i, row in enumerate(raw[:25]):
        cells = [c for c in row if c not in (None, "")]
        if len(cells) < 2:
            continue
        texts = sum(1 for c in cells if isinstance(c, str) and parse_number(c) is None)
        score = texts / max(len(row), 1) + len(cells) / 100
        if texts >= max(2, len(cells) * 0.6) and score > best_score:
            best, best_score = i, score
    return best


def _unique_names(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, n in enumerate(names):
        base = (n or f"Column {i + 1}").strip()[:80] or f"Column {i + 1}"
        key = base.lower()
        if key in seen:
            seen[key] += 1
            base = f"{base} ({seen[key]})"
        else:
            seen[key] = 1
        out.append(base)
    return out


def infer_type(values: list[Any], decimal_comma: bool, date_order: str) -> str:
    present = [v for v in values if v not in (None, "")]
    if not present:
        return "text"
    sample = present[:2000]
    dates = sum(1 for v in sample if isinstance(v, (date, datetime)) or
                (isinstance(v, str) and parse_date(v, date_order) is not None))
    if dates >= 0.9 * len(sample):
        return "date"
    nums = sum(1 for v in sample if parse_number(v, decimal_comma) is not None)
    if nums >= 0.9 * len(sample):
        return "number"
    return "text"


def guess_decimal_comma(raw: list[list[Any]]) -> bool:
    eu = us = 0
    for row in raw[:500]:
        for v in row:
            if isinstance(v, str):
                s = v.strip()
                if _EU_NUMBER.match(s):
                    eu += 1
                elif re.match(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$|^-?\d+\.\d+$", s):
                    us += 1
    return eu > us


def guess_date_order(raw: list[list[Any]]) -> str:
    dmy = mdy = 0
    for row in raw[:500]:
        for v in row:
            if isinstance(v, str):
                m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-]\d{2,4}$", v.strip())
                if m:
                    a, b = int(m[1]), int(m[2])
                    if a > 12 >= b:
                        dmy += 1
                    elif b > 12 >= a:
                        mdy += 1
    return "MDY" if mdy > dmy else "DMY"


def shape(table: Table, *, header_row: int | None = None, decimal_comma: bool | None = None,
          date_order: str | None = None) -> Table:
    raw = table.raw
    hr = detect_header_row(raw) if header_row is None else header_row
    if hr >= len(raw):
        raise Invalid("Header row is beyond the end of the sheet.")
    dc = guess_decimal_comma(raw) if decimal_comma is None else decimal_comma
    order = guess_date_order(raw) if date_order in (None, "auto") else date_order
    header = raw[hr]
    width = max((len(r) for r in raw[hr:hr + 1000]), default=0)
    names = _unique_names([_header_text(header[i]) if i < len(header) else "" for i in range(width)])
    body = [list(r) + [None] * (width - len(r)) for r in raw[hr + 1:]]
    body = [r for r in body if any(v not in (None, "") for v in r)]
    # drop columns that are entirely empty and unnamed
    keep = [i for i in range(width) if names[i] and (not names[i].startswith("Column ") or
                                                      any(r[i] not in (None, "") for r in body))]
    columns = []
    for out_i, src_i in enumerate(keep):
        values = [r[src_i] for r in body]
        columns.append(Column(f"c{out_i}", names[src_i], infer_type(values, dc, order)))
    rows = []
    for r in body:
        rows.append([convert(r[src_i], col.type, dc, order) for src_i, col in zip(keep, columns, strict=True)])
    return Table(table.name, raw, hr, columns, rows)


def _header_text(v: Any) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v).strip()


def convert(v: Any, typ: str, decimal_comma: bool, date_order: str) -> Any:
    if v in (None, ""):
        return None
    if typ == "number":
        return parse_number(v, decimal_comma)
    if typ == "date":
        parsed = parse_date(v, date_order)
        return parsed.isoformat() if parsed else None
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()[:2000]


def period_columns(table: Table) -> list[str]:
    """Column names that look like periods (Jan 2026, 2026-01-31, FY2025, Q1 2026)."""
    out = []
    for c in table.columns:
        n = c.name.strip()
        if parse_date(n) or re.match(r"^(FY\s?\d{2,4}|Q[1-4][\s\-]?\d{2,4}|\d{4}[\s\-]?Q[1-4]|\d{4})$", n, re.I):
            out.append(c.name)
    return out


def unpivot(table: Table, id_columns: list[str], value_name: str = "Amount",
            variable_name: str = "Period") -> Table:
    by_name = {c.name: (i, c) for i, c in enumerate(table.columns)}
    missing = [n for n in id_columns if n not in by_name]
    if missing:
        raise Invalid(f"Unknown columns: {', '.join(missing)}")
    value_cols = [(i, c) for i, c in enumerate(table.columns) if c.name not in id_columns]
    if not value_cols:
        raise Invalid("Choose at least one column to keep and leave the period columns to unpivot.")
    period_dates = {c.name: parse_date(c.name) for _, c in value_cols}
    all_dates = all(period_dates.values())
    cols = [Column(f"c{k}", table.columns[by_name[n][0]].name, table.columns[by_name[n][0]].type)
            for k, n in enumerate(id_columns)]
    cols.append(Column(f"c{len(cols)}", variable_name, "date" if all_dates else "text"))
    cols.append(Column(f"c{len(cols)}", value_name, "number"))
    rows = []
    for r in table.rows:
        ids = [r[by_name[n][0]] for n in id_columns]
        for i, c in value_cols:
            v = r[i]
            if v is None:
                continue
            if not isinstance(v, (int, float)):
                v = parse_number(v)
                if v is None:
                    continue
            period = period_dates[c.name].isoformat() if all_dates else c.name
            rows.append([*ids, period, float(v)])
    return Table(table.name, table.raw, table.header_row, cols, rows)


def preview(table: Table, limit: int = 20) -> dict[str, Any]:
    return {
        "sheet": table.name,
        "header_row": table.header_row,
        "columns": [c.to_dict() for c in table.columns],
        "rows": table.rows[:limit],
        "row_count": len(table.rows),
        "period_columns": period_columns(table),
    }
