"""Files the close produces: the approved journal (in the general-ledger template's columns, ready to import into the
accounting system) and the close pack workbook, where every proposal and variance points at its evidence."""

from __future__ import annotations

import csv
import io
from typing import Any

JOURNAL_COLUMNS = ["Date", "Entry", "Type", "Account code", "Account", "Account class", "Debit", "Credit",
                   "Description", "Party", "Cost centre tags"]
ENTRY_KINDS = {"accrual", "depreciation", "bank_fee", "bank_interest"}


def journal_rows(view: dict[str, Any]) -> list[list[Any]]:
    """One row per line of every approved entry, numbered CLOSE-<month>-001 …"""
    rows: list[list[Any]] = []
    n = 0
    for it in view["items"]:
        if it["kind"] not in ENTRY_KINDS or it["status"] != "approved":
            continue
        n += 1
        entry = f"CLOSE-{view['month']}-{n:03d}"
        d = it["data"]
        day = d.get("date") if it["kind"] in ("bank_fee", "bank_interest") else view["month_end"]
        for ln in d["lines"]:
            rows.append([day, entry, "Journal", ln.get("account_code") or "", ln["account"], ln.get("account_class") or "",
                         ln["debit"] or None, ln["credit"] or None, (it["note"] or d["title"])[:300], None,
                         None])
    return rows


def journal_csv(view: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(JOURNAL_COLUMNS)
    for r in journal_rows(view):
        w.writerow(["" if v is None else v for v in r])
    return buf.getvalue().encode("utf-8-sig")


def _styles() -> dict[str, Any]:
    from openpyxl.styles import Alignment, Font, PatternFill

    return {"head": Font(bold=True, color="FFFFFF"), "fill": PatternFill("solid", fgColor="18202B"),
            "bold": Font(bold=True), "title": Font(bold=True, size=14), "muted": Font(color="5E5A60"),
            "wrap": Alignment(wrap_text=True, vertical="top"), "link": Font(color="1F4FBF", underline="single"),
            "ok": PatternFill("solid", fgColor="E3F1E6"), "warn": PatternFill("solid", fgColor="FBF0D9"),
            "todo": PatternFill("solid", fgColor="F6E0E0")}


def _table(ws: Any, header: list[str], rows: list[list[Any]], widths: list[int], st: dict[str, Any], start: int = 1) -> None:
    from openpyxl.utils import get_column_letter

    for i, h in enumerate(header, 1):
        c = ws.cell(row=start, column=i, value=h)
        c.font, c.fill = st["head"], st["fill"]
    for r, row in enumerate(rows, start + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            if isinstance(v, float):
                c.number_format = "#,##0.00;[Red]-#,##0.00"
            c.alignment = st["wrap"]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def journal_xlsx(view: dict[str, Any]) -> bytes:
    from openpyxl import Workbook

    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    rows = journal_rows(view)
    _table(ws, JOURNAL_COLUMNS, rows, [12, 20, 10, 13, 32, 15, 14, 14, 50, 16, 20], st)
    n = len(rows) + 2
    ws.cell(row=n, column=5, value="Total").font = st["bold"]
    if rows:
        ws.cell(row=n, column=7, value=f"=SUM(G2:G{n - 1})").number_format = "#,##0.00"
        ws.cell(row=n, column=8, value=f"=SUM(H2:H{n - 1})").number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pack_xlsx(view: dict[str, Any], company: str, currency: str) -> bytes:
    """The close pack: Summary (checklist), Proposed entries, Bank reconciliation, Variances and Evidence. Every row
    that rests on ledger lines links to them on the Evidence sheet."""
    from openpyxl import Workbook

    st = _styles()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    entries_ws = wb.create_sheet("Proposed entries")
    bank_ws = wb.create_sheet("Bank reconciliation")
    var_ws = wb.create_sheet("Variances")
    review_ws = wb.create_sheet("To review")
    ev_ws = wb.create_sheet("Evidence")

    # evidence first, so other sheets can link to its rows
    ev_rows: list[list[Any]] = []
    ev_at: dict[str, int] = {}
    for it in view["items"]:
        ev = it["data"].get("evidence") or []
        if not ev:
            continue
        ev_at[it["id"]] = len(ev_rows) + 2
        for e in ev:
            ev_rows.append([it["data"]["title"][:120], e.get("date"), e.get("entry"), e.get("account_code"),
                            e.get("account"), e.get("debit"), e.get("credit"), e.get("party"), e.get("description")])
    _table(ev_ws, ["For", "Date", "Entry", "Account code", "Account", "Debit", "Credit", "Party", "Description"],
           ev_rows, [40, 12, 18, 12, 28, 14, 14, 20, 50], st)

    def link(cell: Any, item_id: str) -> None:
        if item_id in ev_at:
            cell.value = "Evidence"
            cell.hyperlink = f"#Evidence!A{ev_at[item_id]}"
            cell.font = st["link"]

    # summary
    r = view["readiness"]
    summary["A1"] = f"Month-end close: {company}, {view['month']}"
    summary["A1"].font = st["title"]
    summary["A2"] = (f"Ledger “{view['ledger']['name']}”. Amounts in {currency}. Prepared "
                     f"{view.get('prepared_at_text') or ''}{' by ' + view['prepared_by_name'] if view.get('prepared_by_name') else ''}. "
                     f"Status: {view['status']}.")
    summary["A2"].font = st["muted"]
    summary["A4"] = "Ready"
    summary["B4"] = r["score"] / 100
    summary["B4"].number_format = "0%"
    summary["A4"].font = summary["B4"].font = st["bold"]
    _table(summary, ["Check", "Status", "Detail"], [[c["label"], c["status"], c["detail"]] for c in r["checks"]],
           [36, 10, 90], st, start=6)
    for i, c in enumerate(r["checks"], 7):
        summary.cell(row=i, column=2).fill = st.get(c["status"], st["warn"])
    summary.freeze_panes = None
    nb = 7 + len(r["checks"]) + 1
    summary.cell(row=nb, column=1, value="Nothing here was posted to the accounting system. Approved entries are in "
                                         "the journal file for a person to import.").font = st["muted"]

    # proposed entries: one row per line
    erows, links = [], []
    for it in view["items"]:
        if it["kind"] not in ENTRY_KINDS:
            continue
        d = it["data"]
        for j, ln in enumerate(d["lines"]):
            erows.append([d["title"] if j == 0 else "", it["status"] if j == 0 else "", ln.get("account_code"),
                          ln["account"], ln["debit"] or None, ln["credit"] or None,
                          (d["reason"] + (" " + d["warning"] if d.get("warning") else "")) if j == 0 else "",
                          it["note"] if j == 0 else "", ""])
            links.append(it["id"] if j == 0 else None)
    _table(entries_ws, ["Entry", "Status", "Account code", "Account", "Debit", "Credit", "Why", "Note", "Evidence"],
           erows, [42, 11, 12, 28, 14, 14, 70, 30, 10], st)
    for i, item_id in enumerate(links, 2):
        if item_id:
            link(entries_ws.cell(row=i, column=9), item_id)

    # bank reconciliation
    b = view.get("bank")
    if b:
        acc = b["account"]
        lines = [["Bank account", f"{acc.get('code') or ''} {acc['name']}".strip()],
                 ["Statement", f"{view['statement']['filename']} ({b['statement_lines']} lines in {view['month']})"],
                 ["Matched", f"{b['matched']} of {b['statement_lines']} statement lines"]]
        for i, (k, v) in enumerate(lines, 1):
            bank_ws.cell(row=i, column=1, value=k).font = st["bold"]
            bank_ws.cell(row=i, column=2, value=v)
        rec = [["Balance per statement", b.get("closing_balance")],
               ["+ in the ledger, not yet on the statement", b["ledger_only_total"]],
               ["= adjusted bank balance", b.get("adjusted_bank")],
               ["Balance per ledger", b["book_balance"]],
               ["+ on the statement, not in the ledger", b["statement_only_total"]],
               ["= adjusted ledger balance", b.get("adjusted_book")],
               ["Unexplained difference", b.get("difference")]]
        for i, (k, v) in enumerate(rec, 5):
            bank_ws.cell(row=i, column=1, value=k)
            c = bank_ws.cell(row=i, column=2, value=v)
            c.number_format = "#,##0.00;[Red]-#,##0.00"
        row = 13
        bank_ws.cell(row=row, column=1, value="In the ledger, not yet on the statement").font = st["bold"]
        _table(bank_ws, ["Date", "Entry", "Description", "Amount", "Days to month end"],
               [[x["date"], x["entry"], x["description"] or x.get("party") or "", x["amount"], x["age_days"]]
                for x in b["ledger_only"]], [38, 22, 46, 14, 16], st, start=row + 1)
        row += len(b["ledger_only"]) + 3
        bank_ws.cell(row=row, column=1, value="On the statement, not in the ledger").font = st["bold"]
        _table(bank_ws, ["Date", "Statement row", "Description", "Amount", ""],
               [[x["date"], x["row"], x["description"], x["amount"], ""] for x in b["statement_only"]],
               [38, 22, 46, 14, 16], st, start=row + 1)
        bank_ws.freeze_panes = None
    else:
        bank_ws["A1"] = "No bank statement was added for this month."

    # variances
    vrows, vlinks = [], []
    for it in view["items"]:
        if it["kind"] != "variance":
            continue
        d = it["data"]
        vrows.append([d.get("account_code"), d["account"], d["amount"], d["previous"], d["average"], d["difference"],
                      (d["change_pct"] / 100) if d.get("change_pct") is not None else None, d["reason"],
                      it["status"], it["note"], ""])
        vlinks.append(it["id"])
    _table(var_ws, ["Code", "Account", "This month", "Last month", "3-month average", "Difference", "Change", "Why",
                    "Status", "Explanation", "Evidence"], vrows, [10, 28, 14, 14, 15, 14, 9, 70, 11, 40, 10], st)
    for i, item_id in enumerate(vlinks, 2):
        var_ws.cell(row=i, column=7).number_format = "0%"
        link(var_ws.cell(row=i, column=11), item_id)

    # other items to review
    rrows = [[it["data"]["title"], it["data"]["reason"], it["status"], it["note"]]
             for it in view["items"] if it["kind"] in ("unbilled", "bank_unknown")]
    _table(review_ws, ["Item", "Why", "Status", "Note"], rrows, [60, 70, 11, 40], st)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
