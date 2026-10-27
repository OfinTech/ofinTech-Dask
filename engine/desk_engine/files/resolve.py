"""Turns a FileSpec into a *model*: every cell with its value, number format, live formula, readable formula and
source, laid out exactly as it will appear in Excel. The Excel and Word writers, the in-app viewer and the
reviewer all work from this model, so they always agree.

Cell keys in the model (kept short because a workbook can hold thousands of cells):
    v  value            f  format             st status: sourced | unverified | computed | label | empty
    x  Excel formula    e  readable formula   s  source id (S1…)    n  note (where the figure was found)
    d  drill-down (dataset filters behind an aggregated figure)    a  address, e.g. "P&L!C7"
"""

from __future__ import annotations

import re
import time
from typing import TYPE_CHECKING, Any

from openpyxl.utils import get_column_letter

from ..agent.harness import _NUM_RE, _numbers
from ..data.query import Filter, Measure, QuerySpec, run_query
from ..data.store import find_dataset
from ..errors import Invalid
from . import expr as X
from .evidence import Evidence, SourceBook
from .spec import ColumnSpec, FileSpec, KpiSpec, TableSpec, TextSpec, slug

if TYPE_CHECKING:
    from ..app import Engine

MAX_CELLS = 60000
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PLACEHOLDER = re.compile(r"\{\{\s*(.+?)\s*(?:\|\s*(money|number|integer|percent|text))?\s*\}\}")


def fmt_value(v: Any, f: str) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, str) or f in {"text", "date"}:
        return str(v)
    if f == "percent":
        s = f"{abs(v) * 100:,.1f}%"
    elif f == "integer":
        s = f"{abs(v):,.0f}"
    else:
        s = f"{abs(v):,.2f}"
    return f"({s})" if v < 0 and f in {"money", "percent"} else (f"-{s}" if v < 0 else s)


def _pretty_group(value: Any, grain: str | None) -> Any:
    if grain == "month" and isinstance(value, str) and re.match(r"^\d{4}-\d{2}$", value):
        return f"{MONTHS[int(value[5:7]) - 1]} {value[:4]}"
    return value


def _grain_filter(column: str, grain: str | None, value: Any, fy_start: int) -> dict[str, Any] | None:
    if value is None:
        return {"column": column, "op": "is_null"}
    if grain is None or grain == "day":
        return {"column": column, "op": "eq" if grain is None else "between",
                "value": value if grain is None else [f"{value}", f"{value}T23:59:59"]}
    v = str(value)
    if grain == "month":
        return {"column": column, "op": "between", "value": [f"{v}-01", f"{v}-31T23:59:59"]}
    if grain == "year":
        return {"column": column, "op": "between", "value": [f"{v}-01-01", f"{v}-12-31T23:59:59"]}
    if grain == "quarter":
        y, q = v.split("-Q")
        m0 = (int(q) - 1) * 3 + 1
        return {"column": column, "op": "between", "value": [f"{y}-{m0:02d}-01", f"{y}-{m0 + 2:02d}-31T23:59:59"]}
    if grain == "fiscal_year":
        end_year = int(v[2:])
        if fy_start == 1:
            return {"column": column, "op": "between", "value": [f"{end_year}-01-01", f"{end_year}-12-31T23:59:59"]}
        return {"column": column, "op": "between",
                "value": [f"{end_year - 1}-{fy_start:02d}-01", f"{end_year}-{fy_start - 1:02d}-31T23:59:59"]}
    return None  # week: no exact filter


def _guess_format(label: str, values: list[Any]) -> str:
    nums = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not nums:
        return "text"
    low = label.lower()
    if "%" in low or "percent" in low or "margin" in low and "pct" in low or low.endswith(" pct"):
        return "percent"
    if any(w in low for w in ("count", "days", "qty", "quantity", "number of", "headcount", "units")):
        return "integer"
    return "money"


def _to_number(v: Any) -> Any:
    if isinstance(v, str):
        t = v.strip()
        if re.match(r"^\(?-?[\d,]*\.?\d+\)?%?$", t) and any(ch.isdigit() for ch in t):
            neg = t.startswith("(") or t.startswith("-")
            n = float(t.strip("()%-").replace(",", ""))
            n = -n if neg else n
            return n / 100 if t.endswith("%") else n
    return v


class Resolver:
    def __init__(self, engine: Engine, workspace: dict[str, Any], spec: FileSpec, evidence: Evidence) -> None:
        self.engine = engine
        self.workspace = workspace
        self.spec = spec
        self.evidence = evidence
        self.book = SourceBook()
        self.cells: dict[str, dict[str, Any]] = {}
        self.tables: dict[str, dict[str, Any]] = {}
        self.calcs: list[dict[str, Any]] = []
        self._stack: list[str] = []
        self.fy_start = int(workspace.get("fiscal_year_start_month") or 1)

    # -- public ---------------------------------------------------------------------------------------------
    def run(self) -> dict[str, Any]:
        sheets = []
        for sheet in self.spec.sheets:
            blocks = []
            for b in sheet.blocks:
                if isinstance(b, TableSpec):
                    blocks.append(self._table(sheet.name, b))
                elif isinstance(b, TextSpec):
                    blocks.append({"type": b.type, "text": b.text})
                else:
                    blocks.append({"type": "kpis", "spec": b})
            sheets.append({"name": sheet.name, "blocks": blocks})
        self._layout(sheets)
        if len(self.cells) > MAX_CELLS:
            raise Invalid(f"The file would have more than {MAX_CELLS:,} cells; narrow the queries.")
        for sheet in sheets:  # KPI formulas can point at any table, so they are parsed after layout
            for b in sheet["blocks"]:
                if b["type"] == "kpis":
                    self._kpis(sheet["name"], b)
        for addr in list(self.cells):
            self.value(addr)
        for addr, cell in self.cells.items():
            self._compile(addr, cell)
        report = self._report() if self.spec.sections else None
        for sheet in sheets:
            for b in sheet["blocks"]:
                b.pop("spec", None)
                if b["type"] == "table":
                    b.pop("index", None)
        return {
            "version": 1,
            "kind": self.spec.kind,
            "title": self.spec.title,
            "subtitle": self.spec.subtitle,
            "currency": self.workspace.get("base_currency"),
            "created_at": time.time(),
            "sheets": [{"name": s["name"], "blocks": s["blocks"]} for s in sheets],
            "report": report,
            "sources": self.book.items,
            "calculations": self.calcs,
            "summary": self._summary(report),
        }

    # -- tables -----------------------------------------------------------------------------------------------
    def _columns(self, table: TableSpec, base: list[dict[str, Any]]) -> list[dict[str, Any]]:
        overrides = [c for c in table.columns if not c.formula]
        cols = []
        for i, b in enumerate(base):
            o: ColumnSpec | None = overrides[i] if i < len(overrides) else None
            col = dict(b)
            if o is not None:
                col["label"] = o.label
                col["key"] = o.k
                col["format"] = o.format or b["format"]
                if o.total:
                    col["total"] = o.total
                if o.width:
                    col["width"] = o.width
            cols.append(col)
        for c in table.columns:
            if c.formula:
                cols.append({"key": c.k, "label": c.label, "format": c.format or "money", "formula": c.formula,
                             "total": c.total, "width": c.width})
        keys = [c["key"] for c in cols]
        dup = {k for k in keys if keys.count(k) > 1}
        if dup:
            raise Invalid(f"Table '{table.id}' has duplicate column keys {sorted(dup)}; set 'key' explicitly. "
                          f"Keys: {', '.join(keys)}")
        for c in cols:
            if c.get("formula"):
                try:
                    c["ast"] = X.parse(c["formula"])
                except Invalid as exc:
                    raise Invalid(f"Table '{table.id}', column '{c['label']}': {exc.message}") from exc
                for n in X.references(c["ast"]):
                    if isinstance(n, X.Col) and n.key not in keys:
                        raise Invalid(f"Table '{table.id}', column '{c['label']}' uses '{n.key}', which is not a column "
                                      f"of the table. Column keys: {', '.join(keys)}")
            if not c.get("total"):
                c["total"] = ("formula" if c.get("formula") else "sum" if c["format"] in {"money", "number", "integer"}
                              else "none")
        return cols

    def _table(self, sheet: str, t: TableSpec) -> dict[str, Any]:
        if t.id in self.tables:
            raise Invalid(f"Table id '{t.id}' is used twice.")
        if t.query is not None:
            cols, rows = self._query_rows(t)
        else:
            cols, rows = self._literal_rows(t)
        table = {"type": "table", "id": t.id, "title": t.title, "sheet": sheet, "columns": cols, "rows": rows,
                 "total": None, "spec": t}
        table["has_total"] = t.total_row and bool(rows) and any(c["total"] != "none" for c in cols[1:])
        self.tables[t.id] = table
        return table

    def _query_rows(self, t: TableSpec) -> tuple[list[dict[str, Any]], list[list[dict[str, Any]]]]:
        q: QuerySpec = t.query.model_copy(deep=True)  # type: ignore[union-attr]
        if not q.measures:
            raise Invalid(f"Table '{t.id}': the query needs at least one measure.")
        ds = find_dataset(self.engine, self.workspace["id"], q.dataset)
        q.measures = [*q.measures, Measure(agg="count", label="__rows")]
        result = run_query(self.engine, self.workspace, q)
        out_cols = result["columns"]
        ng, nm = len(q.group_by), len(q.measures) - 1
        ref = ds.get("source_ref") or {}
        origin = {"excel": "Excel import", "csv": "CSV import"}.get(ds["source"], ds["source"].replace("_", " ").title())
        if isinstance(ref, dict) and ref.get("filename"):
            origin += f" ({ref['filename']})"
        sid = self.book.add(f"ds:{ds['id']}:{q.model_dump_json()}", kind="dataset",
                            title=f"Dataset “{ds['name']}”", system=origin,
                            at=ds.get("updated_at") or ds.get("created_at"), link=f"data/{ds['id']}",
                            detail=self._describe_query(q, out_cols))
        base_filters = [f.model_dump() for f in t.query.filters]  # type: ignore[union-attr]

        def drill(groups: list[tuple[int, Any]]) -> dict[str, Any]:
            fs = list(base_filters)
            exact = True
            for gi, raw in groups:
                g = q.group_by[gi]
                f = _grain_filter(g.column, g.grain, raw, self.fy_start)
                if f is None:
                    exact = False
                else:
                    fs.append(f)
            return {"dataset": ds["id"], "filters": fs, "exact": exact}

        def group_note(groups: list[tuple[int, Any]], measure: int, count: Any) -> str:
            parts = [f"{q.group_by[gi].column} = {_pretty_group(raw, q.group_by[gi].grain)}" for gi, raw in groups]
            m = q.measures[measure]
            what = (m.label or (f"{m.agg} of {m.column}" if m.column else "count")).strip()
            return " · ".join(parts + [f"{what} over {int(count or 0)} rows"])

        if t.pivot is None:
            base = []
            for i, oc in enumerate(out_cols[: ng + nm]):
                base.append({"key": slug(oc["label"]), "label": oc["label"],
                             "format": "text" if i < ng else oc["format"]})
            cols = self._columns(t, base)
            rows = []
            for r in result["rows"]:
                groups = [(gi, r[gi]) for gi in range(ng)]
                cells = [{"v": _pretty_group(r[gi], q.group_by[gi].grain), "st": "label"} for gi in range(ng)]
                for mi in range(nm):
                    v = r[ng + mi]
                    cells.append({"v": round(v, 10) if isinstance(v, float) else v, "st": "sourced", "s": sid,
                                  "n": group_note(groups, mi, r[-1]), "d": drill(groups)})
                rows.append(cells)
            return cols, self._with_formula_cells(cols, rows)

        # pivot: one column per value of the pivot group
        names = [g.column.lower() for g in q.group_by]
        if t.pivot.column.lower() not in names:
            raise Invalid(f"Table '{t.id}': pivot column '{t.pivot.column}' is not one of the query's group_by columns.")
        p = names.index(t.pivot.column.lower())
        others = [i for i in range(ng) if i != p]
        values: list[Any] = []
        grid: dict[tuple[Any, ...], dict[Any, list[Any]]] = {}
        order: list[tuple[Any, ...]] = []
        for r in result["rows"]:
            pv = r[p]
            if pv not in values:
                values.append(pv)
            key = tuple(r[i] for i in others)
            if key not in grid:
                grid[key] = {}
                order.append(key)
            grid[key][pv] = r
        if q.group_by[p].grain or all(isinstance(v, str) for v in values if v is not None):
            values.sort(key=lambda v: (v is None, str(v)))
        if len(values) > 60:
            raise Invalid(f"Table '{t.id}': the pivot would create {len(values)} columns (at most 60).")
        measure_fmt = out_cols[ng]["format"]
        base = [{"key": slug(out_cols[i]["label"]), "label": out_cols[i]["label"], "format": "text"} for i in others]
        if not others:
            base = [{"key": "item", "label": out_cols[ng]["label"], "format": "text"}]
        pkeys = []
        for v in values:
            k = slug(str(_pretty_group(v, q.group_by[p].grain)))
            while k in pkeys or k in [b["key"] for b in base]:
                k += "_"
            pkeys.append(k)
            base.append({"key": k, "label": str(_pretty_group(v, q.group_by[p].grain)), "format": measure_fmt})
        if t.pivot.row_total and len(values) > 1:
            base.append({"key": "row_total", "label": "Total", "format": measure_fmt, "span": pkeys, "total": "formula"})
        cols = self._columns(t, base)
        for c in cols:
            if c.get("span"):
                c["ast"] = X.Span(tuple(c["span"]))
                c["formula"] = "SUM(" + ", ".join(c["span"]) + ")"
        rows = []
        for key in order:
            label_cells = ([{"v": _pretty_group(v, q.group_by[i].grain), "st": "label"} for v, i in zip(key, others)]
                           if others else [{"v": out_cols[ng]["label"], "st": "label"}])
            cells = list(label_cells)
            for v in values:
                r = grid[key].get(v)
                groups = [(i, kv) for kv, i in zip(key, others)] + [(p, v)]
                if r is None:
                    cells.append({"v": None, "st": "empty"})
                else:
                    val = r[ng]
                    cells.append({"v": round(val, 10) if isinstance(val, float) else val, "st": "sourced", "s": sid,
                                  "n": group_note(groups, 0, r[-1]), "d": drill(groups)})
            rows.append(cells)
        return cols, self._with_formula_cells(cols, rows)

    def _describe_query(self, q: QuerySpec, out_cols: list[dict[str, Any]]) -> str:
        parts = ["; ".join(f"{(m.agg if m.column else 'count')} of {m.column or 'rows'}" for m in q.measures
                           if m.label != "__rows")]
        if q.group_by:
            parts.append("by " + ", ".join(g.column + (f" ({g.grain.replace('_', ' ')})" if g.grain else "")
                                           for g in q.group_by))
        for f in q.filters:
            parts.append(f"where {f.column} {f.op.replace('_', ' ')} {f.value if f.value is not None else ''}".strip())
        return ", ".join(p for p in parts if p)

    def _literal_rows(self, t: TableSpec) -> tuple[list[dict[str, Any]], list[list[dict[str, Any]]]]:
        base_specs = [c for c in t.columns if not c.formula]
        raw_rows = [[_to_number(v) for v in r] + [None] * (len(base_specs) - len(r)) for r in (t.rows or [])]
        base = []
        for i, c in enumerate(base_specs):
            values = [r[i] for r in raw_rows]
            base.append({"key": c.k, "label": c.label, "format": c.format or _guess_format(c.label, values)})
        cols = self._columns(t, base)
        rows = []
        for raw in raw_rows:
            row_label = next((str(v) for v in raw if isinstance(v, str) and v.strip()), "")
            cells = []
            for i, v in enumerate(raw):
                col = cols[i]
                if v is None or v == "":
                    cells.append({"v": None, "st": "empty"})
                elif isinstance(v, str) or col["format"] in {"text", "date"}:
                    cells.append({"v": v, "st": "label"})
                else:
                    if col["format"] == "percent" and abs(v) > 1.5:
                        v = v / 100  # 12.5 meant 12.5 %
                    cells.append(self._trace(float(v), [row_label, col["label"], t.title or ""]))
            rows.append(cells)
        return cols, self._with_formula_cells(cols, rows)

    def _trace(self, value: float, labels: list[str]) -> dict[str, Any]:
        found = self.evidence.match(value, labels) or (self.evidence.match(value * 100, labels)
                                                       if abs(value) < 1.5 else None)
        if not found:
            return {"v": value, "st": "unverified", "n": "No source found for this figure in the conversation."}
        item, line = found
        sid = self.book.add(item.key, kind=item.kind, title=item.title[:1].upper() + item.title[1:],
                            system=item.system, at=item.at, link=item.link, detail=line[:300])
        return {"v": value, "st": "sourced", "s": sid, "n": line[:300]}

    @staticmethod
    def _with_formula_cells(cols: list[dict[str, Any]], rows: list[list[dict[str, Any]]]) -> list[list[dict[str, Any]]]:
        nbase = sum(1 for c in cols if not c.get("ast"))
        out = []
        for r in rows:
            r = r[:nbase] + [{"st": "computed"} for _ in cols[nbase:]]
            out.append(r)
        return out

    # -- layout -------------------------------------------------------------------------------------------------
    def _layout(self, sheets: list[dict[str, Any]]) -> None:
        for sheet in sheets:
            row = 1
            for b in sheet["blocks"]:
                b["row"] = row
                if b["type"] in {"heading", "text", "note"}:
                    row += 2
                    continue
                if b["type"] == "kpis":
                    b["items"] = []
                    row += len(b["spec"].items) + 1
                    continue
                t = b
                if t["title"]:
                    row += 1
                t["header_row"] = row
                ncols = len(t["columns"])
                for ci, c in enumerate(t["columns"]):
                    c["letter"] = get_column_letter(ci + 1)
                t["ref_letter"] = get_column_letter(ncols + 1)
                for ri, cells in enumerate(t["rows"]):
                    r = row + 1 + ri
                    for ci, cell in enumerate(cells):
                        self._register(sheet["name"], f"{t['columns'][ci]['letter']}{r}", cell, t, ri, ci)
                last = row + len(t["rows"])
                t["first_row"], t["last_row"] = row + 1, last
                if t["has_total"]:
                    total_cells = []
                    for ci, c in enumerate(t["columns"]):
                        if ci == 0:
                            cell: dict[str, Any] = {"v": t["spec"].total_label, "st": "label"}
                        elif c["total"] == "none" or (c["total"] == "formula" and not c.get("ast")):
                            cell = {"v": None, "st": "empty"}
                        elif c["total"] == "formula":
                            cell = {"st": "computed"}
                        else:
                            cell = {"st": "computed", "agg": c["total"]}
                        total_cells.append(cell)
                        self._register(sheet["name"], f"{c['letter']}{last + 1}", cell, t, "total", ci)
                    t["total"] = total_cells
                    t["total_row"] = last + 1
                    last += 1
                row = last + 2

    def _register(self, sheet: str, coord: str, cell: dict[str, Any], table: dict[str, Any], ri: Any, ci: int) -> None:
        addr = f"{sheet}!{coord}"
        cell["a"] = addr
        cell["f"] = table["columns"][ci]["format"]
        cell["_t"] = table["id"]
        cell["_r"] = ri
        cell["_c"] = ci
        self.cells[addr] = cell

    # -- references ---------------------------------------------------------------------------------------------
    def _row_index(self, table: dict[str, Any], row: str) -> Any:
        if row.lower() == "total":
            if not table.get("total"):
                raise Invalid(f"Table '{table['id']}' has no total row.")
            return "total"
        wanted = row.strip().lower()
        for i, cells in enumerate(table["rows"]):
            if cells and str(cells[0].get("v") or "").strip().lower() == wanted:
                return i
        labels = [str(c[0].get("v")) for c in table["rows"][:25] if c]
        raise Invalid(f"Table '{table['id']}' has no row “{row}”. Rows: {', '.join(labels)}")

    def _col_index(self, table: dict[str, Any], key: str) -> int:
        keys = [c["key"] for c in table["columns"]]
        k = key.lower()
        if k in keys:
            return keys.index(k)
        labels = [c["label"].lower() for c in table["columns"]]
        if k in labels:
            return labels.index(k)
        raise Invalid(f"Table '{table['id']}' has no column '{key}'. Keys: {', '.join(keys)}")

    def _addr(self, table: dict[str, Any], ri: Any, ci: int) -> str:
        row = table["total_row"] if ri == "total" else table["first_row"] + ri
        return f"{table['sheet']}!{table['columns'][ci]['letter']}{row}"

    def _ref_addr(self, n: X.Ref) -> str:
        table = self.tables.get(n.table)
        if table is None:
            raise Invalid(f"No table with id '{n.table}'. Tables: {', '.join(self.tables)}")
        return self._addr(table, self._row_index(table, n.row), self._col_index(table, n.col))

    def _ref_label(self, n: X.Ref) -> str:
        t = self.tables[n.table]
        col = t["columns"][self._col_index(t, n.col)]["label"]
        row = t["spec"].total_label if n.row.lower() == "total" else n.row
        return f"{row} ({col}{', ' + t['title'] if t.get('title') else ''})"

    def _node_addrs(self, n: Any, ctx: dict[str, Any] | None) -> list[str]:
        if isinstance(n, X.Ref):
            return [self._ref_addr(n)]
        if ctx is None:
            raise Invalid("Use [table:row:column] references here; plain column names only work inside a table.")
        table = self.tables[ctx["_t"]]
        if isinstance(n, X.Col):
            return [self._addr(table, ctx["_r"], self._col_index(table, n.key))]
        return [self._addr(table, ctx["_r"], self._col_index(table, k)) for k in n.keys]  # Span

    # -- values ---------------------------------------------------------------------------------------------------
    def value(self, addr: str) -> Any:
        cell = self.cells[addr]
        if "v" in cell:
            return cell["v"]
        if addr in self._stack:
            raise Invalid(f"The formulas refer to each other in a loop ({' → '.join(self._stack + [addr])}).")
        self._stack.append(addr)
        try:
            if cell.get("agg"):
                table = self.tables[cell["_t"]]
                vals = [self.value(self._addr(table, i, cell["_c"])) for i in range(len(table["rows"]))]
                nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
                agg = cell["agg"]
                if agg == "sum":
                    v = sum(nums)
                elif agg == "count":
                    v = len(nums)
                elif not nums:
                    v = None
                else:
                    v = {"avg": sum(nums) / len(nums), "min": min(nums), "max": max(nums)}[agg]
            else:
                ast = self._ast(cell)
                v = X.evaluate(ast, lambda n: self._ref_value(n, cell))
            cell["v"] = round(v, 10) if isinstance(v, float) else v
        finally:
            self._stack.pop()
        return cell["v"]

    def _ast(self, cell: dict[str, Any]) -> Any:
        if "ast" in cell:
            return cell["ast"]
        table = self.tables[cell["_t"]]
        return table["columns"][cell["_c"]]["ast"]

    def _ref_value(self, n: Any, ctx: dict[str, Any]) -> Any:
        addrs = self._node_addrs(n, ctx if "_t" in ctx else None)
        if isinstance(n, X.Span):
            return [self.value(a) for a in addrs]
        v = self.value(addrs[0])
        return None if isinstance(v, str) else v

    # -- Excel formulas ---------------------------------------------------------------------------------------------
    @staticmethod
    def _quote(sheet: str) -> str:
        return "'" + sheet.replace("'", "''") + "'"

    def _excel_ref(self, addr: str, *, relative_to: str | None = None) -> str:
        sheet, coord = addr.rsplit("!", 1)
        if relative_to and relative_to.rsplit("!", 1)[0] == sheet:
            return coord
        m = re.match(r"([A-Z]+)(\d+)", coord)
        assert m
        return f"{self._quote(sheet)}!${m.group(1)}${m.group(2)}"

    def _compile(self, addr: str, cell: dict[str, Any]) -> None:
        if cell.get("st") != "computed":
            return
        if cell.get("agg"):
            table = self.tables[cell["_t"]]
            letter = table["columns"][cell["_c"]]["letter"]
            fn = {"sum": "SUM", "avg": "AVERAGE", "min": "MIN", "max": "MAX", "count": "COUNT"}[cell["agg"]]
            cell["x"] = f"={fn}({letter}{table['first_row']}:{letter}{table['last_row']})"
            cell["e"] = f"{fn.title()} of {table['columns'][cell['_c']]['label']}"
            return
        ast = self._ast(cell)
        ctx = cell if "_t" in cell else None

        def address(n: Any) -> str:
            addrs = self._node_addrs(n, ctx)
            if isinstance(n, X.Span):
                first, last = addrs[0].rsplit("!", 1)[1], addrs[-1].rsplit("!", 1)[1]
                return f"{first}:{last}"
            if isinstance(n, X.Col):
                return addrs[0].rsplit("!", 1)[1]
            return self._excel_ref(addrs[0], relative_to=addr)

        def label(n: Any) -> str:
            if isinstance(n, X.Col) and ctx:
                t = self.tables[ctx["_t"]]
                return t["columns"][self._col_index(t, n.key)]["label"]
            if isinstance(n, X.Span) and ctx:
                t = self.tables[ctx["_t"]]
                return "sum of " + ", ".join(t["columns"][self._col_index(t, k)]["label"] for k in n.keys)
            return self._ref_label(n)

        cell["x"] = X.to_excel(ast, address)
        cell["e"] = X.to_text(ast, label)

    # -- KPIs -----------------------------------------------------------------------------------------------------
    def _kpis(self, sheet: str, b: dict[str, Any]) -> None:
        spec: KpiSpec = b["spec"]
        for i, item in enumerate(spec.items):
            try:
                ast = X.parse(item.value)
            except Invalid as exc:
                raise Invalid(f"KPI '{item.label}': {exc.message}") from exc
            if any(isinstance(n, (X.Col, X.Span)) for n in X.references(ast)):
                raise Invalid(f"KPI '{item.label}': refer to cells as [table:row:column], e.g. [pl:total:amount].")
            for n in X.references(ast):
                self._ref_addr(n)  # validate early with a helpful message
            addr = f"{sheet}!B{b['row'] + i}"
            cell = {"st": "computed", "f": item.format, "a": addr, "ast": ast}
            self.cells[addr] = cell
            b["items"].append({"label": item.label, "cell": cell})

    # -- report narrative -----------------------------------------------------------------------------------------
    def _calc_for(self, addr: str) -> str:
        for c in self.calcs:
            if c.get("a") == addr:
                return c["id"]
        cell = self.cells[addr]
        cid = f"C{len(self.calcs) + 1}"
        self.calcs.append({"id": cid, "a": addr, "e": cell.get("e"), "x": cell.get("x"), "v": cell.get("v"),
                           "f": cell.get("f")})
        return cid

    def _figure_for_cell(self, addr: str) -> dict[str, Any]:
        cell = self.cells[addr]
        if cell.get("st") == "sourced":
            return {"st": "sourced", "s": cell.get("s"), "addr": addr}
        if cell.get("st") == "unverified":
            return {"st": "unverified", "addr": addr}
        return {"st": "computed", "c": self._calc_for(addr), "addr": addr}

    def _placeholder(self, text: str, fmt: str | None) -> dict[str, Any]:
        try:
            ast = X.parse(text)
        except Invalid as exc:
            raise Invalid(f"In the report text “{{{{{text}}}}}”: {exc.message}") from exc
        refs = X.references(ast)
        if any(isinstance(n, (X.Col, X.Span)) for n in refs):
            raise Invalid(f"In the report text “{{{{{text}}}}}”: refer to cells as [table:row:column].")
        if isinstance(ast, X.Ref):
            addr = self._ref_addr(ast)
            v = self.value(addr)
            f = fmt or self.cells[addr]["f"]
            return {"t": fmt_value(v, f), "fig": self._figure_for_cell(addr)}
        addrs = [self._ref_addr(n) for n in refs]  # noqa: F841 (validates)
        v = X.evaluate(ast, lambda n: self.value(self._ref_addr(n)))
        f = fmt or "number"
        cid = f"C{len(self.calcs) + 1}"
        excel = X.to_excel(ast, lambda n: self._excel_ref(self._ref_addr(n)))
        readable = X.to_text(ast, self._ref_label)
        self.calcs.append({"id": cid, "a": None, "e": readable, "x": excel, "v": v, "f": f})
        return {"t": fmt_value(v, f), "fig": {"st": "computed", "c": cid}}

    def _runs(self, text: str) -> list[dict[str, Any]]:
        runs: list[dict[str, Any]] = []
        pos = 0
        for m in PLACEHOLDER.finditer(text):
            runs.extend(self._literal_runs(text[pos:m.start()]))
            runs.append(self._placeholder(m.group(1), m.group(2)))
            pos = m.end()
        runs.extend(self._literal_runs(text[pos:]))
        return [r for r in runs if r["t"]]

    def _literal_runs(self, text: str) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        pos = 0
        for m in _NUM_RE.finditer(text):
            raw = m.group(0).strip()
            parsed = _numbers(raw)
            if not parsed:
                continue
            value = parsed[0][1]
            a = abs(value)
            if (a < 100 and "." not in raw and "%" not in raw) or (1900 <= a <= 2100 and "." not in raw
                                                                   and "," not in raw):
                continue
            start = m.start() + (len(m.group(0)) - len(m.group(0).lstrip()))
            out.append({"t": text[pos:start]})
            out.append({"t": text[start:m.end()], "fig": self._narrative_figure(raw, value)})
            pos = m.end()
        out.append({"t": text[pos:]})
        return out

    def _narrative_figure(self, raw: str, value: float) -> dict[str, Any]:
        s = raw.lower().replace(",", "")
        decimals = len(s.split(".")[1].rstrip("%kmbnilo ")) if "." in s else 0
        unit = 1.0
        for suffix, factor in (("billion", 1e9), ("million", 1e6), ("bn", 1e9), ("%", 0.01), ("k", 1e3), ("m", 1e6)):
            if s.endswith(suffix):
                unit = factor
                break
        tol = 0.5 * (10 ** -decimals) * unit + 1e-9
        a = abs(value)
        best = None
        for addr, cell in self.cells.items():
            v = cell.get("v")
            if (cell.get("st") in {"sourced", "computed", "unverified"} and isinstance(v, (int, float))
                    and not isinstance(v, bool) and abs(abs(v) - a) <= tol):
                rank = 0 if cell.get("st") == "sourced" else 1
                if best is None or rank < best[0]:
                    best = (rank, addr)
        if best:
            return self._figure_for_cell(best[1])
        traced = self._trace(value, [])
        if traced["st"] == "sourced":
            return {"st": "sourced", "s": traced["s"]}
        return {"st": "unverified"}

    def _report(self) -> dict[str, Any]:
        sections = []
        pid = 0
        for sec in self.spec.sections:
            blocks = []
            for b in sec.blocks:
                if b.type == "paragraph":
                    pid += 1
                    blocks.append({"type": "paragraph", "id": f"P{pid}", "runs": self._runs(b.text)})
                elif b.type == "bullets":
                    items = []
                    for it in b.items:
                        pid += 1
                        items.append({"id": f"P{pid}", "runs": self._runs(it)})
                    blocks.append({"type": "bullets", "items": items})
                else:
                    blocks.append({"type": "table", "table": b.table})
            sections.append({"heading": sec.heading, "blocks": blocks})
        return {"sections": sections}

    # -- summary ----------------------------------------------------------------------------------------------------
    def _summary(self, report: dict[str, Any] | None) -> dict[str, int]:
        counts = {"sourced": 0, "unverified": 0, "computed": 0}
        for cell in self.cells.values():
            st = cell.get("st")
            if st in counts and cell.get("v") is not None:
                counts[st] += 1
        figs = {"sourced": 0, "unverified": 0, "computed": 0}
        for sec in (report or {}).get("sections", []):
            for b in sec["blocks"]:
                paras = [b] if b["type"] == "paragraph" else b.get("items", [])
                for p in paras:
                    for r in p.get("runs", []):
                        if r.get("fig"):
                            figs[r["fig"]["st"]] += 1
        return {"cells_sourced": counts["sourced"], "cells_unverified": counts["unverified"],
                "cells_computed": counts["computed"], "text_sourced": figs["sourced"],
                "text_unverified": figs["unverified"], "text_computed": figs["computed"]}


def clean_model(model: dict[str, Any]) -> dict[str, Any]:
    """Drops working fields (ASTs, internal indexes) so the model can be stored as JSON."""
    for sheet in model["sheets"]:
        for b in sheet["blocks"]:
            if b["type"] == "table":
                for c in b["columns"]:
                    c.pop("ast", None)
                for cells in b["rows"] + ([b["total"]] if b.get("total") else []):
                    for cell in cells:
                        for k in ("_t", "_r", "_c", "ast"):
                            cell.pop(k, None)
            if b["type"] == "kpis":
                for it in b["items"]:
                    it["cell"].pop("ast", None)
    return model


def resolve(engine: Engine, workspace: dict[str, Any], spec: FileSpec, evidence: Evidence) -> dict[str, Any]:
    return clean_model(Resolver(engine, workspace, spec, evidence).run())


__all__ = ["Filter", "fmt_value", "resolve"]
