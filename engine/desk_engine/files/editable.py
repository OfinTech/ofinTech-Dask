"""The user's own Excel and Word files: read them, then edit them the way a careful colleague would.

The original file is kept untouched as version 1. Every edit saves a new version that keeps the file's own layout,
styles and formulas, and records what changed and where each new figure came from:

* a number written into a cell or a paragraph is matched to the conversation's evidence (tool results, documents,
  the user's messages); a match is noted in the edit log and as an Excel comment, an unmatched number is shaded
  yellow (Excel) or highlighted (Word) and listed as "without a source";
* formulas are written as formulas, after a check that refuses external links, DDE and functions that reach outside
  the workbook;
* macro-enabled and encrypted workbooks are refused rather than silently stripped.

openpyxl does not keep charts, images or pivot caches of an existing workbook; files with those are flagged before
the first edit so nobody loses them unknowingly.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import TYPE_CHECKING, Any

from ..errors import Invalid
from .limits import check_office_zip

if TYPE_CHECKING:
    from .evidence import Evidence

MAX_BYTES = 25 * 1024 * 1024
MAX_OPS = 200
PLACEHOLDER = re.compile(r"\{\{\s*[\w .\-]{1,60}\s*\}\}|\[\[\s*[\w .\-]{1,60}\s*\]\]|<<\s*[\w .\-]{1,60}\s*>>")
BAD_FORMULA = re.compile(r"\b(WEBSERVICE|FILTERXML|CALL|REGISTER(\.ID)?|EXEC|HYPERLINK|RTD|DDE|INFO|CELL|IMPORT\w*)\s*\(",
                         re.I)
CELL_REF = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
YELLOW = "FFF4CC"


def detect_format(name: str, data: bytes) -> str:
    low = name.lower()
    if len(data) > MAX_BYTES:
        raise Invalid("Files up to 25 MB can be edited.")
    if low.endswith((".xlsm", ".xltm", ".xlsb", ".xls")):
        raise Invalid("Macro-enabled and old-format workbooks cannot be edited here. Save it as .xlsx and try again.",
                      code="file_format")
    if low.endswith(".doc") or low.endswith(".docm"):
        raise Invalid("Save the document as .docx and try again.", code="file_format")
    if not data.startswith(b"PK"):
        raise Invalid("This file is not a readable .xlsx or .docx (it may be password-protected).", code="file_format")
    check_office_zip(data)
    try:
        names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
    except zipfile.BadZipFile as exc:
        raise Invalid("This file is damaged or not an Office file.", code="file_format") from exc
    if "xl/workbook.xml" in names:
        if "xl/vbaProject.bin" in names:
            raise Invalid("This workbook contains macros; save a copy as .xlsx without macros first.",
                          code="file_format")
        return "xlsx"
    if "word/document.xml" in names:
        return "docx"
    raise Invalid("Only Excel (.xlsx) and Word (.docx) files can be edited.", code="file_format")


def at_risk(data: bytes, fmt: str) -> list[str]:
    """Parts of an existing file that an edit would drop (openpyxl/python-docx limits)."""
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    if fmt != "xlsx":
        return []
    out = []
    if any(n.startswith("xl/charts/") for n in names):
        out.append("charts")
    if any(n.startswith("xl/media/") for n in names):
        out.append("images")
    if any(n.startswith("xl/pivotCache") for n in names):
        out.append("pivot tables")
    return out


def _num(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _text(v: Any, limit: int = 80) -> str:
    s = str(v)
    return s if len(s) <= limit else s[: limit - 1] + "…"


# -- Excel -----------------------------------------------------------------------------------------------------------
def _load_xlsx(data: bytes, values: bool = False) -> Any:
    from openpyxl import load_workbook

    try:
        return load_workbook(io.BytesIO(data), data_only=values)
    except Exception as exc:  # noqa: BLE001 - any parse failure is the user's file
        raise Invalid("The workbook could not be opened.", code="file_format") from exc


def outline_xlsx(data: bytes) -> dict[str, Any]:
    wb = _load_xlsx(data)
    sheets = []
    for ws in wb.worksheets:
        formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        sheets.append({"name": ws.title, "range": ws.dimensions, "rows": ws.max_row, "columns": ws.max_column,
                       "formulas": formulas})
    return {"format": "xlsx", "sheets": sheets}


def inspect_xlsx(data: bytes, sheet: str | None = None, cell_range: str | None = None, limit: int = 400) -> str:
    """A compact listing: one line per non-empty cell, formulas shown with their last saved value."""
    wb = _load_xlsx(data)
    vals = _load_xlsx(data, values=True)
    lines = []
    names = [ws.title for ws in wb.worksheets]
    if sheet and sheet not in names:
        raise Invalid(f"No sheet named “{sheet}”. Sheets: {', '.join(names)}.")
    for ws in wb.worksheets:
        lines.append(f"Sheet “{ws.title}” ({ws.dimensions}, {ws.max_row} rows × {ws.max_column} columns)")
    targets = [wb[sheet]] if sheet else wb.worksheets
    shown = 0
    for ws in targets:
        lines.append(f"\n## {ws.title}")
        if cell_range:
            try:
                got = ws[cell_range]
            except ValueError as exc:
                raise Invalid(f"“{cell_range}” is not a cell range such as A1:F40.") from exc
            rows = ((got,),) if not isinstance(got, tuple) else (got,) if got and not isinstance(got[0], tuple) else got
        else:
            rows = ws.iter_rows()
        for row in rows:
            parts = []
            for c in row:
                if c.value is None:
                    continue
                if isinstance(c.value, str) and c.value.startswith("="):
                    saved = vals[ws.title][c.coordinate].value
                    parts.append(f"{c.coordinate}={_text(c.value, 120)}" + (f" → {_text(saved)}" if saved is not None else ""))
                else:
                    parts.append(f"{c.coordinate}: {_text(c.value)}")
            if parts:
                lines.append(" | ".join(parts))
                shown += len(parts)
            if shown >= limit:
                lines.append(f"… more cells; inspect a smaller range (e.g. sheet='{ws.title}', range='A1:H60').")
                return "\n".join(lines)
    return "\n".join(lines)


def check_formula(f: str) -> str:
    f = f.strip()
    if not f.startswith("="):
        f = "=" + f
    if len(f) > 2000:
        raise Invalid("That formula is too long.")
    if BAD_FORMULA.search(f) or "|" in f or re.search(r"\[[^\]]*\.(xls|xlsx|xlsm|csv)\]", f, re.I) or "://" in f:
        raise Invalid("Formulas may only refer to this workbook; links to other files, web services or DDE are not "
                      "allowed.", code="formula_refused")
    if f.count("(") != f.count(")"):
        raise Invalid(f"Unbalanced brackets in {f}.")
    return f


def _labels(ws: Any, cell: Any) -> list[str]:
    """Row and column labels around a cell (the first text to its left and the first text above it)."""
    out = []
    for col in range(cell.column - 1, 0, -1):
        v = ws.cell(row=cell.row, column=col).value
        if isinstance(v, str) and not v.startswith("="):
            out.append(v)
            break
    for row in range(cell.row - 1, 0, -1):
        v = ws.cell(row=row, column=cell.column).value
        if isinstance(v, str) and not v.startswith("="):
            out.append(v)
            break
    return out


def _sheet(wb: Any, name: str | None) -> Any:
    if not name:
        return wb.worksheets[0]
    if name not in wb.sheetnames:
        raise Invalid(f"No sheet named “{name}”. Sheets: {', '.join(wb.sheetnames)}.")
    return wb[name]


def _write_value(ws: Any, cell: Any, value: Any, formula: str | None, evidence: Evidence, log: list[dict[str, Any]],
                 user_value: bool = False) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    entry: dict[str, Any] = {"op": "set", "target": f"{ws.title}!{cell.coordinate}", "before": cell.value}
    if formula:
        cell.value = check_formula(formula)
        entry |= {"after": cell.value, "status": "formula"}
    else:
        if isinstance(value, str) and value.startswith("="):
            raise Invalid("Use `formula` for formulas, not `value`.")
        cell.value = value
        entry["after"] = value
        n = _num(value)
        if n is None or user_value:
            entry["status"] = "text" if n is None else "entered"
        else:
            hit = evidence.match(n, _labels(ws, cell))
            if hit:
                item, line = hit
                entry |= {"status": "sourced", "source": item.title, "system": item.system, "line": line[:300],
                          "link": item.link}
                cell.comment = Comment(f"Source: {item.title}\n{line[:200]}", "OfinTech Desk")
            else:
                entry["status"] = "unverified"
                cell.fill = PatternFill("solid", fgColor=YELLOW)
                cell.comment = Comment("No source found for this figure in the conversation. Check it.", "OfinTech Desk")
    log.append(entry)


def apply_xlsx(data: bytes, ops: list[dict[str, Any]], evidence: Evidence) -> tuple[bytes, list[dict[str, Any]]]:
    from copy import copy

    wb = _load_xlsx(data)
    log: list[dict[str, Any]] = []
    for op in ops:
        kind = op.get("op")
        if kind == "set_cell":
            ws = _sheet(wb, op.get("sheet"))
            ref = str(op.get("cell") or "").upper().replace("$", "")
            if not CELL_REF.match(ref):
                raise Invalid(f"“{op.get('cell')}” is not a cell address such as C7.")
            _write_value(ws, ws[ref], op.get("value"), op.get("formula"), evidence, log)
        elif kind == "append_row":
            ws = _sheet(wb, op.get("sheet"))
            values = op.get("values") or []
            if not isinstance(values, list) or not values:
                raise Invalid("append_row needs `values`, a list with one entry per column.")
            after = op.get("after_row")
            if after:
                row = int(after) + 1
                ws.insert_rows(row)
            else:
                row = ws.max_row + 1
            src = row - 1 if row > 1 else None
            for i, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=i)
                if src:  # copy the style of the row above (number format, font, borders)
                    above = ws.cell(row=src, column=i)
                    if above.has_style:
                        cell._style = copy(above._style)
                if v is None or v == "":
                    continue
                if isinstance(v, str) and v.startswith("="):
                    _write_value(ws, cell, None, v, evidence, log)
                else:
                    _write_value(ws, cell, v, None, evidence, log)
            log.append({"op": "append_row", "target": f"{ws.title}!{row}", "after": len(values), "status": "structure"})
        elif kind == "add_sheet":
            name = str(op.get("name") or "").strip()[:31]
            if not name or re.search(r"[\[\]:*?/\\]", name):
                raise Invalid("A sheet name needs 1–31 characters and none of [ ] : * ? / \\.")
            if name in wb.sheetnames:
                raise Invalid(f"There is already a sheet named “{name}”.")
            ws = wb.create_sheet(name)
            for r, row in enumerate(op.get("rows") or [], start=1):
                for c, v in enumerate(row, start=1):
                    if v is None or v == "":
                        continue
                    if isinstance(v, str) and v.startswith("="):
                        _write_value(ws, ws.cell(row=r, column=c), None, v, evidence, log)
                    else:
                        _write_value(ws, ws.cell(row=r, column=c), v, None, evidence, log)
            log.append({"op": "add_sheet", "target": name, "status": "structure"})
        elif kind == "rename_sheet":
            ws = _sheet(wb, op.get("sheet"))
            new = str(op.get("name") or "").strip()[:31]
            if not new or new in wb.sheetnames:
                raise Invalid("Give a new, unused sheet name.")
            log.append({"op": "rename_sheet", "target": ws.title, "after": new, "status": "structure"})
            ws.title = new
        else:
            raise Invalid(f"Unknown Excel edit “{kind}”. Use set_cell, append_row, add_sheet or rename_sheet.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), log


# -- Word ------------------------------------------------------------------------------------------------------------
def _load_docx(data: bytes) -> Any:
    from docx import Document

    try:
        return Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise Invalid("The document could not be opened.", code="file_format") from exc


def _paragraphs(doc: Any) -> list[Any]:
    return list(doc.paragraphs)


def outline_docx(data: bytes) -> dict[str, Any]:
    doc = _load_docx(data)
    paras = _paragraphs(doc)
    text = "\n".join(p.text for p in paras) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    return {"format": "docx", "paragraphs": len(paras), "tables": len(doc.tables),
            "headings": [p.text[:80] for p in paras if p.style is not None and p.style.name.lower().startswith("heading")
                         and p.text.strip()][:30],
            "placeholders": sorted(set(PLACEHOLDER.findall(text)))[:60]}


def inspect_docx(data: bytes, start: int = 0, limit: int = 120) -> str:
    doc = _load_docx(data)
    paras = _paragraphs(doc)
    o = outline_docx(data)
    lines = [f"{len(paras)} paragraphs, {len(doc.tables)} tables."]
    if o["placeholders"]:
        lines.append("Placeholders to fill: " + ", ".join(o["placeholders"]))
    for i, p in enumerate(paras[start:start + limit], start=start):
        if not p.text.strip():
            continue
        style = p.style.name if p.style is not None else ""
        tag = f" [{style}]" if style and style != "Normal" else ""
        lines.append(f"P{i}{tag}: {_text(p.text, 400)}")
    if start + limit < len(paras):
        lines.append(f"… more paragraphs; call again with start={start + limit}.")
    for ti, t in enumerate(doc.tables):
        lines.append(f"\nTable T{ti} ({len(t.rows)} rows × {len(t.columns)} columns)")
        for ri, r in enumerate(t.rows[:40]):
            lines.append(f"  R{ri}: " + " | ".join(_text(c.text, 60) for c in r.cells))
        if len(t.rows) > 40:
            lines.append("  …")
    return "\n".join(lines)


def _set_paragraph_text(p: Any, text: str) -> None:
    """Replace a paragraph's text and keep the formatting of its first run."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def _replace_in_paragraph(p: Any, find: str, repl: str) -> int:
    if find not in p.text:
        return 0
    # fast path: the text sits inside one run (keeps all formatting)
    n = 0
    for r in p.runs:
        if find in r.text:
            n += r.text.count(find)
            r.text = r.text.replace(find, repl)
    if n:
        return n
    n = p.text.count(find)  # split across runs (Word does this often): rewrite into the first run
    _set_paragraph_text(p, p.text.replace(find, repl))
    return n


def _all_paragraphs(doc: Any) -> list[Any]:
    out = list(doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                out.extend(c.paragraphs)
    for s in doc.sections:
        for part in (s.header, s.footer):
            out.extend(part.paragraphs)
    return out


def _check_numbers(text: str, evidence: Evidence, labels: list[str]) -> tuple[list[dict[str, Any]], list[str]]:
    from ..agent.harness import _numbers

    sourced, missing = [], []
    for raw, val in _numbers(text):
        if abs(val) < 100 and float(val).is_integer():  # section numbers, counts, payment days
            continue
        hit = evidence.match(val, labels)
        if hit:
            sourced.append({"figure": raw, "source": hit[0].title, "line": hit[1][:300], "link": hit[0].link})
        else:
            missing.append(raw)
    return sourced, missing


def _highlight(p: Any, figures: list[str]) -> None:
    from docx.enum.text import WD_COLOR_INDEX

    for r in p.runs:
        if any(f in r.text for f in figures):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _logged(op: str, target: str, before: str, after: str, evidence: Evidence, labels: list[str],
            paras: list[Any]) -> dict[str, Any]:
    sourced, missing = _check_numbers(after, evidence, labels)
    for p in paras:
        _highlight(p, missing)
    return {"op": op, "target": target, "before": before[:500], "after": after[:500], "sources": sourced,
            "unverified": missing, "status": "unverified" if missing else ("sourced" if sourced else "text")}


def apply_docx(data: bytes, ops: list[dict[str, Any]], evidence: Evidence) -> tuple[bytes, list[dict[str, Any]]]:
    doc = _load_docx(data)
    log: list[dict[str, Any]] = []
    for op in ops:
        kind = op.get("op")
        if kind in {"fill", "replace_text"}:
            pairs = (op.get("values") or {}) if kind == "fill" else {str(op.get("find") or ""): op.get("replace", "")}
            if not pairs or any(not k for k in pairs):
                raise Invalid("fill needs `values` {placeholder: text}; replace_text needs `find` and `replace`.")
            for find, repl in pairs.items():
                repl = "" if repl is None else str(repl)
                if kind == "replace_text" or PLACEHOLDER.fullmatch(find.strip()):
                    candidates = [find.strip() if kind == "fill" else find]
                else:  # a bare name fills its placeholder only, never the same word in the text
                    candidates = ["{{" + find + "}}", "{{ " + find + " }}", "[[" + find + "]]", "<<" + find + ">>"]
                hits, touched = 0, []
                for p in _all_paragraphs(doc):
                    for c in candidates:
                        k = _replace_in_paragraph(p, c, repl)
                        if k:
                            hits += k
                            touched.append(p)
                if not hits:
                    raise Invalid(f"“{find}” does not appear in the document. Inspect it to see the exact text.")
                entry = _logged(kind, find, find, repl, evidence, [find.strip("{}[]<> ")], touched)
                entry["count"] = hits
                log.append(entry)
        elif kind == "set_paragraph":
            paras = _paragraphs(doc)
            i = int(op.get("index", -1))
            if not 0 <= i < len(paras):
                raise Invalid(f"There is no paragraph P{i}; the document has P0–P{len(paras) - 1}.")
            before = paras[i].text
            _set_paragraph_text(paras[i], str(op.get("text") or ""))
            log.append(_logged("set_paragraph", f"P{i}", before, paras[i].text, evidence, [], [paras[i]]))
        elif kind == "insert_paragraph":
            paras = _paragraphs(doc)
            text = str(op.get("text") or "")
            style = op.get("style")
            after = op.get("after")
            if after is None or int(after) >= len(paras) - 1:
                p = doc.add_paragraph(text)
                where = "end"
            else:
                anchor = paras[int(after) + 1]
                p = anchor.insert_paragraph_before(text)
                where = f"after P{int(after)}"
            if style:
                try:
                    p.style = doc.styles[style]
                except KeyError:
                    pass  # the template has no such style; keep Normal
            log.append(_logged("insert_paragraph", where, "", text, evidence, [], [p]))
        elif kind == "set_table_cell":
            ti, ri, ci = int(op.get("table", -1)), int(op.get("row", -1)), int(op.get("column", -1))
            if not 0 <= ti < len(doc.tables):
                raise Invalid(f"There is no table T{ti}.")
            t = doc.tables[ti]
            if not (0 <= ri < len(t.rows) and 0 <= ci < len(t.columns)):
                raise Invalid(f"Table T{ti} has {len(t.rows)} rows and {len(t.columns)} columns (counting from 0).")
            cell = t.cell(ri, ci)
            before = cell.text
            text = str(op.get("text") if op.get("text") is not None else "")
            if cell.paragraphs:
                _set_paragraph_text(cell.paragraphs[0], text)
                for extra in cell.paragraphs[1:]:
                    extra._element.getparent().remove(extra._element)
            else:
                cell.text = text
            labels = [t.cell(ri, 0).text, t.cell(0, ci).text]
            log.append(_logged("set_table_cell", f"T{ti} R{ri} C{ci}", before, text, evidence, labels, cell.paragraphs))
        elif kind == "add_table_row":
            ti = int(op.get("table", -1))
            if not 0 <= ti < len(doc.tables):
                raise Invalid(f"There is no table T{ti}.")
            t = doc.tables[ti]
            values = [str("" if v is None else v) for v in (op.get("values") or [])]
            row = t.add_row()
            for ci, v in enumerate(values[: len(row.cells)]):
                row.cells[ci].text = v
            paras = [p for c in row.cells for p in c.paragraphs]
            log.append(_logged("add_table_row", f"T{ti} R{len(t.rows) - 1}", "", " | ".join(values), evidence,
                               values[:1], paras))
        else:
            raise Invalid(f"Unknown Word edit “{kind}”. Use fill, replace_text, set_paragraph, insert_paragraph, "
                          "set_table_cell or add_table_row.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), log


def summarize(log: list[dict[str, Any]]) -> dict[str, int]:
    out = {"edits": len(log), "sourced": 0, "unverified": 0, "formulas": 0}
    for e in log:
        if e.get("status") == "sourced":
            out["sourced"] += 1
        elif e.get("status") == "unverified":
            out["unverified"] += 1
        elif e.get("status") == "formula":
            out["formulas"] += 1
    return out
