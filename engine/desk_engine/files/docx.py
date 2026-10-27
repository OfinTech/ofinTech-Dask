"""Writes a report model as a Word document.

Every figure in the narrative carries a marker: [S2] for a source, [C1] for a calculation. Markers link to the
"Sources and calculations" section at the end, which links on to the app. Figures without a source are highlighted
and carry a comment, and reviewer findings are real Word comments on the paragraphs they concern.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .resolve import fmt_value

AUTHOR = "OfinTech Desk"
REVIEWER = "OfinTech Reviewer"
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def _internal_link(paragraph: Any, text: str, anchor: str, *, superscript: bool = True) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    if superscript:
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _external_link(paragraph: Any, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _bookmark(paragraph: Any, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _shade(cell: Any, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def render_docx(model: dict[str, Any], *, findings: list[dict[str, Any]], meta: dict[str, Any],
                app_url: Callable[[str], str | None]) -> bytes:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    open_findings = [f for f in findings if f.get("status", "open") == "open"]
    by_target: dict[str, list[dict[str, Any]]] = {}
    for f in open_findings:
        by_target.setdefault(f["target"], []).append(f)
    tables = {b["id"]: b for s in model["sheets"] for b in s["blocks"] if b["type"] == "table"}
    used_sources: list[str] = []
    used_calcs: list[str] = []

    def comment(runs: Any, f: dict[str, Any]) -> None:
        text = f"{f['severity'].title()}: {f['title']}" + (f"\n{f['body']}" if f.get("body") else "")
        doc.add_comment(runs, text=text, author=REVIEWER if f.get("origin") == "reviewer" else AUTHOR,
                        initials="OR" if f.get("origin") == "reviewer" else "OD")

    title = doc.add_heading(model["title"], level=0)
    for f in by_target.get("file", []):
        comment(title.runs, f)
    if model.get("subtitle"):
        doc.add_paragraph(model["subtitle"], style="Subtitle")
    p = doc.add_paragraph()
    r = p.add_run(f"Prepared by {meta.get('author') or 'OfinTech Desk'} on "
                  f"{datetime.fromtimestamp(model['created_at']).strftime('%d %B %Y')}, version {meta.get('version', 1)}. "
                  "Figures are marked with their source [S] or calculation [C], listed at the end. "
                  "Highlighted figures have no source and must be checked.")
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED

    def runs_into(par: Any, runs: list[dict[str, Any]]) -> None:
        for run in runs:
            fig = run.get("fig")
            w = par.add_run(run["t"])
            if not fig:
                continue
            if fig["st"] == "unverified":
                w.font.highlight_color = WD_COLOR_INDEX.YELLOW
                doc.add_comment(w, text="No source found for this figure. Check it before the report is shared.",
                                author=AUTHOR, initials="OD")
                continue
            ref = fig.get("s") if fig["st"] == "sourced" else fig.get("c")
            if fig["st"] == "sourced" and fig.get("s") and fig["s"] not in used_sources:
                used_sources.append(fig["s"])
            if fig["st"] == "computed" and fig.get("c") and fig["c"] not in used_calcs:
                used_calcs.append(fig["c"])
            if ref:
                _internal_link(par, f"[{ref}]", f"ref_{ref}")

    for sec in (model.get("report") or {}).get("sections", []):
        doc.add_heading(sec["heading"], level=1)
        for b in sec["blocks"]:
            if b["type"] == "paragraph":
                par = doc.add_paragraph()
                runs_into(par, b["runs"])
                for f in by_target.get(b["id"], []):
                    if par.runs:
                        comment(par.runs, f)
            elif b["type"] == "bullets":
                for it in b["items"]:
                    par = doc.add_paragraph(style="List Bullet")
                    runs_into(par, it["runs"])
                    for f in by_target.get(it["id"], []):
                        if par.runs:
                            comment(par.runs, f)
            elif b["type"] == "table" and b["table"] in tables:
                _table(doc, tables[b["table"]], by_target, comment, used_sources, model)

    # Sources and calculations
    doc.add_heading("Sources and calculations", level=1)
    bid = 100
    sources = {s["id"]: s for s in model["sources"]}
    for sid in sorted(used_sources, key=lambda x: int(x[1:])) + [s for s in sources if s not in used_sources]:
        src = sources.get(sid)
        if not src:
            continue
        par = doc.add_paragraph()
        _bookmark(par, f"ref_{sid}", bid)
        bid += 1
        r = par.add_run(f"[{sid}] ")
        r.bold = True
        par.add_run(f"{src['title']} — {src.get('system') or ''}")
        if src.get("at"):
            par.add_run(f", retrieved {datetime.fromtimestamp(src['at']).strftime('%d %b %Y %H:%M')}")
        par.add_run(". ")
        if src.get("detail"):
            d = par.add_run(str(src["detail"])[:400] + " ")
            d.font.color.rgb = MUTED
        url = app_url(src["link"]) if src.get("link") else None
        if url:
            _external_link(par, "Open", url)
    calcs = {c["id"]: c for c in model.get("calculations", [])}
    for cid in used_calcs:
        calc = calcs.get(cid)
        if not calc:
            continue
        par = doc.add_paragraph()
        _bookmark(par, f"ref_{cid}", bid)
        bid += 1
        r = par.add_run(f"[{cid}] ")
        r.bold = True
        where = f" (workbook cell {calc['a']})" if calc.get("a") else ""
        par.add_run(f"{calc.get('e') or 'Calculation'} = {fmt_value(calc.get('v'), calc.get('f') or 'number')}{where}.")
    note = doc.add_paragraph()
    n = note.add_run("The companion Excel workbook holds every table with live formulas and the same source references.")
    n.font.size = Pt(8.5)
    n.font.color.rgb = MUTED
    buf = io.BytesIO()
    doc.core_properties.author = AUTHOR
    doc.core_properties.title = model["title"]
    doc.save(buf)
    return buf.getvalue()


def _table(doc: Any, t: dict[str, Any], by_target: dict[str, list[dict[str, Any]]], comment: Any,
           used_sources: list[str], model: dict[str, Any]) -> None:
    if t.get("title"):
        cap = doc.add_paragraph()
        cr = cap.add_run(t["title"])
        cr.bold = True
    rows = t["rows"] + ([t["total"]] if t.get("total") else [])
    table = doc.add_table(rows=1 + len(rows), cols=len(t["columns"]))
    table.style = "Table Grid"
    for ci, col in enumerate(t["columns"]):
        cell = table.cell(0, ci)
        cell.text = ""
        run = cell.paragraphs[0].add_run(col["label"])
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, "E7EDF3")
    sids: list[str] = []
    flagged: list[dict[str, Any]] = []
    for ri, cells in enumerate(rows, start=1):
        is_total = t.get("total") is not None and ri == len(rows)
        for ci, c in enumerate(cells):
            wc = table.cell(ri, ci)
            wc.text = ""
            par = wc.paragraphs[0]
            run = par.add_run(fmt_value(c.get("v"), c.get("f", "text")))
            run.font.size = Pt(9)
            run.bold = is_total
            if c.get("f") not in (None, "text"):
                par.alignment = 2  # right
            if c.get("st") == "unverified":
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if c.get("s") and c["s"] not in sids:
                sids.append(c["s"])
            flagged.extend(by_target.get(c.get("a", ""), []))
    for s in sids:
        if s not in used_sources:
            used_sources.append(s)
    formulas = [f"{col['label']} = {t['rows'][0][ci].get('e') or col['formula']}"
                for ci, col in enumerate(t["columns"]) if col.get("formula") and t["rows"]]
    cap = doc.add_paragraph()
    parts = []
    if sids:
        parts.append("Source: ")
    r = cap.add_run("".join(parts))
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    for i, s in enumerate(sids[:12]):
        _internal_link(cap, s + (", " if i < min(len(sids), 12) - 1 else ""), f"ref_{s}", superscript=False)
    tail = [" ·"] if sids else []
    if formulas:
        tail.append(" Formulas: " + "; ".join(formulas) + ".")
    tail.append(f" Live formulas: workbook sheet “{t['sheet']}”.")
    r2 = cap.add_run("".join(tail))
    r2.font.size = Pt(8)
    r2.font.color.rgb = MUTED
    for f in flagged:
        comment(cap.runs, {**f, "title": f"{f['target']}: {f['title']}"})
