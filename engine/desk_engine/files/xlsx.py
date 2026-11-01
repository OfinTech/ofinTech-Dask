"""Writes a file model as an Excel workbook.

* Computed cells are live formulas (Excel recalculates them when the file opens).
* Modelling colour convention: blue figures are inputs taken from a source, black figures are formulas.
* Every input carries a note naming its source; the Ref column links each row to the Sources sheet.
* Figures without a source are shaded yellow. Reviewer findings are notes on the cells they concern and are
  listed on the Review sheet.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.hyperlink import Hyperlink

from .resolve import fmt_value

NUMBER_FORMATS = {
    "money": '#,##0.00;(#,##0.00);"–"',
    "number": "#,##0.00",
    "integer": "#,##0",
    "percent": '0.0%;(0.0%);"–"',
    "date": "yyyy-mm-dd",
}
INPUT = Font(color="1F4E79")
INPUT_BOLD = Font(color="1F4E79", bold=True)
BOLD = Font(bold=True)
MUTED = Font(color="6B7280", size=9)
LINK = Font(color="1F4E79", underline="single", size=9)
HEADER_FILL = PatternFill("solid", fgColor="E7EDF3")
UNVERIFIED_FILL = PatternFill("solid", fgColor="FFF2CC")
FINDING_FILL = {"high": PatternFill("solid", fgColor="F8D7DA"), "medium": PatternFill("solid", fgColor="FFE8CC"),
                "low": PatternFill("solid", fgColor="E2ECF7")}
THIN = Side(style="thin", color="9CA3AF")
DOUBLE = Side(style="double", color="111827")
AUTHOR = "OfinTech Desk"


def _q(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def _set(cell: Any, value: Any) -> None:
    from .cells import set_cell

    set_cell(cell, value)  # text that looks like a formula stays text, whatever its first character


def _value(v: Any, f: str) -> Any:
    if f == "date" and isinstance(v, str):
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return v
    return v


def _note(lines: list[str]) -> Comment:
    text = "\n".join(lines)[:1500]
    c = Comment(text, AUTHOR)
    c.width = 320
    c.height = min(60 + 16 * text.count("\n") + len(text) // 3, 260)
    return c


def render_xlsx(model: dict[str, Any], *, findings: list[dict[str, Any]], meta: dict[str, Any],
                app_url: Callable[[str], str | None]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = AUTHOR
    wb.properties.title = model["title"]
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    sources_sheet = "Sources"
    source_rows: dict[str, int] = {}
    # Sources rows are known in advance so cells can link to them
    first_source_row = 8
    for i, s in enumerate(model["sources"]):
        source_rows[s["id"]] = first_source_row + i
    by_target: dict[str, list[dict[str, Any]]] = {}
    for f in findings:
        by_target.setdefault(f["target"], []).append(f)

    for sheet in model["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        widths: dict[int, int] = {}

        def width(col: int, text: Any) -> None:
            widths[col] = max(widths.get(col, 8), min(len(str(text)) + 2, 60))

        freeze = None
        for b in sheet["blocks"]:
            r = b["row"]
            if b["type"] == "heading":
                c = ws.cell(r, 1)
                _set(c, b["text"])
                c.font = Font(bold=True, size=13)
            elif b["type"] in {"text", "note"}:
                c = ws.cell(r, 1)
                _set(c, b["text"])
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if b["type"] == "note":
                    c.font = Font(italic=True, color="4B5563")
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                ws.row_dimensions[r].height = max(15, 15 * (len(b["text"]) // 110 + 1))
            elif b["type"] == "kpis":
                for i, it in enumerate(b["items"]):
                    lc = ws.cell(r + i, 1)
                    _set(lc, it["label"])
                    lc.font = BOLD
                    width(1, it["label"])
                    vc = ws.cell(r + i, 2)
                    _write_cell(ws, vc, it["cell"], by_target, model, source_rows, sources_sheet)
                    width(2, fmt_value(it["cell"].get("v"), it["cell"]["f"]))
            elif b["type"] == "table":
                if b.get("title"):
                    c = ws.cell(r, 1)
                    _set(c, b["title"])
                    c.font = Font(bold=True, size=11)
                h = b["header_row"]
                if freeze is None and h <= 4:
                    freeze = f"A{h + 1}"
                ncols = len(b["columns"])
                for ci, col in enumerate(b["columns"], start=1):
                    c = ws.cell(h, ci)
                    _set(c, col["label"])
                    c.font = BOLD
                    c.fill = HEADER_FILL
                    c.border = Border(bottom=THIN)
                    c.alignment = Alignment(horizontal="left" if col["format"] == "text" else "right", wrap_text=True)
                    width(ci, col["label"])
                    if col.get("width"):
                        widths[ci] = col["width"]
                rc = ws.cell(h, ncols + 1)
                _set(rc, "Ref")
                rc.font = MUTED
                rc.fill = HEADER_FILL
                rc.border = Border(bottom=THIN)
                widths[ncols + 1] = 8
                for cells in b["rows"]:
                    sids: list[str] = []
                    row_no = 0
                    for ci, cell in enumerate(cells, start=1):
                        c = ws.cell(_row(cell["a"]), ci)
                        row_no = c.row
                        _write_cell(ws, c, cell, by_target, model, source_rows, sources_sheet)
                        width(ci, fmt_value(cell.get("v"), cell["f"]))
                        if cell.get("s") and cell["s"] not in sids:
                            sids.append(cell["s"])
                    if sids and row_no:
                        rc = ws.cell(row_no, ncols + 1)
                        _set(rc, ", ".join(sids[:4]))
                        rc.font = LINK
                        rc.hyperlink = Hyperlink(ref=rc.coordinate, location=f"{_q(sources_sheet)}!A{source_rows[sids[0]]}",
                                                 tooltip="Open the source")
                if b.get("total"):
                    for ci, cell in enumerate(b["total"], start=1):
                        c = ws.cell(_row(cell["a"]), ci)
                        _write_cell(ws, c, cell, by_target, model, source_rows, sources_sheet)
                        c.font = Font(bold=True, color=c.font.color if c.font else None)
                        c.border = Border(top=THIN, bottom=DOUBLE)
                        width(ci, fmt_value(cell.get("v"), cell["f"]))
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w
        if freeze:
            ws.freeze_panes = freeze
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    open_findings = [f for f in findings if f.get("status", "open") == "open"]
    if open_findings:
        _review_sheet(wb, open_findings)
    _sources_sheet(wb, model, meta, app_url, first_source_row)
    if model.get("calculations"):
        _calculations(wb[sources_sheet], model, first_source_row + len(model["sources"]) + 2)
    for name in ("Review", sources_sheet):
        if name in wb.sheetnames:
            ws = wb[name]
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(addr: str) -> int:
    coord = addr.rsplit("!", 1)[1]
    return int("".join(ch for ch in coord if ch.isdigit()))


def _write_cell(ws: Any, c: Any, cell: dict[str, Any], by_target: dict[str, list[dict[str, Any]]],
                model: dict[str, Any], source_rows: dict[str, int], sources_sheet: str) -> None:
    st = cell.get("st")
    f = cell.get("f", "text")
    if st == "computed" and cell.get("x"):
        c.value = cell["x"]
    else:
        _set(c, _value(cell.get("v"), f))
    if f in NUMBER_FORMATS and st != "label":
        c.number_format = NUMBER_FORMATS[f]
    if st in {"sourced", "unverified"}:
        c.font = INPUT
    if f != "text" and st != "label":
        c.alignment = Alignment(horizontal="right")
    lines: list[str] = []
    if st == "sourced" and cell.get("s"):
        src = next((s for s in model["sources"] if s["id"] == cell["s"]), None)
        if src:
            lines.append(f"{src['id']}: {src['title']} — {src['system']}")
        if cell.get("n"):
            lines.append(cell["n"])
    elif st == "unverified":
        c.fill = UNVERIFIED_FILL
        lines.append("No source found for this figure in the conversation. Check it before relying on it.")
    for f_ in by_target.get(cell.get("a", ""), []):
        if f_.get("status", "open") != "open":
            continue
        lines.append(f"Review ({f_['severity']}): {f_['title']}" + (f"\n{f_['body']}" if f_.get("body") else ""))
        if st != "unverified":
            c.fill = FINDING_FILL.get(f_["severity"], FINDING_FILL["low"])
    if lines:
        c.comment = _note(lines)


def _review_sheet(wb: Workbook, findings: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Review")
    _set(ws.cell(1, 1), "Review findings")
    ws.cell(1, 1).font = Font(bold=True, size=13)
    _set(ws.cell(2, 1), "Open points from the automatic checks and the independent reviewer. Resolve them in the app.")
    ws.cell(2, 1).font = MUTED
    headers = ["Severity", "Where", "Finding", "Details", "From"]
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(4, ci)
        _set(c, h)
        c.font = BOLD
        c.fill = HEADER_FILL
        c.border = Border(bottom=THIN)
    order = {"high": 0, "medium": 1, "low": 2}
    for i, f in enumerate(sorted(findings, key=lambda x: order.get(x["severity"], 3)), start=5):
        sev = ws.cell(i, 1)
        _set(sev, f["severity"].title())
        sev.fill = FINDING_FILL.get(f["severity"], FINDING_FILL["low"])
        where = ws.cell(i, 2)
        target = f["target"]
        _set(where, target if target != "file" else "Whole file")
        if "!" in target:
            sheet, coord = target.rsplit("!", 1)
            if sheet in wb.sheetnames:
                where.hyperlink = Hyperlink(ref=where.coordinate, location=f"{_q(sheet)}!{coord}")
                where.font = LINK
        _set(ws.cell(i, 3), f["title"])
        _set(ws.cell(i, 4), f.get("body") or "")
        ws.cell(i, 4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="top")
        _set(ws.cell(i, 5), "Reviewer" if f.get("origin") == "reviewer" else "Automatic check")
    for col, w in zip("ABCDE", (10, 16, 46, 70, 16), strict=True):
        ws.column_dimensions[col].width = w


def _sources_sheet(wb: Workbook, model: dict[str, Any], meta: dict[str, Any], app_url: Callable[[str], str | None],
                   first_row: int) -> None:
    ws = wb.create_sheet("Sources")
    s = model.get("summary", {})
    total = s.get("cells_sourced", 0) + s.get("cells_unverified", 0)
    rows = [
        ("Sources and checks", None),
        (model["title"], None),
        (f"Prepared by {meta.get('author') or 'OfinTech Desk'} in {meta.get('workspace') or 'the workspace'}, "
         f"version {meta.get('version', 1)}, {datetime.fromtimestamp(model['created_at']).strftime('%Y-%m-%d %H:%M')}", None),
        (f"{total} input figures: {s.get('cells_sourced', 0)} traced to a source, {s.get('cells_unverified', 0)} "
         f"without a source; {s.get('cells_computed', 0)} figures are live formulas.", None),
        ("Blue figures are inputs (hover for their source); black figures are formulas.", None),
    ]
    for i, (text, _) in enumerate(rows, start=1):
        c = ws.cell(i, 1)
        _set(c, text)
        c.font = Font(bold=True, size=13) if i == 1 else (BOLD if i == 2 else MUTED)
    h = first_row - 1
    for ci, label in enumerate(["Ref", "Source", "System", "Retrieved", "What was used", "Open"], start=1):
        c = ws.cell(h, ci)
        _set(c, label)
        c.font = BOLD
        c.fill = HEADER_FILL
        c.border = Border(bottom=THIN)
    for i, src in enumerate(model["sources"]):
        r = first_row + i
        _set(ws.cell(r, 1), src["id"])
        ws.cell(r, 1).font = BOLD
        _set(ws.cell(r, 2), src["title"])
        _set(ws.cell(r, 3), src.get("system") or "")
        at = src.get("at")
        _set(ws.cell(r, 4), datetime.fromtimestamp(at).strftime("%Y-%m-%d %H:%M") if at else "")
        _set(ws.cell(r, 5), src.get("detail") or "")
        ws.cell(r, 5).alignment = Alignment(wrap_text=True, vertical="top")
        url = app_url(src["link"]) if src.get("link") else None
        if url:
            c = ws.cell(r, 6)
            _set(c, "Open in OfinTech Desk")
            c.hyperlink = url
            c.font = LINK
    for col, w in zip("ABCDEF", (6, 46, 26, 17, 80, 22), strict=True):
        ws.column_dimensions[col].width = w


def _calculations(ws: Any, model: dict[str, Any], row: int) -> None:
    _set(ws.cell(row, 1), "Calculations quoted in the report")
    ws.cell(row, 1).font = BOLD
    for ci, label in enumerate(["Ref", "Calculation", "Formula", "Value"], start=1):
        c = ws.cell(row + 1, ci)
        _set(c, label)
        c.font = BOLD
        c.fill = HEADER_FILL
    for i, calc in enumerate(model["calculations"], start=row + 2):
        _set(ws.cell(i, 1), calc["id"])
        _set(ws.cell(i, 2), calc.get("e") or "")
        vc = ws.cell(i, 4)
        if calc.get("a"):  # a figure from a table: point at that cell
            sheet, coord = calc["a"].rsplit("!", 1)
            col = "".join(ch for ch in coord if ch.isalpha())
            row_no = "".join(ch for ch in coord if ch.isdigit())
            _set(ws.cell(i, 3), f"{calc['a']}: {(calc.get('x') or '').lstrip('=')}")
            vc.value = f"={_q(sheet)}!${col}${row_no}"
        elif calc.get("x"):
            _set(ws.cell(i, 3), calc["x"].lstrip("="))
            vc.value = calc["x"]
        else:
            vc.value = calc.get("v")
        vc.number_format = NUMBER_FORMATS.get(calc.get("f") or "number", "#,##0.00")
