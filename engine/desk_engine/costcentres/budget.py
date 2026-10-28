"""Reading a budget from the customer's own Excel or Word file.

The file is read as it is: we look for a header row with an amount column (``Budget``, ``الموازنة``,
``Presupuesto``, ``Montant``…), take the label (and an account code, when there is one) from the same rows, skip
total and subtotal rows, and remember the sheet and cell of every figure so each budget line points back to the file.
A total row in the file is used as a check, not as a line. Lines are then matched to ledger accounts by code or by
name; people can correct the match.
"""

from __future__ import annotations

import io
import re
from typing import Any

from ..errors import Invalid
from .rules import norm

BUDGET_WORDS = ["budget", "approved budget", "الموازنة", "موازنة", "الميزانية", "ميزانية", "المعتمد", "presupuesto",
                "budget alloue", "budget approuve"]
AMOUNT_WORDS = ["amount", "value", "cost", "total cost", "المبلغ", "القيمة", "التكلفة", "importe", "monto", "coste",
                "costo", "montant", "cout", "valeur"]
LABEL_WORDS = ["item", "description", "line", "category", "account", "cost item", "البند", "الوصف", "البيان",
               "الحساب", "التصنيف", "partida", "concepto", "descripcion", "categoria", "cuenta", "poste", "libelle",
               "categorie", "compte"]
CODE_WORDS = ["code", "account code", "account no", "acct", "gl code", "رمز", "الرمز", "رقم الحساب", "كود",
              "codigo", "n de cuenta", "numero de compte", "n compte"]
TOTAL_WORDS = ["total", "grand total", "subtotal", "sub-total", "sub total", "الاجمالي", "اجمالي", "المجموع",
               "مجموع", "الإجمالي", "إجمالي", "suma", "total general", "sous-total", "somme"]
MAX_LINES = 500


def _is(text: str, words: list[str]) -> bool:
    t = norm(text)
    return bool(t) and any(t == norm(w) or t.startswith(norm(w) + " ") or t.endswith(" " + norm(w)) for w in words)


def _exact(text: str, words: list[str]) -> bool:
    """The whole header is the word, possibly followed by a unit: ``Budget``, ``Budget (SAR)``, ``المبلغ - ريال``."""
    t = re.sub(r"\s*[\(\[\-–].*$", "", norm(text))
    return bool(t) and any(t == norm(w) for w in words)


def _is_total(label: str) -> bool:
    t = norm(label)
    return any(t == norm(w) or t.startswith(norm(w) + " ") or t.startswith(norm(w) + ":") for w in TOTAL_WORDS)


def _num(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().translate(str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")).replace(",", "").replace(" ", "")
        neg = s.startswith("(") and s.endswith(")")
        s = s.strip("()")
        s = re.sub(r"^[^\d\-.]+|[^\d.]+$", "", s)  # currency symbols and codes around the number
        try:
            return -float(s) if neg else float(s)
        except ValueError:
            return None
    return None


def _text(v: Any) -> str:
    return str(v).strip() if v is not None and not isinstance(v, (int, float)) else ""


def _code(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s[:40] if s else None


def _header(rows: list[list[Any]]) -> tuple[int, int, int | None, int | None] | None:
    """(header row, amount column, label column, code column) in the first 40 rows."""
    for r, row in enumerate(rows[:40]):
        texts = [_text(v) for v in row]
        scores = [(4 if _exact(t, BUDGET_WORDS) else 3 if _exact(t, AMOUNT_WORDS) else
                   2 if _is(t, BUDGET_WORDS) else 1 if _is(t, AMOUNT_WORDS) and not _is(t, LABEL_WORDS) else 0)
                  for t in texts]
        if not any(scores) or sum(1 for t in texts if t) < 2:
            continue  # a header row has at least a label and an amount heading; a title line has one cell
        amount = max(range(len(scores)), key=lambda c: (scores[c], -c))
        code = next((c for c, t in enumerate(texts) if c != amount and _is(t, CODE_WORDS)), None)
        label = next((c for c, t in enumerate(texts) if c not in (amount, code) and _is(t, LABEL_WORDS)), None)
        if label is None:  # nearest text column to the left of the amounts
            label = next((c for c in range(amount - 1, -1, -1) if c != code and texts[c]), None)
        return r, amount, label, code
    return None


def _guess(rows: list[list[Any]]) -> tuple[int, int, int | None, int | None] | None:
    """No header row: the column with most numbers holds amounts, the text column with most values holds labels."""
    width = max((len(r) for r in rows), default=0)
    nums = [sum(1 for r in rows if c < len(r) and _num(r[c]) is not None and not isinstance(r[c], str)) for c in range(width)]
    texts = [sum(1 for r in rows if c < len(r) and _text(r[c]) and _num(r[c]) is None) for c in range(width)]
    if not width or max(nums) < 2:
        return None
    amount = max(range(width), key=lambda c: (nums[c], -c))
    label = max((c for c in range(width) if c != amount), key=lambda c: (texts[c], -c), default=None)
    return -1, amount, label, None


def _lines(rows: list[list[Any]], refs: list[list[str]], found: tuple[int, int, int | None, int | None],
           formulas: set[str] | None = None) -> dict[str, Any]:
    head, amount_col, label_col, code_col = found
    lines: list[dict[str, Any]] = []
    totals: list[dict[str, Any]] = []
    warnings: list[str] = []
    blank = 0
    for r in range(head + 1, len(rows)):
        row = rows[r]
        cell = lambda c: row[c] if c is not None and c < len(row) else None  # noqa: E731
        label = _text(cell(label_col)) or (_code(cell(label_col)) or "")
        raw = cell(amount_col)
        ref = refs[r][amount_col] if amount_col < len(refs[r]) else ""
        if not label and raw in (None, ""):
            blank += 1
            if blank >= 5 and lines:
                break
            continue
        blank = 0
        value = _num(raw)
        if value is None:
            if formulas and ref in formulas:
                warnings.append(f"{ref} has a formula with no saved result; open the file in Excel, save it and "
                                "import it again.")
            continue
        if _is_total(label) or (not label and lines):
            totals.append({"label": label or "(no label)", "amount": value, "cell": ref})
            continue
        if not label:
            continue
        lines.append({"label": label[:200], "amount": round(value, 2), "code": _code(cell(code_col)), "cell": ref})
        if len(lines) >= MAX_LINES:
            warnings.append(f"Only the first {MAX_LINES} lines were read.")
            break
    return {"lines": lines, "totals": totals, "warnings": warnings}


def _check(result: dict[str, Any]) -> dict[str, Any]:
    total = round(sum(x["amount"] for x in result["lines"]), 2)
    result["total"] = total
    grand = result["totals"][-1] if result["totals"] else None
    result["file_total"] = grand
    if grand is not None:
        if abs(grand["amount"] - total) <= 0.01:
            result["check"] = f"The lines add up to the file's total in {grand['cell']} ({grand['amount']:,.2f})."
        else:
            result["warnings"].append(f"The lines add up to {total:,.2f} but the file's total in {grand['cell']} is "
                                      f"{grand['amount']:,.2f}; check for subtotal rows or lines outside the table.")
    return result


def parse_xlsx(data: bytes, sheet: str | None = None) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
        wbf = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:  # noqa: BLE001 - any parse failure is the user's file
        raise Invalid("The workbook could not be opened.", code="file_format") from exc
    names = [ws.title for ws in wb.worksheets]
    if sheet and sheet not in names:
        raise Invalid(f"No sheet named “{sheet}”. Sheets: {', '.join(names)}.")
    tried = []
    for name in ([sheet] if sheet else names):
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 3000), values_only=True)]
        refs = [[f"{name}!{get_column_letter(c + 1)}{r + 1}" for c in range(len(row))] for r, row in enumerate(rows)]
        formulas = {f"{name}!{c.coordinate}" for row in wbf[name].iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")}
        found = _header(rows)
        guessed = found is None
        found = found or _guess(rows)
        if not found:
            tried.append(name)
            continue
        result = _lines(rows, refs, found, formulas)
        if not result["lines"]:
            tried.append(name)
            continue
        head, amount_col, label_col, code_col = found
        result |= {"format": "xlsx", "sheet": name, "sheets": names,
                   "columns": {"amount": get_column_letter(amount_col + 1),
                               "label": get_column_letter(label_col + 1) if label_col is not None else None,
                               "code": get_column_letter(code_col + 1) if code_col is not None else None},
                   "header_row": head + 1 if head >= 0 else None}
        if guessed:
            result["warnings"].insert(0, f"No header row was found on “{name}”; column {result['columns']['label']} "
                                         f"was read as labels and column {result['columns']['amount']} as amounts.")
        return _check(result)
    raise Invalid("No budget table was found. Put the lines in one table with a header row that has a label column "
                  "and an amount column (for example “Item” and “Budget”).", code="no_budget_table",
                  details={"sheets": tried})


def parse_docx(data: bytes) -> dict[str, Any]:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise Invalid("The document could not be opened.", code="file_format") from exc
    for t, table in enumerate(doc.tables, start=1):
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        refs = [[f"Table {t}, row {r + 1}, column {c + 1}" for c in range(len(row))] for r, row in enumerate(rows)]
        found = _header(rows)
        if not found:
            continue
        result = _lines(rows, refs, found)
        if result["lines"]:
            result |= {"format": "docx", "table": t, "header_row": found[0] + 1}
            return _check(result)
    raise Invalid("No budget table was found in the document. The budget must be a Word table with a header row that "
                  "has a label column and an amount column.", code="no_budget_table")


def _tokens(s: str) -> set[str]:
    return {w for w in re.split(r"[^\w]+", norm(s)) if len(w) >= 2 and not w.isdigit()}


def match_accounts(lines: list[dict[str, Any]], accounts: list[tuple[str | None, str]]) -> list[dict[str, Any]]:
    """Adds ``accounts`` (and how they were found) to each budget line: the file's account code when there is one,
    otherwise the ledger account whose name is closest to the label (at least half the words in common)."""
    known_codes = {str(c) for c, _ in accounts if c}
    out = []
    for ln in lines:
        code = ln.get("code")
        if code and any(k.startswith(code) for k in known_codes):
            out.append(ln | {"accounts": [code], "matched_by": "code"})
            continue
        words = _tokens(ln["label"])
        best, score = None, 0.0
        for c, name in accounts:
            other = _tokens(name)
            if not words or not other:
                continue
            s = len(words & other) / len(words | other)
            if norm(name) == norm(ln["label"]):
                s = 1.0
            if s > score:
                best, score = (c or name), s
        if best and score >= 0.5:
            out.append(ln | {"accounts": [best], "matched_by": "name", "match_score": round(score, 2)})
        else:
            out.append(ln | {"accounts": [], "matched_by": None})
    return out
