"""Excel templates people fill in and upload under Cost centres, and the reader for an uploaded general ledger.

* ``budget_template()``: one table of budget lines (code, item, budget, notes) with a total, read back by
  ``budget.parse_xlsx``.
* ``ledger_template()``: the general-ledger columns with example entries and a guide sheet, read back by
  ``read_ledger_file`` into the dataset cost centres, KPIs and forecasts use.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from ..connectors.accounting import ACCOUNT_CLASSES, GL_COLUMNS
from ..errors import Invalid

GL_NAMES = [c["name"] for c in GL_COLUMNS]
CLASS_ALIASES = {  # what people type -> the canonical class
    "cash": "Bank", "bank and cash": "Bank", "accounts receivable": "Receivable", "receivables": "Receivable",
    "debtors": "Receivable", "accounts payable": "Payable", "payables": "Payable", "creditors": "Payable",
    "assets": "Asset", "liabilities": "Liability", "revenue": "Income", "sales": "Income", "cogs": "Cost of sales",
    "cost of goods sold": "Cost of sales", "direct costs": "Cost of sales", "expenses": "Expense", "opex": "Expense",
    "other expenses": "Other expense", "other incomes": "Other income",
}


def _style(ws: Any, header_row: int, widths: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="18202B")
        c.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def budget_template(name: str = "Project budget", currency: str = "USD") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Code", "Item", f"Budget ({currency})", "Notes"])
    _style(ws, 1, [12, 40, 18, 48])
    examples = [("6100", "Site rent", 36000, "Account code 6100 in the ledger"),
                ("5100", "Subcontractors", 120000, "Code prefix 51 takes 5100, 5110…"),
                (None, "Contingency", 15000, "No code: match it to an account on the page, or leave it")]
    for row in examples:
        ws.append(list(row))
    for r in range(2, 2 + len(examples)):
        ws.cell(row=r, column=3).number_format = "#,##0.00"
    total_row = 2 + len(examples) + 8  # room for more lines
    ws.cell(row=total_row, column=2, value="Total").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row - 1})").font = Font(bold=True)
    ws.cell(row=total_row, column=3).number_format = "#,##0.00"
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Budget must be a number", errorTitle="Budget")
    ws.add_data_validation(dv)
    dv.add(f"C2:C{total_row - 1}")
    guide = wb.create_sheet("How to fill in")
    for line in [
        f"{name}: budget lines for a cost centre in OfinTech Desk",
        "",
        "1. One line per budget item on the Budget sheet. Keep the header row as it is.",
        "2. Code: the ledger account code the line covers (a prefix such as 51 covers 5100, 5110…). Leave it empty when there is none;",
        "   the page then matches the line by its name, and you can correct the match.",
        f"3. Budget: a number in {currency}. Do not type currency symbols.",
        "4. Total: the formula at the bottom is checked against the sum of the lines when the file is read.",
        "5. Rows that start with Total or Subtotal are not read as lines.",
        "6. Upload the file under Cost centres → Read budget from a file. Every line keeps the cell it came from.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ledger_template(currency: str = "USD") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "General ledger"
    ws.append(GL_NAMES)
    _style(ws, 1, [12, 16, 16, 13, 30, 15, 13, 13, 40, 24, 30])
    examples = [
        (date(2026, 3, 10), "BILL-0142", "Supplier bill", "5100", "Subcontractors", "Cost of sales", 10000, 0,
         "Steel works, phase 1", "Gulf Steel Co.", "Project: Riyadh tower"),
        (date(2026, 3, 10), "BILL-0142", "Supplier bill", "2000", "Accounts payable", "Payable", 0, 10000,
         "Steel works, phase 1", "Gulf Steel Co.", None),
        (date(2026, 4, 5), "PAY-0088", "Supplier payment", "2000", "Accounts payable", "Payable", 10000, 0,
         "Payment of BILL-0142", "Gulf Steel Co.", None),
        (date(2026, 4, 5), "PAY-0088", "Supplier payment", "1010", "Operating account", "Bank", 0, 10000,
         "Payment of BILL-0142", "Gulf Steel Co.", None),
        (date(2026, 4, 30), "JV-0031", "Journal", "6500", "Depreciation", "Expense", 1200, 0,
         "Fleet depreciation April", None, "Department: Operations (60%) | Department: Sales (40%)"),
        (date(2026, 4, 30), "JV-0031", "Journal", "1510", "Accumulated depreciation", "Asset", 0, 1200,
         "Fleet depreciation April", None, None),
    ]
    for row in examples:
        ws.append(list(row))
    for r in range(2, 2 + len(examples)):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=7).number_format = "#,##0.00"
        ws.cell(row=r, column=8).number_format = "#,##0.00"
    dv = DataValidation(type="list", formula1='"' + ",".join(ACCOUNT_CLASSES) + '"', allow_blank=False,
                        error="Choose one of the account classes listed on the Guide sheet", errorTitle="Account class")
    ws.add_data_validation(dv)
    dv.add("F2:F5000")
    guide = wb.create_sheet("Guide")
    lines = [
        "General ledger for OfinTech Desk: one row per journal line",
        "",
        "Keep the header row exactly as it is and delete the example rows before you upload.",
        "",
        "Date            yyyy-mm-dd (or an Excel date).",
        "Entry           the journal entry number. Lines of the same entry share it, and each entry balances (debits = credits).",
        "Type            Supplier bill, Customer invoice, Payment, Journal… (free text, optional).",
        "Account code    the account's code in your chart of accounts (optional but recommended).",
        "Account         the account's name.",
        "Account class   one of: " + ", ".join(ACCOUNT_CLASSES) + ".",
        "                Bank = bank and cash accounts; Payable = the supplier control account; Receivable = the customer control account.",
        f"Debit, Credit   positive amounts in {currency}; one of them is usually 0.",
        "Description     the line's narrative (optional). Words here can be used by cost-centre rules.",
        "Party           the customer or supplier (optional). Needed to match payments to the right bills.",
        "Cost centre tags  Project: Riyadh tower | Department: Sales (40%) — any dimension your system has (optional).",
        "",
        "Upload the file under Cost centres → Import a general ledger file, or under Data → Import. Cost centres,",
        "financial analysis and forecasts then read it. Upload it again after each month-end to replace the data.",
    ]
    for line in lines:
        guide.append([line])
    guide.column_dimensions["A"].width = 130
    guide["A1"].font = Font(bold=True, size=13)
    for r in range(5, 5 + 11):
        guide.cell(row=r, column=1).font = Font(name="Consolas")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm(s: Any) -> str:
    return " ".join(str(s or "").strip().lower().replace("_", " ").split())


def _cls(v: Any) -> str | None:
    t = _norm(v)
    for c in ACCOUNT_CLASSES:
        if t == c.lower():
            return c
    return CLASS_ALIASES.get(t)


def _date(v: Any) -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    from ..data.importer import parse_date

    parsed = parse_date(v) if v not in (None, "") else None
    return parsed.isoformat() if parsed else None


def _num(v: Any) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    from ..data.importer import parse_number

    n = parse_number(v)
    if n is None:
        raise Invalid(f"“{v}” is not a number.", code="bad_ledger")
    return float(n)


def read_ledger_file(data: bytes, filename: str) -> dict[str, Any]:
    """Finds the sheet with the general-ledger header, checks every row and returns clean rows for the dataset.

    Checks: dates, numbers, account classes (aliases such as 'Accounts receivable' are accepted), an account on every
    line, and that every entry balances. Problems name the row, so they can be fixed in the file."""
    from ..data.importer import read_workbook

    tables = read_workbook(data, filename)
    want = [n.lower() for n in GL_NAMES]
    chosen = None
    header_at = 0
    for t in tables:
        for i, row in enumerate(t.raw[:30]):
            names = [_norm(v) for v in row]
            if names[:len(want)] == want:
                chosen, header_at = t, i
                break
        if chosen:
            break
    if chosen is None:
        raise Invalid("No general-ledger sheet was found. The first row must have these columns in this order: "
                      + ", ".join(GL_NAMES) + ". Download the template to see an example.", code="not_a_ledger")
    rows: list[list[Any]] = []
    problems: list[str] = []
    sums: dict[str, list[float]] = {}
    for n, raw in enumerate(chosen.raw[header_at + 1:], start=header_at + 2):
        raw = list(raw) + [None] * (len(GL_NAMES) - len(raw))
        if all(v in (None, "") for v in raw[:len(GL_NAMES)]):
            continue
        day = _date(raw[0])
        entry = str(raw[1]).strip() if raw[1] not in (None, "") else ""
        code = raw[3]
        if isinstance(code, float) and code.is_integer():
            code = int(code)
        code = str(code).strip() if code not in (None, "") else None
        account = str(raw[4]).strip() if raw[4] not in (None, "") else ""
        cls = _cls(raw[5])
        try:
            debit, credit = _num(raw[6]), _num(raw[7])
        except Invalid as exc:
            problems.append(f"row {n}: {exc.message}")
            continue
        if not day:
            problems.append(f"row {n}: the date is missing or not a date")
        if not entry:
            problems.append(f"row {n}: the entry number is missing")
        if not account:
            problems.append(f"row {n}: the account name is missing")
        if cls is None:
            problems.append(f"row {n}: unknown account class “{raw[5]}”")
        if debit < 0 or credit < 0:
            problems.append(f"row {n}: debit and credit are positive numbers")
        if len(problems) >= 20:
            break
        sums.setdefault(entry, [0.0, 0.0])
        sums[entry][0] += debit
        sums[entry][1] += credit
        rows.append([day, entry, str(raw[2]).strip()[:60] if raw[2] not in (None, "") else None, code, account, cls,
                     round(debit, 2), round(credit, 2), str(raw[8]).strip()[:500] if raw[8] not in (None, "") else None,
                     str(raw[9]).strip()[:200] if raw[9] not in (None, "") else None,
                     str(raw[10]).strip()[:300] if raw[10] not in (None, "") else None])
    unbalanced = [e for e, (d, c) in sums.items() if abs(d - c) > 0.01]
    if problems:
        raise Invalid("The file has problems: " + "; ".join(problems[:20]) + (" …" if len(problems) >= 20 else ""),
                      code="bad_ledger", details={"problems": problems[:20]})
    if not rows:
        raise Invalid("The sheet has no journal lines under the header row.", code="bad_ledger")
    dates = [r[0] for r in rows]
    return {"rows": rows, "sheet": chosen.name, "lines": len(rows), "entries": len(sums), "first_date": min(dates),
            "last_date": max(dates), "unbalanced_entries": unbalanced[:20], "unbalanced_count": len(unbalanced)}
